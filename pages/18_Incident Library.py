import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import os
import base64
import re
import json
from io import BytesIO
from urllib.request import Request, urlopen
from openpyxl import load_workbook
from streamlit_autorefresh import st_autorefresh
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo

# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="PSM Digital Dashboard",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ============================================================
# REMOVE STREAMLIT TOP SPACE
# ============================================================

st.markdown(
    """
    <style>

    html,
    body {
        margin: 0 !important;
        padding: 0 !important;
    }

    [data-testid="stAppViewContainer"] {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }

    [data-testid="stAppViewContainer"] > .main {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }

    [data-testid="stHeader"] {
        display: none !important;
        height: 0 !important;
        min-height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    [data-testid="stToolbar"] {
        display: none !important;
    }

    [data-testid="stDecoration"] {
        display: none !important;
        height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    .block-container {
        padding-top: 0 !important;
        margin-top: 0 !important;
        padding-bottom: 0 !important;
        padding-left: 0 !important;
        padding-right: 0 !important;
        max-width: 100% !important;
    }

    .stApp {
        margin-top: 0 !important;
        padding-top: 0 !important;
    }

    iframe {
        display: block !important;
        margin-top: -8px !important;
        padding-top: 0 !important;
        border: 0 !important;
    }

    </style>
    """,
    unsafe_allow_html=True
)

# ============================================================
# BASE DIRECTORY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent


# ============================================================
# IMAGE TO BASE64
# ============================================================

def image_to_base64(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        return ""

    try:

        with open(file_path, "rb") as file:

            return base64.b64encode(
                file.read()
            ).decode("utf-8")

    except Exception:

        return ""


# ============================================================
# COMPANY LOGO
# ============================================================

logo_path = BASE_DIR / "jsw_jfe_logo.jpg"

logo_base64 = image_to_base64(logo_path)

# ============================================================
# FILE CHECK
# ============================================================

if not logo_base64:
    st.error(
        "jsw_jfe_logo.jpg not found. "
        "Keep jsw_jfe_logo.jpg in the same folder as this Python file."
    )

# ============================================================
# DATE AND TIME
# ============================================================

now = datetime.now(ZoneInfo("Asia/Kolkata"))

current_date = now.strftime(
    "%d %b %Y"
).upper()

current_time = now.strftime(
    "%I:%M %p"
)

# ============================================================
# HEADER HTML
# ============================================================

header_html = """

<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<style>


/* ============================================================
   MAIN HEADER
   ============================================================ */

.psm-header {

    position: relative;
    width: 100%;
    padding: 0 !important;

    height: 90px;

    overflow: hidden;

    background:
        linear-gradient(
            90deg,
            #031d34 0%,
            #052b49 42%,
            #07385c 74%,
            #052b49 100%
        );

    border-radius: 7px;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.18);

}


/* ============================================================
   LEFT LOGO AREA
   ============================================================ */

.psm-left {

    position: absolute;

    left: 0;

    top: 0;

    width: 100%;

    height: 95px;

    display: flex;

    align-items: center;

    padding-left: 4px;

    box-sizing: border-box;

    z-index: 20;

    pointer-events: none;

}


/* ============================================================
   LOGO PANEL
   ONLY VERTICAL POSITION CHANGED
   ============================================================ */

.logo-panel {

    width: 195px;

    height: 68px;

    background: #ffffff;

    border-radius: 5px;

    padding: 4px;

    display: flex;

    align-items: center;

    justify-content: center;

    flex-shrink: 0;

    box-sizing: border-box;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.20);

    position: relative;

    top: -6px;

    left: 5px;

}


/* ============================================================
   COMPANY LOGO
   ============================================================ */

.company-logo {

    width: 100%;

    height: 100%;

    object-fit: contain;

    object-position: center;

    display: block;

}


/* ============================================================
   LEFT VERTICAL DIVIDER
   ============================================================ */

.vertical-line {

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

    margin-left: 18px;

    margin-right: 20px;

    flex-shrink: 0;

}


/* ============================================================
   CENTER TITLE AREA
   ============================================================ */

.title-area {

    position: absolute;

    left: 50%;

    top: 0;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: center;

    text-align: center;

    min-width: max-content;

    box-sizing: border-box;

    transform: translateX(-50%);

}


/* ============================================================
   MAIN TITLE
   ============================================================ */

.main-title {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 27px;

    font-weight: 900;

    line-height: 1;

    letter-spacing: 0.3px;

    white-space: nowrap;

    margin: 0;

    padding: 0;

}


/* ============================================================
   ORANGE TITLE PART
   ============================================================ */

.main-title-orange {

    color: #f28c00;

}


/* ============================================================
   SUBTITLE
   ============================================================ */

.subtitle {

    color: #ffffff;

    font-family:
        Arial,
        sans-serif;

    font-size: 12px;

    font-weight: 400;

    letter-spacing: 3.6px;

    margin-top: 8px;

    line-height: 1;

    white-space: nowrap;

}

/* ============================================================
   SUB-SUBTITLE / TAGLINE
   ============================================================ */

.tagline {

    color:
        rgba(255,255,255,0.82);

    font-family:
        Arial,
        sans-serif;

    font-size: 7px;

    font-weight: 500;

    letter-spacing: 2.2px;

    margin-top: 6px;

    line-height: 1;

    white-space: nowrap;

}


/* ============================================================
   RIGHT DATE / TIME AREA
   ============================================================ */

.psm-right {

    position: absolute;

    right: 16px;

    top: 0;

    width: 15%;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: flex-end;

    text-align: right;

    color: #ffffff;

    z-index: 30;

    padding-left: 18px;

    box-sizing: border-box;

}


/* ============================================================
   RIGHT VERTICAL DIVIDER
   ============================================================ */

.psm-right::before {

    content: "";

    position: absolute;

    left: 15px;

    top: 6px;

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

}


/* ============================================================
   DATE
   ============================================================ */

.date {

    color:
        rgba(255,255,255,0.95);

    font-family:
        Arial,
        sans-serif;

    font-size: 11px;

    font-weight: 400;

    letter-spacing: 0.7px;

    line-height: 1;

    margin: 0;

    padding: 0;

}


/* ============================================================
   TIME
   ============================================================ */

.time {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 22px;

    font-weight: 800;

    margin-top: 4px;

    line-height: 1;

    padding: 0;

}


/* ============================================================
   RIGHT HORIZONTAL LINE
   ============================================================ */

.right-line {

    width: 80px;

    height: 2px;

    background:
        rgba(255,255,255,0.75);

    margin-top: 7px;

    flex-shrink: 0;

}


/* ============================================================
   ORANGE BOTTOM BAR
   ============================================================ */

.orange-bar {

    position: absolute;

    left: 0;

    bottom: 0;

    width: 100%;

    height: 7px;

    background: #f28c00;

    z-index: 50;

}


</style>

</head>


<body>


<!-- ============================================================
     MAIN HEADER
     ============================================================ -->

<div class="psm-header">


    <!-- ========================================================
         LEFT LOGO AREA
         ======================================================== -->

    <div class="psm-left">


        <!-- ====================================================
             LOGO
             ==================================================== -->

        <div class="logo-panel">

            <img
                class="company-logo"
                src="data:image/jpeg;base64,LOGO_IMAGE_BASE64"
                alt="JSW JFE Steel Limited"
            >

        </div>


        <!-- ====================================================
             LEFT VERTICAL LINE
             ==================================================== -->

        <div class="vertical-line"></div>


        <!-- ====================================================
             CENTER TITLE GROUP
             ==================================================== -->

        <div class="title-area">


            <!-- MAIN TITLE -->

            <div class="main-title">

                INCIDENT LEARNING LIBRARY

                <span class="main-title-orange"></span>

            </div>


            <!-- SUBTITLE -->

            <div class="subtitle">

                PSM DIGITAL DASHBOARD

            </div>


            <!-- SUB-SUBTITLE -->

            <div class="tagline">

                PEOPLE
                &nbsp; | &nbsp;
                PROCESS
                &nbsp; | &nbsp;
                RISK
                &nbsp; | &nbsp;
                COMPLIANCE

            </div>


        </div>


    </div>


    <!-- ========================================================
         RIGHT DATE / TIME
         ======================================================== -->

    <div class="psm-right">


        <div class="date">

            CURRENT_DATE_VALUE

        </div>


        <div class="time">

            CURRENT_TIME_VALUE

        </div>


        <div class="right-line"></div>


    </div>


    <!-- ========================================================
         ORANGE BOTTOM BAR
         ======================================================== -->

    <div class="orange-bar"></div>


</div>


</body>

</html>

"""

# ============================================================
# INSERT LOGO
# ============================================================

header_html = header_html.replace(
    "LOGO_IMAGE_BASE64",
    logo_base64
)

# ============================================================
# INSERT DATE
# ============================================================

header_html = header_html.replace(
    "CURRENT_DATE_VALUE",
    current_date
)

# ============================================================
# INSERT TIME
# ============================================================

header_html = header_html.replace(
    "CURRENT_TIME_VALUE",
    current_time
)

# ============================================================
# DISPLAY HEADER
# ============================================================

components.html(
    header_html,
    height=114,
    scrolling=False
)

import re
import math
import html
import requests
import pandas as pd
import streamlit as st

from io import BytesIO
from openpyxl import load_workbook
from streamlit_autorefresh import st_autorefresh


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="JSW JFE Steel Process Safety Incident Learning Library",
    page_icon="📘",
    layout="wide",
    initial_sidebar_state="collapsed"
)


# ============================================================
# CONFIGURATION
# ============================================================

SPREADSHEET_ID = "1--X0TT5Ts92EKAxrhV-fQgqeTHBX3rDVc1Egg74MewM"

SHEET_NAME = "INCIDENT LIBRARY"

ROWS_PER_PAGE = 6

# Automatically refresh every 30 seconds
REFRESH_INTERVAL = 30 * 1000


# ============================================================
# AUTOMATIC REFRESH
# ============================================================

st_autorefresh(
    interval=REFRESH_INTERVAL,
    key="incident_library_auto_refresh"
)


# ============================================================
# GLOBAL CSS
# ============================================================

st.markdown(
    """
    <style>

    .stApp {
        background: #ffffff;
    }

    .block-container {

        max-width: 1500px !important;

        padding-top: 0.15rem !important;

        padding-bottom: 0.2rem !important;

        padding-left: 1.1rem !important;

        padding-right: 1.1rem !important;
    }

    header {
        visibility: hidden;
        height: 0 !important;
    }

    footer {
        visibility: hidden;
        height: 0 !important;
    }


    /* REMOVE EXCESS STREAMLIT SPACING */

    div[data-testid="stVerticalBlock"] {
        gap: 0.15rem;
    }

    div[data-testid="stHorizontalBlock"] {
        gap: 0.7rem;
    }


    /* ======================================================
       SEARCH BOX
       ====================================================== */

    div[data-testid="stTextInput"] {

        margin: 0 !important;
        padding: 0 !important;
    }

    div[data-testid="stTextInput"] > div {

        margin: 0 !important;
        padding: 0 !important;
    }

    div[data-testid="stTextInput"] input {
    
        height: 38px !important;

        min-height: 38px !important;

        padding-left: 11px !important;

        padding-right: 10px !important;

        border:
            1px solid #d7e2ec !important;

        border-radius:
            8px !important;

        background:
            #ffffff !important;

        color:
            #657487 !important;

        font-size:
            13px !important;

        box-shadow:
            0 1px 5px
            rgba(30,80,130,0.04) !important;
    }

    div[data-testid="stTextInput"] input:focus {

        border-color:
            #9dbdd5 !important;

        box-shadow:
            0 0 0 1px
            #9dbdd5 !important;
    }


    /* ======================================================
       PAGINATION
       ====================================================== */

    div[data-testid="stButton"] {

        margin: 0 !important;

        padding: 0 !important;
    }

    div[data-testid="stButton"] button {

        width:
            34px !important;

        min-width:
            34px !important;

        height:
            31px !important;

        min-height:
            31px !important;

        padding:
            0 !important;

        border:
            1px solid #d4dfe9 !important;

        border-radius:
            7px !important;

        background:
            #ffffff !important;

        color:
            #60748a !important;

        font-size:
            17px !important;

        font-weight:
            400 !important;
    }

    div[data-testid="stButton"] button:hover {

        background:
            #f3f8fc !important;

        color:
            #174f8f !important;

        border-color:
            #a7c3d9 !important;
    }

    div[data-testid="stButton"] button:disabled {

        opacity:
            0.35 !important;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# DOWNLOAD GOOGLE SHEET
# ============================================================

@st.cache_data(
    ttl=20,
    show_spinner=False
)
def download_sheet():

    url = (
        "https://docs.google.com/spreadsheets/d/"
        + SPREADSHEET_ID
        + "/export?format=xlsx"
    )

    response = requests.get(
        url,
        timeout=40,
        headers={
            "User-Agent": "Mozilla/5.0"
        }
    )

    response.raise_for_status()

    return response.content


# ============================================================
# GET URL FROM HYPERLINK FORMULA
# ============================================================

def get_formula_url(value):

    if value is None:
        return ""

    value = str(value).strip()

    match = re.search(
        r'HYPERLINK\s*\(\s*"([^"]+)"',
        value,
        re.IGNORECASE
    )

    if match:
        return match.group(1).strip()

    return ""


# ============================================================
# GET LINK FROM CELL
# ============================================================

def get_cell_link(cell):

    # Normal Excel hyperlink
    try:

        if cell.hyperlink:

            if cell.hyperlink.target:

                return str(
                    cell.hyperlink.target
                ).strip()

    except Exception:

        pass


    # HYPERLINK formula
    formula_link = get_formula_url(
        cell.value
    )

    if formula_link:
        return formula_link


    # Direct URL
    if isinstance(
        cell.value,
        str
    ):

        value = cell.value.strip()

        if (
            value.startswith("http://")
            or value.startswith("https://")
        ):

            return value


    return ""


# ============================================================
# LOAD INCIDENT LIBRARY
# ============================================================

@st.cache_data(
    ttl=20,
    show_spinner=False
)
def load_data():

    content = download_sheet()

    workbook = load_workbook(
        filename=BytesIO(content),
        data_only=False
    )

    if SHEET_NAME not in workbook.sheetnames:

        raise Exception(
            "INCIDENT LIBRARY tab was not found. "
            "Available tabs: "
            + ", ".join(
                workbook.sheetnames
            )
        )

    ws = workbook[SHEET_NAME]

    max_row = ws.max_row
    max_col = ws.max_column

    if max_row < 2:

        raise Exception(
            "INCIDENT LIBRARY is empty."
        )


    # ========================================================
    # HEADERS
    # ========================================================

    headers = []

    for col in range(
        1,
        max_col + 1
    ):

        value = ws.cell(
            row=1,
            column=col
        ).value

        if value is None:

            headers.append("")

        else:

            headers.append(
                re.sub(
                    r"\s+",
                    " ",
                    str(value)
                ).strip()
            )


    valid_cols = [
        i
        for i, header in enumerate(headers)
        if header
    ]

    headers = [
        headers[i]
        for i in valid_cols
    ]


    # ========================================================
    # DATA + LINKS
    # ========================================================

    records = []
    links = []

    for row_number in range(
        2,
        max_row + 1
    ):

        row_values = []
        row_links = []

        for index in valid_cols:

            column_number = index + 1

            cell = ws.cell(
                row=row_number,
                column=column_number
            )

            value = cell.value

            if value is None:

                text = ""

            else:

                text = str(
                    value
                ).strip()


            # -----------------------------------------------
            # HYPERLINK DISPLAY TEXT
            # -----------------------------------------------

            display_match = re.search(
                r'HYPERLINK\s*\(\s*"[^"]+"\s*,\s*"([^"]*)"',
                text,
                re.IGNORECASE
            )

            if display_match:

                text = display_match.group(1)


            row_values.append(text)

            row_links.append(
                get_cell_link(cell)
            )


        # Ignore completely empty rows

        if any(
            str(x).strip()
            for x in row_values
        ):

            records.append(
                row_values
            )

            links.append(
                row_links
            )


    df = pd.DataFrame(
        records,
        columns=headers
    )

    links_df = pd.DataFrame(
        links,
        columns=headers
    )

    return df, links_df


# ============================================================
# NORMALIZE COLUMN NAME
# ============================================================

def normalize(value):

    return re.sub(
        r"[^a-z0-9]",
        "",
        str(value).lower()
    )


# ============================================================
# FIND COLUMN
# ============================================================

def find_column(
    columns,
    names
):

    normalized_columns = {
        normalize(column): column
        for column in columns
    }


    # Exact match

    for name in names:

        key = normalize(name)

        if key in normalized_columns:

            return normalized_columns[key]


    # Partial match

    for column in columns:

        column_key = normalize(
            column
        )

        for name in names:

            name_key = normalize(
                name
            )

            if (
                name_key in column_key
                or
                column_key in name_key
            ):

                return column


    return None


# ============================================================
# SERIAL NUMBER
# ============================================================

def format_serial(
    value,
    fallback
):

    if value is None:

        return str(fallback)

    text = str(
        value
    ).strip()

    if not text:

        return str(fallback)

    if re.fullmatch(
        r"\d+\.0",
        text
    ):

        text = text[:-2]

    return text


# ============================================================
# DATE
# ============================================================

def format_date(value):

    if value is None:

        return ""

    if not str(value).strip():

        return ""

    try:

        date_value = pd.to_datetime(
            value,
            errors="coerce"
        )

        if not pd.isna(
            date_value
        ):

            return date_value.strftime(
                "%d %b %Y"
            )

    except Exception:

        pass

    return str(
        value
    ).strip()


# ============================================================
# SAFE HTML
# ============================================================

def safe(value):

    return html.escape(
        str(value),
        quote=False
    )


# ============================================================
# COMPACT VIEW BUTTON
# ============================================================

def view_button(url):

    if not url:

        return """
        <span style="
            color:#8a97a6;

            font-size:11px;

            font-family:
                Arial,
                Helvetica,
                sans-serif;
        ">
            Not available
        </span>
        """


    url = html.unescape(
        str(url).strip()
    )

    safe_url = html.escape(
        url,
        quote=True
    )


    return f"""
    <a
        href="{safe_url}"
        target="_blank"
        rel="noopener noreferrer"

        style="
            display:flex;

            align-items:center;

            justify-content:center;

            width:82px;

            height:26px;

            margin:0 auto;

            box-sizing:border-box;

            background:#0b7fd3;

            border:none;

            border-radius:5px;

            color:#ffffff;

            text-decoration:none;

            font-family:
                Arial,
                Helvetica,
                sans-serif;

            font-size:12px;

            font-weight:700;

            line-height:26px;

            text-align:center;
        "
    >
        View
    </a>
    """


# ============================================================
# LOAD DATA
# ============================================================

try:

    df, links_df = load_data()

except Exception as error:

    st.error(
        "Unable to load INCIDENT LIBRARY."
    )

    st.error(
        str(error)
    )

    st.stop()


# ============================================================
# IDENTIFY COLUMNS
# ============================================================

columns = list(
    df.columns
)


sno_col = find_column(
    columns,
    [
        "S. No.",
        "S No",
        "S.No",
        "Sr No",
        "Serial No",
        "Sl No"
    ]
)


incident_col = find_column(
    columns,
    [
        "Incident Name",
        "Incident",
        "Incident Title",
        "Name"
    ]
)


date_col = find_column(
    columns,
    [
        "Date",
        "Incident Date",
        "Event Date"
    ]
)


report_col = find_column(
    columns,
    [
        "Complete Report",
        "Complete Report Link",
        "Full Report",
        "Report"
    ]
)


summary_col = find_column(
    columns,
    [
        "One Pager Summary",
        "One Pager",
        "One-Pager Summary",
        "Summary"
    ]
)


# ============================================================
# VALIDATION
# ============================================================

required = {
    "Incident Name": incident_col,
    "Date": date_col,
    "Complete Report": report_col,
    "One Pager Summary": summary_col
}


missing = [
    name
    for name, value in required.items()
    if value is None
]


if missing:

    st.error(
        "Required columns are missing."
    )

    for item in missing:

        st.write(
            "• " + item
        )

    st.write(
        "Columns detected:"
    )

    st.code(
        "\n".join(
            columns
        )
    )

    st.stop()


# ============================================================
# ABOUT THIS LIBRARY
# ============================================================

st.html(
    """
    <div style="
        width:100%;

        box-sizing:border-box;

        background:
            linear-gradient(
                135deg,
                #f8fbff 0%,
                #edf6ff 55%,
                #e1f0ff 100%
            );

        border-left:
            5px solid #ff3038;

        border-radius:
            0 12px 12px 0;

        padding:
            10px 24px 11px 24px;

        margin:
            0 0 8px 0;

        box-shadow:
            0 2px 9px
            rgba(20,80,140,0.05);

        font-family:
            Arial,
            Helvetica,
            sans-serif;
    ">

        <div style="
            color:#123f96;

            font-size:25px;

            font-weight:800;

            line-height:1.05;

            margin:0;
        ">
            About This Library
        </div>


        <div style="
            width:76px;

            height:3px;

            background:#ff3038;

            margin:3px 0 6px 0;
        "></div>


        <div style="
            color:#123f9b;

            font-size:13 px;

            line-height:1.38;

            margin:0;
        ">

            As part of our continued commitment to strengthen
            <strong>
                Process Safety Management (PSM)
            </strong>
            culture and proactively prevent major process safety
            incidents, we are pleased to share the
            <strong>
                “JSW JFE Steel Process Safety Incident Learning Library”
            </strong>
            at
            <strong>
                JSW JFE Steel Ltd., Sambalpur.
            </strong>

            <br>

            This library contains significant process safety
            incidents from around the world, enabling us to learn
            from past events, understand the root causes, and
            apply the lessons learned to prevent similar incidents
            in our operations.

        </div>

    </div>
    """
)


# ============================================================
# INCIDENT REGISTER + SEARCH
# ============================================================

title_col, search_col = st.columns(
    [2.6, 1],
    gap="medium"
)


with title_col:

    st.html(
        """
        <div style="
            color:#123f96;

            font-family:
                Arial,
                Helvetica,
                sans-serif;

            font-size:27px;

            font-weight:800;

            line-height:1.05;

            margin:0;
        ">
            Incident Register
        </div>


        <div style="
            width:76px;

            height:3px;

            background:#ff3038;

            margin:3px 0 0 0;
        "></div>
        """
    )


with search_col:

    search = st.text_input(
        "Search",
        placeholder=(
            "🔍  Search incidents by name, date, etc..."
        ),
        label_visibility="collapsed",
        key="incident_search"
    )


# ============================================================
# FILTER DATA
# ============================================================

filtered_df = df.copy()

filtered_links = links_df.copy()


if search.strip():

    query = (
        search
        .strip()
        .lower()
    )


    mask = (
        filtered_df
        .astype(str)
        .apply(
            lambda col:
                col
                .str
                .lower()
                .str
                .contains(
                    query,
                    regex=False,
                    na=False
                )
        )
        .any(axis=1)
    )


    filtered_df = (
        filtered_df
        .loc[mask]
        .reset_index(drop=True)
    )


    filtered_links = (
        filtered_links
        .loc[mask]
        .reset_index(drop=True)
    )


# ============================================================
# PAGINATION
# ============================================================

total_records = len(
    filtered_df
)


total_pages = max(
    1,
    math.ceil(
        total_records /
        ROWS_PER_PAGE
    )
)


if "incident_page" not in st.session_state:

    st.session_state.incident_page = 1


if "last_incident_search" not in st.session_state:

    st.session_state.last_incident_search = ""


if (
    st.session_state.last_incident_search
    != search
):

    st.session_state.incident_page = 1

    st.session_state.last_incident_search = search


if (
    st.session_state.incident_page
    > total_pages
):

    st.session_state.incident_page = (
        total_pages
    )


current_page = (
    st.session_state.incident_page
)


# ============================================================
# CURRENT PAGE
# ============================================================

start = (
    current_page - 1
) * ROWS_PER_PAGE


end = (
    start +
    ROWS_PER_PAGE
)


page_df = (
    filtered_df
    .iloc[start:end]
    .reset_index(drop=True)
)


page_links = (
    filtered_links
    .iloc[start:end]
    .reset_index(drop=True)
)


# ============================================================
# BUILD TABLE
# ============================================================

rows_html = ""


for i in range(
    len(page_df)
):

    row = page_df.iloc[i]


    # Serial number

    if sno_col:

        serial = format_serial(
            row[sno_col],
            start + i + 1
        )

    else:

        serial = str(
            start + i + 1
        )


    # Incident

    incident = safe(
        row[incident_col]
    )


    # Date

    date = safe(
        format_date(
            row[date_col]
        )
    )


    # Complete report

    report_url = ""

    if report_col in page_links.columns:

        value = page_links.iloc[i][
            report_col
        ]

        if pd.notna(value):

            report_url = str(
                value
            ).strip()


    # One pager

    summary_url = ""

    if summary_col in page_links.columns:

        value = page_links.iloc[i][
            summary_col
        ]

        if pd.notna(value):

            summary_url = str(
                value
            ).strip()


    report_button = view_button(
        report_url
    )


    summary_button = view_button(
        summary_url
    )


    row_background = (
        "#ffffff"
        if i % 2 == 0
        else "#f4f8fb"
    )


    rows_html += f"""
    <tr style="
        background:{row_background};
    ">

        <td style="
            width:8%;
            height:34px;
            padding:3px 5px;
            text-align:center;
            vertical-align:middle;
            color:#174b96;
            font-size:13px;
            border-right:1px solid #d8e4ee;
            border-bottom:1px solid #d8e4ee;
        ">
            {safe(serial)}
        </td>


        <td style="
            width:34%;
            height:34px;
            padding:3px 9px;
            text-align:left;
            vertical-align:middle;
            color:#174b96;
            font-size:13px;
            line-height:1.18;
            border-right:1px solid #d8e4ee;
            border-bottom:1px solid #d8e4ee;
        ">
            {incident}
        </td>


        <td style="
            width:17%;
            height:34px;
            padding:3px 5px;
            text-align:center;
            vertical-align:middle;
            color:#174b96;
            font-size:13px;
            border-right:1px solid #d8e4ee;
            border-bottom:1px solid #d8e4ee;
        ">
            {date}
        </td>


        <td style="
            width:20.5%;
            height:34px;
            padding:2px 4px;
            text-align:center;
            vertical-align:middle;
            border-right:1px solid #d8e4ee;
            border-bottom:1px solid #d8e4ee;
        ">
            {report_button}
        </td>


        <td style="
            width:20.5%;
            height:34px;
            padding:2px 4px;
            text-align:center;
            vertical-align:middle;
            border-bottom:1px solid #d8e4ee;
        ">
            {summary_button}
        </td>

    </tr>
    """


# ============================================================
# TABLE DISPLAY
# ============================================================

if not page_df.empty:

    table_html = f"""
    <div style="
        width:100%;
        box-sizing:border-box;
        background:#ffffff;
        border:1px solid #dfe9f2;
        border-radius:11px;
        padding:5px;
        margin:5px 0 0 0;
        box-shadow:
            0 2px 10px
            rgba(20,80,135,0.07);
        overflow:hidden;
    ">

        <table style="
            width:100%;
            table-layout:fixed;
            border-collapse:separate;
            border-spacing:0;
            border-radius:7px;
            overflow:hidden;
            font-family:
                Arial,
                Helvetica,
                sans-serif;
        ">

            <thead>

                <tr>

                    <th style="
                        width:8%;
                        height:39px;
                        padding:7px 5px;
                        background:#c8e3f5;
                        color:#164b91;
                        font-size:13px;
                        font-weight:800;
                        text-align:center;
                        border-right:1px solid #a8cbe3;
                        border-bottom:1px solid #a8cbe3;
                    ">
                        S. No.
                    </th>


                    <th style="
                        width:34%;
                        height:39px;
                        padding:7px 5px;
                        background:#c8e3f5;
                        color:#164b91;
                        font-size:13px;
                        font-weight:800;
                        text-align:center;
                        border-right:1px solid #a8cbe3;
                        border-bottom:1px solid #a8cbe3;
                    ">
                        Incident Name
                    </th>


                    <th style="
                        width:17%;
                        height:39px;
                        padding:7px 5px;
                        background:#c8e3f5;
                        color:#164b91;
                        font-size:13px;
                        font-weight:800;
                        text-align:center;
                        border-right:1px solid #a8cbe3;
                        border-bottom:1px solid #a8cbe3;
                    ">
                        Date
                    </th>


                    <th style="
                        width:20.5%;
                        height:39px;
                        padding:7px 5px;
                        background:#c8e3f5;
                        color:#164b91;
                        font-size:13px;
                        font-weight:800;
                        text-align:center;
                        border-right:1px solid #a8cbe3;
                        border-bottom:1px solid #a8cbe3;
                    ">
                        Complete Report
                    </th>


                    <th style="
                        width:20.5%;
                        height:39px;
                        padding:7px 5px;
                        background:#c8e3f5;
                        color:#164b91;
                        font-size:13px;
                        font-weight:800;
                        text-align:center;
                        border-bottom:1px solid #a8cbe3;
                    ">
                        One Pager Summary
                    </th>

                </tr>

            </thead>


            <tbody>

                {rows_html}

            </tbody>

        </table>

    </div>
    """


    st.html(
        table_html
    )


else:

    st.html(
        """
        <div style="
            margin-top:5px;
            padding:25px;
            text-align:center;
            border:1px solid #dfe8f1;
            border-radius:9px;
            color:#7c8b9b;
            font-size:13px;
        ">
            No incidents found.
        </div>
        """
    )


# ============================================================
# PAGINATION FOOTER
# ============================================================

blank_col, info_col, previous_col, next_col = st.columns(
    [7.7, 1.65, 0.35, 0.35],
    gap="small"
)


# ============================================================
# PAGE INFORMATION
# ============================================================

with info_col:

    st.html(
        f"""
        <div style="
            text-align:right;
            white-space:nowrap;
            padding-top:4px;
            font-family:
                Arial,
                Helvetica,
                sans-serif;
            font-size:12px;
            color:#68788b;
            font-weight:600;
        ">

            <span style="
                color:#315b91;
                font-weight:700;
            ">
                Page {current_page} of {total_pages}
            </span>

            <span style="
                margin-left:7px;
            ">
                • {total_records} records
            </span>

        </div>
        """
    )


# ============================================================
# PREVIOUS
# ============================================================

with previous_col:

    previous_clicked = st.button(
        "‹",
        key="incident_previous",
        disabled=(
            current_page <= 1
        )
    )


if previous_clicked:

    st.session_state.incident_page = max(
        1,
        current_page - 1
    )

    st.rerun()


# ============================================================
# NEXT
# ============================================================

with next_col:

    next_clicked = st.button(
        "›",
        key="incident_next",
        disabled=(
            current_page >= total_pages
        )
    )


if next_clicked:

    st.session_state.incident_page = min(
        total_pages,
        current_page + 1
    )

    st.rerun()

