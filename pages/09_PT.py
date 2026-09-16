import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import os
import base64
import re
import json
from io import BytesIO
from urllib.request import Request, urlopen
from openpyxl import load_workbook
from streamlit_autorefresh import st_autorefresh
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="PSM Dashboard - PT",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)


# =========================================================
# SESSION STATE
# =========================================================

if "status_filter" not in st.session_state:
    st.session_state.status_filter = "All"

if "page_number" not in st.session_state:
    st.session_state.page_number = 1

if "department_selector" not in st.session_state:
    st.session_state.department_selector = "All Departments"




# =========================================================
# AUTO REFRESH
# =========================================================

st_autorefresh(
    interval=100000,
    key="psm_auto_refresh"
)


# =========================================================
# GOOGLE SHEET - SHEET2
# =========================================================

SPREADSHEET_ID = "1--X0TT5Ts92EKAxrhV-fQgqeTHBX3rDVc1Egg74MewM"

SHEET2_CSV_URL = (
    f"https://docs.google.com/spreadsheets/d/"
    f"{SPREADSHEET_ID}"
    f"/gviz/tq?tqx=out:csv&sheet=PT"
)


# These are the exact Google Sheet column names used by the PT register.
STATUS_COLUMN = "Status (Ongoing/Completed)"
LINK_COLUMN = "Attach PT Softcopy Link"
APPROVAL_COLUMN = "Approved  (Yes/No)"


def normalize_column_name(value):
    """Normalize Google Sheet headers so line breaks/multiple spaces do not matter."""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ").replace("\n", " ")).strip().lower()


@st.cache_data(ttl=60)
def get_pt_data():

    try:
        data = pd.read_csv(SHEET2_CSV_URL)

        data.columns = (
            data.columns
            .astype(str)
            .str.replace("\xa0", " ", regex=False)
            .str.replace("\n", " ", regex=False)
            .str.strip()
        )

        for col in data.columns:

            if data[col].dtype == "object":

                data[col] = (
                    data[col]
                    .astype(str)
                    .str.replace("\xa0", " ", regex=False)
                    .str.strip()
                )

        data = data.replace(
            {
                "nan": "",
                "NaN": "",
                "NAN": ""
            }
        )

        return data

    except Exception as exc:

        st.error(
            f"Unable to load Google Sheet PT: {exc}"
        )

        return pd.DataFrame()


@st.cache_data(ttl=60)
def get_pt_document_links():
    """Read the actual hyperlink target from the PT worksheet."""
    links = {}
    try:
        xlsx_url = (
            f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export"
            f"?format=xlsx"
        )
        request = Request(xlsx_url, headers={"User-Agent": "Mozilla/5.0"})
        with urlopen(request, timeout=30) as response:
            workbook_bytes = response.read()

        workbook = load_workbook(
            filename=BytesIO(workbook_bytes),
            read_only=False,
            data_only=False
        )
        if "PT" not in workbook.sheetnames:
            workbook.close()
            return links

        worksheet = workbook["PT"]
        headers = {}
        for cell in worksheet[1]:
            if cell.value is not None:
                header = str(cell.value).replace("\xa0", " ").strip()
                headers[header] = cell.column

        link_col = headers.get(LINK_COLUMN)
        pt_col = headers.get("PT No.")
        if not link_col or not pt_col:
            workbook.close()
            return links

        formula_pattern = re.compile(
            r'=HYPERLINK\s*\(\s*["\']([^"\']+)["\']',
            re.IGNORECASE
        )

        for row_number in range(2, worksheet.max_row + 1):
            pt_cell = worksheet.cell(row=row_number, column=pt_col)
            link_cell = worksheet.cell(row=row_number, column=link_col)
            pt_number = "" if pt_cell.value is None else str(pt_cell.value).strip()
            if not pt_number:
                continue

            document_link = ""
            if link_cell.hyperlink and link_cell.hyperlink.target:
                document_link = str(link_cell.hyperlink.target).strip()

            if not document_link and isinstance(link_cell.value, str):
                match = formula_pattern.search(link_cell.value)
                if match:
                    document_link = match.group(1).strip()

            if not document_link and isinstance(link_cell.value, str):
                candidate = link_cell.value.strip()
                if re.match(r"^https?://", candidate, re.IGNORECASE):
                    document_link = candidate

            if document_link:
                links[pt_number] = document_link

        workbook.close()
    except Exception:
        return links
    return links


df = get_pt_data()
pt_document_links = get_pt_document_links()



# =========================================================
# ONLY COLUMNS REQUIRED FROM GOOGLE SHEET: PT
# =========================================================

required_columns = [
    "PT No.",
    "Department",
    "Name of PT",
    "Status (Ongoing/Completed)",
    "Attach PT Softcopy Link"
]

# Resolve the real Google Sheet headers by normalized text.
# This handles headers such as "Status  (Ongoing/Completed)"
# or headers containing a line break/extra spaces.
_normalized_columns = {normalize_column_name(col): col for col in df.columns}

_column_aliases = {
    "PT No.": ["PT No."],
    "Department": ["Department"],
    "Name of PT": ["Name of PT"],
    "Status (Ongoing/Completed)": [
        "Status (Ongoing/Completed)",
        "Status  (Ongoing/Completed)"
    ],
    "Attach PT Softcopy Link": ["Attach PT Softcopy Link"]
}

_resolved_columns = {}
for _required in required_columns:
    _resolved = None
    for _alias in _column_aliases.get(_required, [_required]):
        _resolved = _normalized_columns.get(normalize_column_name(_alias))
        if _resolved:
            break
    _resolved_columns[_required] = _resolved

if _resolved_columns["Status (Ongoing/Completed)"]:
    STATUS_COLUMN = _resolved_columns["Status (Ongoing/Completed)"]

if _resolved_columns["Attach PT Softcopy Link"]:
    LINK_COLUMN = _resolved_columns["Attach PT Softcopy Link"]

required_columns = [
    _resolved_columns[col] if _resolved_columns[col] else col
    for col in required_columns
]

# =========================================================
# CHECK DATA
# =========================================================
# =========================================================

if df.empty:

    st.error(
        "No data found in Google Sheet: PT."
    )

    st.stop()


missing_columns = [
    column
    for column in required_columns
    if column not in df.columns
]

if missing_columns:

    st.error(
        "Some required columns are missing from Google Sheet: PT."
    )

    st.write("Missing columns:")
    st.write(missing_columns)

    st.write("Columns found in PT:")
    st.write(df.columns.tolist())

    st.stop()


# =========================================================
# GLOBAL CSS
# LIGHT 3D INDUSTRIAL THEME
# =========================================================

st.markdown(
    """
<style>

/* =====================================================
   REFERENCE-STYLE WHITE / NAVY INDUSTRIAL THEME
   VISUAL ONLY — NO DATA / LOGIC CHANGES
   ===================================================== */

* {
    box-sizing: border-box;
}

html,
body,
.stApp,
[data-testid="stAppViewContainer"],
[data-testid="stAppViewContainer"] > .main {
    margin: 0 !important;
    padding: 0 !important;
    height: 100vh !important;
    max-height: 100vh !important;
    overflow: hidden !important;
}

#MainMenu,
header,
footer,
[data-testid="stHeader"],
[data-testid="stToolbar"] {
    display: none !important;
}

[data-testid="stMainBlockContainer"],
[data-testid="stAppViewBlockContainer"],
.block-container {
    width: 100% !important;
    max-width: none !important;
    margin: 0 !important;
    padding: 0 6px !important;
}

[data-testid="stAppViewContainer"] > .main > div {
    padding: 0 !important;
}

iframe {
    display: block !important;
    border: 0 !important;
    margin: 0 !important;
    padding: 0 !important;
}


/* =====================================================
   MAIN BACKGROUND
   ===================================================== */

.stApp {
    background:
        linear-gradient(
            180deg,
            #ffffff 0%,
            #f7faff 55%,
            #eef4fa 100%
        ) !important;

    color: #092d5c !important;

    font-family:
        Arial,
        Helvetica,
        sans-serif !important;
}

.stApp * {
    font-family:
        Arial,
        Helvetica,
        sans-serif;
}


/* =====================================================
   SPACING
   ===================================================== */

[data-testid="stVerticalBlock"] {
    gap: 0.00rem !important;
}

[data-testid="stHorizontalBlock"] {
    gap: 8px !important;
}


/* =====================================================
   FILTER LABELS
   ===================================================== */

[data-testid="stSelectbox"] label {
    color: #092d5c !important;
    font-size: 12px !important;
    font-weight: 900 !important;
    letter-spacing: .35px !important;
    margin-bottom: 3px !important;
    padding-left: 4px !important;
}


/* =====================================================
   SELECT BOX
   ===================================================== */

div[data-baseweb="select"] > div {
    height: 38px !important;
    min-height: 38px !important;
    border-radius: 6px !important;

    background:
        #ffffff !important;

    border:
        1.5px solid #a9bfd8 !important;

    box-shadow:
        0 2px 5px rgba(8,45,92,.10),
        inset 0 1px 0 rgba(255,255,255,.95) !important;
}

div[data-baseweb="select"]:hover > div {
    border-color: #176fc1 !important;
    box-shadow:
        0 3px 8px rgba(8,76,135,.16) !important;
}

div[data-baseweb="select"] * {
    color: #092d5c !important;
    font-size: 12px !important;
    font-weight: 700 !important;
}

div[data-baseweb="select"] svg {
    fill: #0a4e91 !important;
}


/* =====================================================
   MONTH + DEPARTMENT — ALIGN WITH RESET FILTERS
   ===================================================== */

/*
   IMPORTANT:
   Move only the two selectbox widgets.
   The 22px spacers and the Reset Filters 46px spacer
   remain unchanged, so the Reset Filters position is
   not affected.
*/

/* Month */
div[data-testid="stColumn"]:has(.month-filter-anchor)
div[data-testid="stSelectbox"],
div[data-testid="column"]:has(.month-filter-anchor)
div[data-testid="stSelectbox"] {
    transform: translateY(-3px) !important;
}

/* Department */
div[data-testid="stColumn"]:has(.department-filter-anchor)
div[data-testid="stSelectbox"],
div[data-testid="column"]:has(.department-filter-anchor)
div[data-testid="stSelectbox"] {
    transform: translateY(-3px) !important;
}


/* =====================================================
   TEXT INPUT
   ===================================================== */

div[data-testid="stTextInput"] input {
    height: 40px !important;
    min-height: 40px !important;
    border-radius: 6px !important;

    background:
        #ffffff !important;

    border:
        1.5px solid #a9bfd8 !important;

    color: #092d5c !important;

    font-size: 12px !important;
    font-weight: 600 !important;

    box-shadow:
        0 2px 5px rgba(8,45,92,.09),
        inset 0 1px 2px rgba(0,0,0,.025) !important;
}

div[data-testid="stTextInput"] input:focus {
    border-color: #126bc0 !important;

    box-shadow:
        0 0 0 1px #126bc0,
        0 3px 9px rgba(18,107,192,.15) !important;
}

div[data-testid="stTextInput"] input::placeholder {
    color: #657990 !important;
    opacity: 1 !important;
}


/* =====================================================
   3D INDUSTRIAL BUTTONS
   ===================================================== */

div.stButton > button {
    height: 36px !important;
    min-height: 36px !important;

    border-radius: 6px !important;

    background:
        linear-gradient(
            180deg,
            #ffffff 0%,
            #e8f0f8 100%
        ) !important;

    border:
        1.5px solid #9db7d2 !important;

    color: #07366d !important;

    font-size: 12px !important;
    font-weight: 900 !important;

    box-shadow:
        0 3px 0 #7897b6,
        0 5px 9px rgba(6,48,91,.13),
        inset 0 1px 0 rgba(255,255,255,.95) !important;

    transition:
        transform .12s ease,
        box-shadow .12s ease,
        background .12s ease !important;
}

div.stButton > button:hover {
    border-color: #126bc0 !important;

    color: #ffffff !important;

    background:
        linear-gradient(
            180deg,
            #1685db 0%,
            #075ca8 100%
        ) !important;

    transform:
        translateY(-1px) !important;

    box-shadow:
        0 4px 0 #06477f,
        0 7px 13px rgba(4,74,135,.24),
        inset 0 1px 0 rgba(255,255,255,.28) !important;
}

div.stButton > button:active {
    transform:
        translateY(2px) !important;

    box-shadow:
        0 1px 0 #06477f,
        0 3px 6px rgba(4,74,135,.18) !important;
}

div.stButton > button:disabled {
    color: #8293a7 !important;
    background: #eef3f7 !important;
    border-color: #c5d2df !important;
    box-shadow: none !important;
}


/* =====================================================
   PT REGISTER TOOLBAR — ONLY THESE 4 BUTTONS
   ALL / COMPLETED / ONGOING / REFRESH DATA
   NORMAL = WHITE SHINING
   HOVER = DEEP OCEAN BLUE
   ===================================================== */

/* The toolbar is the horizontal block containing the
   Search input. Columns 2–5 are the four buttons. */

[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(2) button,
[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(3) button,
[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(4) button,
[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(5) button {

    background: #ffffff !important;
    background-image: none !important;

    color:
        #075985 !important;

    border:
        1.5px solid #b8cfe0 !important;

    box-shadow:
        0 2px 4px rgba(0,0,0,.12),
        inset 0 1px 0 #ffffff !important;

    transition:
        background .15s ease,
        color .15s ease,
        border-color .15s ease,
        transform .15s ease,
        box-shadow .15s ease !important;
}


/* Mouse over ONLY the four toolbar buttons */

[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(2) button:hover,
[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(3) button:hover,
[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(4) button:hover,
[data-testid="stHorizontalBlock"]:has(
    [data-testid="stTextInput"]
) > [data-testid="column"]:nth-child(5) button:hover {

    background:
        linear-gradient(
            180deg,
            #0b6f9f 0%,
            #064f73 52%,
            #043d5c 100%
        ) !important;

    color:
        #ffffff !important;

    border-color:
        #064f73 !important;

    transform:
        translateY(-1px) !important;

    box-shadow:
        0 4px 0 #032f46,
        0 8px 15px rgba(4,79,115,.30),
        inset 0 1px 0 rgba(255,255,255,.28) !important;
}


/* =====================================================
   KPI CARDS
   ===================================================== */

.kpi-card {
    position: relative;
    height: 125px;
    overflow: hidden;

    background:
        linear-gradient(
            145deg,
            #ffffff 0%,
            #ffffff 72%,
            #edf4fa 100%
        );

    border:
        1.5px solid #c2d3e4;

    border-top:
        4px solid #176fc1;

    border-radius: 8px;

    padding: 17px 16px;

    box-shadow:
        0 4px 10px rgba(6,48,91,.12),
        0 1px 2px rgba(6,48,91,.08),
        inset 0 1px 0 rgba(255,255,255,.98);

    transition:
        transform .15s ease,
        box-shadow .15s ease;
}

.kpi-card:hover {
    transform: translateY(-2px);

    box-shadow:
        0 7px 16px rgba(6,48,91,.17),
        0 2px 4px rgba(6,48,91,.08),
        inset 0 1px 0 rgba(255,255,255,1);
}

.kpi-card.completed {
    border-top-color: #19a657;
}


.kpi-card.ongoing {
    border-top-color: #f18d05;
}

.kpi-card.approved {
    border-top-color: #176fc1;
}

.kpi-card.pending {
    border-top-color: #d94b4b;
}


/* =====================================================
   KPI ICONS
   ===================================================== */

.kpi-icon {
    display: none !important;
}




/* =====================================================
   TOTAL PT — REMOVE ICON ONLY
   ===================================================== */

.kpi-card.total .kpi-icon {
    display: none;
}

.kpi-card.total .kpi-content {
    margin-left: 0 !important;
    width: 100% !important;
    height: 100% !important;

    display: flex !important;
    flex-direction: column !important;
    align-items: center !important;
    justify-content: center !important;

    text-align: center !important;
}


/* =====================================================
   KPI TEXT — HIGH CONTRAST
   ===================================================== */

/* Keep all KPI text centered horizontally and vertically like TOTAL PT. */
.kpi-card .kpi-content {
    margin-left: 0 !important;
    width: 100% !important;
    height: 100% !important;
    display: flex !important;
    flex-direction: column !important;
    align-items: center !important;
    justify-content: center !important;
    text-align: center !important;
}

.kpi-content {
    margin-left: 78px;
}

.kpi-label {
    color: #092d5c;
    font-size: 14px;
    font-weight: 900;
    letter-spacing: .15px;

    text-align: center;
}

.kpi-card.completed .kpi-label {
    color: #08783c;
}


.kpi-card.ongoing .kpi-label {
    color: #b96700;
}

.kpi-card.approved .kpi-label {
    color: #0a4e91;
}

.kpi-card.pending .kpi-label {
    color: #b33a3a;
}

/* Keep the two document KPI titles on a single line. */
.kpi-card.approved .kpi-label,
.kpi-card.pending .kpi-label {
    white-space: nowrap !important;
    font-size: 13px !important;
    letter-spacing: 0 !important;
}

.kpi-value {
    font-size: 42px;
    line-height: 1;
    font-weight: 900;
    margin-top: 6px;
    color: #0a4e91;

    text-align: center;
}

.kpi-value.green {
    color: #159447 !important;
}


.kpi-value.orange {
    color: #f0a000 !important;
}

.kpi-value.approved {
    color: #176fc1 !important;
}

.kpi-value.pending {
    color: #d94b4b !important;
}

.kpi-description {
    color: #304d6d;
    font-size: 11px;
    font-weight: 700;
    margin-top: 7px;

    text-align: center;
}

.kpi-pattern {
    display: none !important;
}



.kpi-arrow {
    display: none !important;
}


/* =====================================================
   PT REGISTER PANEL
   ===================================================== */

.register-wrap {
    background: #ffffff;

    border:
        1.5px solid #b7cce1;

    border-radius:
        7px 7px 0 0;

    overflow: hidden;

    box-shadow:
        0 4px 10px rgba(7,45,82,.12);
}

.register-title {
    height: 40px;

    display: flex;
    align-items: center;

    padding: 0 17px;

    color: #ffffff;

    font-size: 18px;
    font-weight: 900;
    letter-spacing: .25px;

    background:
        linear-gradient(
            180deg,
            #0a4f91 0%,
            #063b70 100%
        );

    border-bottom:
        2px solid #176fc1;

    box-shadow:
        inset 0 1px 0 rgba(255,255,255,.16);
}

.register-icon {
    margin-right: 9px;
    color: #ffffff;
}


/* =====================================================
   TABLE HEADER — REFERENCE MATCH
   ===================================================== */

.table-head {
    min-height: 43px;

    display: flex;
    align-items: center;
    justify-content: center;

    background:
        linear-gradient(
            180deg,
            #0b4f91 0%,
            #063c73 100%
        );

    color: #ffffff;

    border-right:
        1px solid #8caecc;

    border-top:
        1px solid #2879ba;

    border-bottom:
        1px solid #052f5b;

    font-size: 12px;
    line-height: 1.15;
    font-weight: 900;

    text-align: center;
    padding: 5px 3px;

    box-shadow:
        inset 0 1px 0 rgba(255,255,255,.16);
}


/* =====================================================
   TABLE CELLS — DARK BLUE CLEAR TEXT
   ===================================================== */

.table-cell {
    min-height: 43px;

    display: flex;
    align-items: center;
    justify-content: center;

    background:
        #ffffff;

    border-right:
        1px solid #c8d6e4;

    border-bottom:
        1px solid #c8d6e4;

    color:
        #092d5c;

    font-size:
        11px;

    line-height:
        1.18;

    font-weight:
        600;

    text-align:
        center;

    padding:
        5px 4px;

    word-break:
        break-word;
}

.table-cell.alt {
    background:
        #f3f7fb;
}

.table-cell.left {
    justify-content:
        flex-start;

    text-align:
        left;

    font-weight:
        650;
}
/* =====================================================
   STATUS — TEXT ONLY
   ===================================================== */

.status-pill {
    display: inline-flex;
    align-items: center;
    justify-content: center;

    min-width: auto;
    padding: 0;

    border-radius: 0;

    background: transparent !important;
    border: none !important;
    box-shadow: none !important;

    font-size: 11px;
    font-weight: 900;
    letter-spacing: .1px;

    white-space: nowrap;
}


/* COMPLETED — GREEN TEXT ONLY */

.status-completed {
    background: transparent !important;
    border: none !important;
    color: #16A34A !important;
    box-shadow: none !important;
}


/* ONGOING — ORANGE TEXT ONLY */

.status-ongoing {
    background: transparent !important;
    border: none !important;
    color: #EA8A00 !important;
    box-shadow: none !important;
}
/* =====================================================
   STREAMLIT TABLE ACTION BUTTONS
   ===================================================== */

.table-cell + div button,
div[data-testid="column"] div.stButton > button {
    font-size: 11px !important;
    font-weight: 900 !important;
}
/* =====================================================
   RECORD BAR / PAGINATION
   ===================================================== */

.record-bar {
    height: 38px;

    display: flex;
    align-items: center;

    padding: 0 12px;

    background:
        linear-gradient(
            180deg,
            #ffffff 0%,
            #edf3f8 100%
        );

    color:
        #173f6d;

    font-size:
        11px;

    font-weight:
        800;

    border-top:
        1px solid #c4d4e3;

    border-bottom:
        1px solid #c4d4e3;

    box-shadow:
        inset 0 1px 0 rgba(255,255,255,.9);
}


/* =====================================================
   DOWNLOAD BUTTON
   ===================================================== */

div.stDownloadButton > button {
    border-radius: 6px !important;

    background:
        linear-gradient(
            180deg,
            #ffffff,
            #e9f1f8
        ) !important;

    border:
        1.5px solid #9eb8d2 !important;

    color:
        #083c76 !important;

    font-size:
        11px !important;

    font-weight:
        900 !important;

    box-shadow:
        0 3px 0 #7895b1,
        0 5px 8px rgba(8,53,94,.12) !important;
}

div.stDownloadButton > button:hover {
    color: #ffffff !important;

    background:
        linear-gradient(
            180deg,
            #1685db,
            #075ca8
        ) !important;

    border-color:
        #075ca8 !important;
}


/* =====================================================
   INFO / ALERT
   ===================================================== */

div[data-testid="stAlert"] {
    border-radius: 6px !important;

    color: #123b68 !important;

    box-shadow:
        0 2px 7px rgba(20,70,100,.08) !important;
}


/* =====================================================
   FOOTER
   ===================================================== */

.footer {
    height: 34px;

    display: flex;
    align-items: center;
    justify-content: center;

    color: #ffffff;

    background:
        linear-gradient(
            180deg,
            #0a4f91 0%,
            #063563 100%
        );

    font-size:
        11px;

    font-weight:
        800;

    border-top:
        2px solid #176fc1;

    box-shadow:
        0 -2px 8px rgba(0,0,0,.12);
}


/* =====================================================
   SCROLLBAR
   ===================================================== */

::-webkit-scrollbar {
    width: 8px;
    height: 8px;
}

::-webkit-scrollbar-track {
    background: #e9f0f6;
}

::-webkit-scrollbar-thumb {
    background: #8daac4;
    border-radius: 8px;
}

::-webkit-scrollbar-thumb:hover {
    background: #527fa6;
}

</style>
""",
    unsafe_allow_html=True
)

# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="PSM Digital Dashboard",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="collapsed"
)


# ============================================================
# REMOVE STREAMLIT TOP SPACE
# ============================================================

st.markdown(
    """
    <style>

    html,
    body {
        margin: 0 !important;
        padding: 0 !important;
    }

    [data-testid="stAppViewContainer"] {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }

    [data-testid="stAppViewContainer"] > .main {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }


    [data-testid="stDecoration"] {
        display: none !important;
        height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    .block-container {
        padding-top: 0 !important;
        margin-top: 0 !important;
        padding-bottom: 0 !important;
        padding-left: 0 !important;
        padding-right: 0 !important;
        max-width: 100% !important;
    }

    .stApp {
        margin-top: 0 !important;
        padding-top: 0 !important;
    }

    iframe {
        display: block !important;
        margin-top: -0px !important;
        padding-top: 0 !important;
        border: 0 !important;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# BASE DIRECTORY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent


# ============================================================
# IMAGE TO BASE64
# ============================================================

def image_to_base64(file_path):

    file_path = Path(file_path)

    if not file_path.exists():
        return ""

    try:

        with open(file_path, "rb") as file:

            return base64.b64encode(
                file.read()
            ).decode("utf-8")

    except Exception:

        return ""


# ============================================================
# COMPANY LOGO
# ============================================================

logo_path = BASE_DIR / "jsw_jfe_logo.jpg"

logo_base64 = image_to_base64(logo_path)


# ============================================================
# FILE CHECK
# ============================================================

if not logo_base64:

    st.error(
        "jsw_jfe_logo.jpg not found. "
        "Keep jsw_jfe_logo.jpg in the same folder as this Python file."
    )


# ============================================================
# DATE AND TIME
# ============================================================

now = datetime.now(ZoneInfo("Asia/Kolkata"))

current_date = now.strftime(
    "%d %b %Y"
).upper()

current_time = now.strftime(
    "%I:%M %p"
)


# ============================================================
# HEADER HTML
# ============================================================

header_html = """

<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<style>


/* ============================================================
   MAIN HEADER
   ============================================================ */

.psm-header {

    position: relative;

    width: 100%;

    height: 90px;

    overflow: hidden;

    background:
        linear-gradient(
            90deg,
            #031d34 0%,
            #052b49 42%,
            #07385c 74%,
            #052b49 100%
        );

    border-radius: 7px;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.18);

}


/* ============================================================
   LEFT LOGO AREA
   ============================================================ */

.psm-left {

    position: absolute;

    left: 0;

    top: 0;

    width: 100%;

    height: 95px;

    display: flex;

    align-items: center;

    padding-left: 4px;

    box-sizing: border-box;

    z-index: 20;

    pointer-events: none;

}


/* ============================================================
   LOGO PANEL
   ONLY VERTICAL POSITION CHANGED
   ============================================================ */

.logo-panel {

    width: 195px;

    height: 68px;

    background: #ffffff;

    border-radius: 5px;

    padding: 4px;

    display: flex;

    align-items: center;

    justify-content: center;

    flex-shrink: 0;

    box-sizing: border-box;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.20);

    position: relative;

    top: -5px;

    left: 5px;

}


/* ============================================================
   COMPANY LOGO
   ============================================================ */

.company-logo {

    width: 100%;

    height: 100%;

    object-fit: contain;

    object-position: center;

    display: block;

}


/* ============================================================
   LEFT VERTICAL DIVIDER
   ============================================================ */

.vertical-line {

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

    margin-left: 18px;

    margin-right: 20px;

    flex-shrink: 0;

}


/* ============================================================
   CENTER TITLE AREA
   ============================================================ */

.title-area {

    position: absolute;

    left: 50%;

    top: 0;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: center;

    text-align: center;

    min-width: max-content;

    box-sizing: border-box;

    transform: translateX(-50%);

}


/* ============================================================
   MAIN TITLE
   ============================================================ */

.main-title {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 27px;

    font-weight: 900;

    line-height: 1;

    letter-spacing: 0.3px;

    white-space: nowrap;

    margin: 0;

    padding: 0;

}


/* ============================================================
   ORANGE TITLE PART
   ============================================================ */

.main-title-orange {

    color: #f28c00;

}


/* ============================================================
   SUBTITLE
   ============================================================ */

.subtitle {

    color: #ffffff;

    font-family:
        Arial,
        sans-serif;

    font-size: 12px;

    font-weight: 400;

    letter-spacing: 3.6px;

    margin-top: 8px;

    line-height: 1;

    white-space: nowrap;

}

/* ============================================================
   SUB-SUBTITLE / TAGLINE
   ============================================================ */

.tagline {

    color:
        rgba(255,255,255,0.82);

    font-family:
        Arial,
        sans-serif;

    font-size: 7px;

    font-weight: 500;

    letter-spacing: 2.2px;

    margin-top: 6px;

    line-height: 1;

    white-space: nowrap;

}


/* ============================================================
   RIGHT DATE / TIME AREA
   ============================================================ */

.psm-right {

    position: absolute;

    right: 16px;

    top: 0;

    width: 15%;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: flex-end;

    text-align: right;

    color: #ffffff;

    z-index: 30;

    padding-left: 18px;

    box-sizing: border-box;

}


/* ============================================================
   RIGHT VERTICAL DIVIDER
   ============================================================ */

.psm-right::before {

    content: "";

    position: absolute;

    left: 15px;

    top: 6px;

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

}


/* ============================================================
   DATE
   ============================================================ */

.date {

    color:
        rgba(255,255,255,0.95);

    font-family:
        Arial,
        sans-serif;

    font-size: 11px;

    font-weight: 400;

    letter-spacing: 0.7px;

    line-height: 1;

    margin: 0;

    padding: 0;

}


/* ============================================================
   TIME
   ============================================================ */

.time {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 22px;

    font-weight: 800;

    margin-top: 4px;

    line-height: 1;

    padding: 0;

}


/* ============================================================
   RIGHT HORIZONTAL LINE
   ============================================================ */

.right-line {

    width: 80px;

    height: 2px;

    background:
        rgba(255,255,255,0.75);

    margin-top: 7px;

    flex-shrink: 0;

}


/* ============================================================
   ORANGE BOTTOM BAR
   ============================================================ */

.orange-bar {

    position: absolute;

    left: 0;

    bottom: 0;

    width: 100%;

    height: 7px;

    background: #f28c00;

    z-index: 50;

}


</style>

</head>


<body>


<!-- ============================================================
     MAIN HEADER
     ============================================================ -->

<div class="psm-header">


    <!-- ========================================================
         LEFT LOGO AREA
         ======================================================== -->

    <div class="psm-left">


        <!-- ====================================================
             LOGO
             ==================================================== -->

        <div class="logo-panel">

            <img
                class="company-logo"
                src="data:image/jpeg;base64,LOGO_IMAGE_BASE64"
                alt="JSW JFE Steel Limited"
            >

        </div>


        <!-- ====================================================
             LEFT VERTICAL LINE
             ==================================================== -->

        <div class="vertical-line"></div>


        <!-- ====================================================
             CENTER TITLE GROUP
             ==================================================== -->

        <div class="title-area">


            <!-- MAIN TITLE -->

            <div class="main-title">

                PROCESS TECHNOLOGY (PT)

                <span class="main-title-orange"></span>

            </div>


            <!-- SUBTITLE -->

            <div class="subtitle">

                PSM DIGITAL DASHBOARD

            </div>


            <!-- SUB-SUBTITLE -->

            <div class="tagline">

                PEOPLE
                &nbsp; | &nbsp;
                PROCESS
                &nbsp; | &nbsp;
                RISK
                &nbsp; | &nbsp;
                COMPLIANCE

            </div>


        </div>


    </div>


    <!-- ========================================================
         RIGHT DATE / TIME
         ======================================================== -->

    <div class="psm-right">


        <div class="date">

            CURRENT_DATE_VALUE

        </div>


        <div class="time">

            CURRENT_TIME_VALUE

        </div>


        <div class="right-line"></div>


    </div>


    <!-- ========================================================
         ORANGE BOTTOM BAR
         ======================================================== -->

    <div class="orange-bar"></div>


</div>


</body>

</html>

"""


# ============================================================
# INSERT LOGO
# ============================================================

header_html = header_html.replace(
    "LOGO_IMAGE_BASE64",
    logo_base64
)


# ============================================================
# INSERT DATE
# ============================================================

header_html = header_html.replace(
    "CURRENT_DATE_VALUE",
    current_date
)


# ============================================================
# INSERT TIME
# ============================================================

header_html = header_html.replace(
    "CURRENT_TIME_VALUE",
    current_time
)


# ============================================================
# DISPLAY HEADER
# ============================================================

components.html(
    header_html,
    height=114,
    scrolling=False
)
#======================================================================================================================
# =========================================================
# RESET FILTER
# =========================================================

def reset_pt_filters():
    st.session_state.status_filter = "All"
    st.session_state.page_number = 1
    st.session_state.department_selector = "All Departments"


# =========================================================
# FILTER SECTION
# =========================================================

filter_month, filter_department, filter_reset = st.columns(
    [1.0, 1.0, 0.34],
    gap="small"
)

with filter_month:
    st.markdown(
        "<div class='month-filter-anchor' style='height:22px;'></div>",
        unsafe_allow_html=True
    )

    selected_month = st.selectbox(
        "Month",
        [
            "August 2026",
            "July 2026",
            "June 2026",
            "May 2026",
            "April 2026",
            "March 2026"
        ],
        index=0
    )

with filter_department:
    st.markdown(
        "<div class='department-filter-anchor' style='height:22px;'></div>",
        unsafe_allow_html=True
    )

    department_values = (
        df["Department"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    department_options = ["All Departments"] + sorted(
        [
            value for value in department_values.unique().tolist()
            if value and value.lower() != "nan"
        ],
        key=lambda x: x.lower()
    )

    selected_department = st.selectbox(
        "Department",
        department_options,
        key="department_selector"
    )

with filter_reset:
    st.markdown(
        "<div style='height:46px;'></div>",
        unsafe_allow_html=True
    )

    st.button(
        "↻ Reset Filters",
        use_container_width=True,
        key="reset_pt_filters_button",
        on_click=reset_pt_filters
    )


# =========================================================
# FILTER DATA
# =========================================================

filtered_df = df.copy()

if selected_department != "All Departments":
    filtered_df = filtered_df[
        filtered_df["Department"]
        .fillna("")
        .astype(str)
        .str.strip()
        == selected_department
    ]

filtered_df[STATUS_COLUMN] = (
    filtered_df[STATUS_COLUMN]
    .fillna("")
    .astype(str)
    .str.strip()
    .str.lower()
)


# =========================================================
# KPI
# =========================================================

total_pt = len(filtered_df)

completed = int(
    (filtered_df[STATUS_COLUMN] == "completed").sum()
)

ongoing = int(
    (filtered_df[STATUS_COLUMN] == "ongoing").sum()
)

# =========================================================
# APPROVAL KPI
# Source column: Approved  (Yes/No)
# This column is NOT a required column.
# It is used only for the two approval KPI boxes.
# =========================================================

def _find_approval_column(dataframe):
    target = normalize_column_name("Approved  (Yes/No)")
    for column in dataframe.columns:
        if normalize_column_name(column) == target:
            return column
    return None


_APPROVAL_SOURCE_COLUMN = _find_approval_column(df)

if _APPROVAL_SOURCE_COLUMN is not None:
    approval_values = (
        df[_APPROVAL_SOURCE_COLUMN]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.lower()
    )

    approved_documents = int(
        (approval_values == "yes").sum()
    )

    pending_documents = int(
        (approval_values == "no").sum()
    )
else:
    approved_documents = 0
    pending_documents = 0

completion_percentage = (
    completed / total_pt * 100
    if total_pt
    else 0
)

# Five KPI boxes are displayed in one row.
k1, k2, k3, k4, k5 = st.columns(5, gap="small")

cards = [
    ("", "TOTAL PT", total_pt, "", "blue", "total"),
    (
        "",
        "APPROVED",
        approved_documents,
        "",
        "approved",
        "approved"
    ),
    (
        "",
        "PENDING",
        pending_documents,
        "",
        "pending",
        "pending"
    ),
    (
        "",
        "COMPLETED",
        completed,
        f"",
        "green",
        "completed"
    ),
    (
        "",
        "ONGOING",
        ongoing,
        "",
        "orange",
        "ongoing"
    )
]

for column, card in zip([k1, k2, k3, k4, k5], cards):
    icon, label, value, description, color, extra = card

    with column:
        st.html(
            f"""
<div class="kpi-card {extra}">
    <div class="kpi-icon">{icon}</div>
    <div class="kpi-content">
        <div class="kpi-label">{label}</div>
        <div class="kpi-value {color}">{value}</div>
        <div class="kpi-description">{description}</div>
    </div>
    <div class="kpi-pattern"></div>
    <div class="kpi-arrow">&gt;</div>
</div>
"""
        )


# =========================================================
# PT REGISTER
# =========================================================

st.html(
    """
<div class="register-wrap">
    <div class="register-title">
        <span class="register-icon">▣</span>
        PT REGISTER
    </div>
</div>
"""
)


# =========================================================
# TOOLBAR
# =========================================================

search_col, all_col, completed_col, ongoing_col, refresh_col = st.columns(
    [2.8, 0.55, 0.85, 0.75, 1.05],
    gap="small"
)

with search_col:
    search_text = st.text_input(
        "Search",
        placeholder="Search PT No., Name of PT, Department...",
        label_visibility="collapsed",
        key="pt_search"
    )

with all_col:
    if st.button("All", use_container_width=True, key="pt_all"):
        st.session_state.status_filter = "All"
        st.session_state.page_number = 1
        st.rerun()

with completed_col:
    if st.button("Completed", use_container_width=True, key="pt_completed"):
        st.session_state.status_filter = "Completed"
        st.session_state.page_number = 1
        st.rerun()

with ongoing_col:
    if st.button("Ongoing", use_container_width=True, key="pt_ongoing"):
        st.session_state.status_filter = "Ongoing"
        st.session_state.page_number = 1
        st.rerun()

with refresh_col:
    if st.button("↻ Refresh Data", use_container_width=True, key="pt_refresh"):
        st.cache_data.clear()
        st.rerun()


# =========================================================
# SEARCH + STATUS
# =========================================================

display_df = filtered_df.copy()

if search_text.strip():
    q = search_text.strip().lower()

    search_mask = (
        display_df["PT No."]
        .fillna("")
        .astype(str)
        .str.lower()
        .str.contains(q, regex=False)
        |
        display_df["Department"]
        .fillna("")
        .astype(str)
        .str.lower()
        .str.contains(q, regex=False)
        |
        display_df["Name of PT"]
        .fillna("")
        .astype(str)
        .str.lower()
        .str.contains(q, regex=False)
        |
        display_df[STATUS_COLUMN]
        .fillna("")
        .astype(str)
        .str.lower()
        .str.contains(q, regex=False)
    )

    display_df = display_df[search_mask]

if st.session_state.status_filter != "All":
    display_df = display_df[
        display_df[STATUS_COLUMN]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.lower()
        ==
        st.session_state.status_filter.lower()
    ]


# =========================================================
# PAGINATION
# =========================================================

ROWS_PER_PAGE = 5
total_entries = len(display_df)

total_pages = max(
    1,
    (total_entries + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE
)

if st.session_state.page_number > total_pages:
    st.session_state.page_number = total_pages

page_number = st.session_state.page_number

start_index = (page_number - 1) * ROWS_PER_PAGE
end_index = start_index + ROWS_PER_PAGE

page_df = display_df.iloc[start_index:end_index].copy()


# =========================================================
# TABLE CSS
# =========================================================

st.markdown(
    """
<style>
.pt-simple-table {
    width: 100%;
    overflow-x: auto;
    border: 1px solid #c8d6e4;
    background: #ffffff;
}

.pt-simple-inner {
    min-width: 900px;
}

.pt-simple-header,
.pt-simple-row {
    display: grid;
    grid-template-columns:
        0.60fr
        1.10fr
        1.45fr
        2.20fr
        1.20fr
        1.30fr;
}

.pt-simple-header {
    min-height: 44px;
    background: linear-gradient(
        180deg,
        #0b4f91 0%,
        #063c73 100%
    );
}

.pt-simple-header > div {
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 5px 4px;
    color: #ffffff;
    border-right: 1px solid #8caecc;
    border-bottom: 1px solid #052f5b;
    font-size: 12px;
    font-weight: 900;
    line-height: 1.15;
    text-align: center;
}

.pt-simple-cell {
    min-height: 43px;
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 5px 7px;
    background: #ffffff;
    color: #092d5c;
    border-right: 1px solid #c8d6e4;
    border-bottom: 1px solid #c8d6e4;
    font-size: 14px;
    font-weight: 600;
    line-height: 1.18;
    text-align: center;
    word-break: break-word;
}

.pt-simple-cell.alt {
    background: #f3f7fb;
}

.pt-simple-cell.left {
    justify-content: flex-start;
    text-align: left;
}

.pt-simple-status-completed {
    color: #16A34A;
    font-weight: 900;
}

.pt-simple-status-ongoing {
    color: #EA8A00;
    font-weight: 900;
}

.pt-view-link {
    width: 100%;
    min-height: 43px;
    display: flex;
    align-items: center;
    justify-content: center;
    background: #ffffff;
    color: #174b87 !important;
    text-decoration: none !important;
    font-size: 11px;
    font-weight: 900;
    border: 0;
    padding: 0;
    margin: 0;
    cursor: pointer;
    font-family: Arial, Helvetica, sans-serif;
}

.pt-view-link.alt {
    background: #f3f7fb;
}

.pt-view-link:hover {
    background: #e8f1f9;
    color: #075ca8 !important;
}

.pt-no-link {
    width: 100%;
    min-height: 43px;
    display: flex;
    align-items: center;
    justify-content: center;
    background: #ffffff;
    color: #9aaabd;
    font-size: 11px;
    font-weight: 700;
}

.pt-no-link.alt {
    background: #f3f7fb;
}

.pt-record-bar {
    min-height: 38px;
    display: flex;
    align-items: center;
    padding: 0 12px;
    background: linear-gradient(
        180deg,
        #ffffff 0%,
        #edf3f8 100%
    );
    color: #173f6d;
    font-size: 11px;
    font-weight: 800;
    border-top: 1px solid #c4d4e3;
    border-bottom: 1px solid #c4d4e3;
}

.pt-pagination {
    height: 38px !important;
    min-height: 38px !important;
    display: flex;
    align-items: center;
    justify-content: center;
    color: #173f6d;
    font-size: 11px;
    font-weight: 800;
    margin: 0 !important;
    padding: 0 !important;
}

.pt-pagination-row [data-testid="stHorizontalBlock"] {
    align-items: center !important;
}

.pt-pagination-row [data-testid="stVerticalBlock"] {
    justify-content: center !important;
    gap: 0 !important;
}

.pt-pagination-row div.stButton {
    margin: 0 !important;
    padding: 0 !important;
}

.pt-pagination-row div.stButton > button {
    height: 36px !important;
    min-height: 36px !important;
    margin: 0 !important;
}
</style>
""",
    unsafe_allow_html=True
)


# =========================================================
# TABLE HEADER
# =========================================================

st.html(
    """
<div class="pt-simple-table">
<div class="pt-simple-inner">
<div class="pt-simple-header">
    <div>Sr No.</div>
    <div>PT No.</div>
    <div>Department</div>
    <div>Name of PT</div>
    <div>Status</div>
    <div>View Document</div>
</div>
</div>
</div>
"""
)


# =========================================================
# TABLE DATA
#
# The URL is fetched from the SAME GOOGLE SHEET ROW:
# Attach PT Softcopy Link
#
# That source column is hidden from the visible table.
# =========================================================

rows_html = [
    '<div class="pt-simple-table">',
    '<div class="pt-simple-inner">'
]

for row_no, (_, row) in enumerate(page_df.iterrows()):

    alt = "alt" if row_no % 2 else ""

    sr_no = start_index + row_no + 1

    pt_no = str(row["PT No."]).strip()
    department = str(row["Department"]).strip()
    name_pt = str(row["Name of PT"]).strip()
    status = str(row[STATUS_COLUMN]).strip()

    # =====================================================
    # FETCH DOCUMENT LINK FROM GOOGLE SHEET ROW
    # =====================================================

    # Fetch the actual URL from the same Google Sheet row.
    document_link = pt_document_links.get(pt_no, "")

    # Fallback when the sheet cell itself contains a plain URL.
    if not document_link:
        raw_link = row[LINK_COLUMN]
        if not pd.isna(raw_link):
            candidate_link = str(raw_link).strip()
            if re.match(r"^https?://", candidate_link, re.IGNORECASE):
                document_link = candidate_link

    # HTML escaping for displayed values.
    def esc(value):
        return (
            str(value)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
            .replace("'", "&#39;")
        )

    if status.lower() == "completed":
        status_html = (
            '<span class="pt-simple-status-completed">'
            'COMPLETED'
            '</span>'
        )

    elif status.lower() == "ongoing":
        status_html = (
            '<span class="pt-simple-status-ongoing">'
            'ONGOING'
            '</span>'
        )

    else:
        status_html = esc(status) if status else "—"

    if document_link:

        safe_link = (
            document_link
            .replace("&", "&amp;")
            .replace('"', "&quot;")
            .replace("<", "%3C")
            .replace(">", "%3E")
        )

        # A normal anchor is used deliberately. Streamlit does not intercept
        # this click, so the browser opens the exact document URL in a new tab.
        view_html = f"""
<a
    class=\"pt-view-link {alt}\"
    href=\"{safe_link}\"
    target=\"_blank\"
    rel=\"noopener noreferrer\"
>
    ◉ View
</a>
"""

    else:

        view_html = f"""
<div class=\"pt-no-link {alt}\">
    No Link
</div>
"""

    rows_html.append(
        f"""
<div class="pt-simple-row">

    <div class="pt-simple-cell {alt}">
        {esc(sr_no)}
    </div>

    <div class="pt-simple-cell {alt}">
        {esc(pt_no)}
    </div>

    <div class="pt-simple-cell {alt} left">
        {esc(department)}
    </div>

    <div class="pt-simple-cell {alt} left">
        {esc(name_pt)}
    </div>

    <div class="pt-simple-cell {alt}">
        {status_html}
    </div>

    <div class="pt-simple-cell {alt}">
        {view_html}
    </div>

</div>
"""
    )

rows_html.append("</div></div>")

st.html(
    "".join(rows_html)
)


# =========================================================
# RECORD COUNT + PAGINATION
# =========================================================

first_entry = start_index + 1 if total_entries else 0
last_entry = min(end_index, total_entries)

if total_pages > 1:

    # Keep Showing text, Previous, Page X of Y and Next
    # on the same horizontal line and vertically centered.
    st.markdown(
        '<div class="pt-pagination-row">',
        unsafe_allow_html=True
    )

    pg_showing, pg_prev, pg_page, pg_next, pg_end = st.columns(
        [2.25, 1.05, 1.45, 1.05, 2.20],
        gap="small"
    )

    with pg_showing:
        st.markdown(
            f"""
<div class="pt-record-bar">
    Showing {first_entry} - {last_entry}
    of {total_entries} entries
</div>
""",
            unsafe_allow_html=True
        )

    with pg_prev:
        if st.button(
            "‹ Previous",
            use_container_width=True,
            key="pt_previous_page",
            disabled=page_number <= 1
        ):
            st.session_state.page_number -= 1
            st.rerun()

    with pg_page:
        st.markdown(
            f"""
<div class="pt-pagination">
    Page {page_number} of {total_pages}
</div>
""",
            unsafe_allow_html=True
        )

    with pg_next:
        if st.button(
            "Next ›",
            use_container_width=True,
            key="pt_next_page",
            disabled=page_number >= total_pages
        ):
            st.session_state.page_number += 1
            st.rerun()

    st.markdown(
        '</div>',
        unsafe_allow_html=True
    )

else:

    st.markdown(
        f"""
<div class="pt-record-bar">
    Showing {first_entry} - {last_entry}
    of {total_entries} entries
</div>
""",
        unsafe_allow_html=True
    )

# =========================================================
# FOOTER
# =========================================================

st.markdown(
    """
<div class="footer">
    PROCESS SAFETY MANAGEMENT • DIGITAL OPERATIONS
</div>
""",
    unsafe_allow_html=True
)
