import streamlit as st
import html
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Process Safety Incident Dashboard",
    layout="wide",
    initial_sidebar_state="collapsed"
)
# ============================================================
# STYLE
# ============================================================
st.markdown(
    """
    <style>

    .stApp {
        background:#f3f8fc;
    }

    #MainMenu, footer {
        visibility:hidden;
    }



.block-container {
    padding:0rem 0.35rem 0rem 0.35rem !important;
    margin-top:-35px !important;
    margin-bottom:0px !important;
    max-width:100%;
}

/* REMOVE BOTTOM SPACE */
[data-testid="stAppViewContainer"] {
    padding-bottom:0px !important;
}

[data-testid="stMainBlockContainer"] {
    padding-bottom:0px !important;
    margin-bottom:0px !important;
}

section.main {
    padding-bottom:0px !important;
    margin-bottom:0px !important;
}

   div[data-testid="stMetric"] {
    background:#ffffff;
    border:1px solid #cbddea;
    border-radius:7px;
    padding:6px 6px !important;
    min-height:72px;
    overflow:visible !important;
}

    div[data-testid="stMetricLabel"] {
    font-size:9px !important;
    font-weight:800 !important;
    color:#20384f !important;
    white-space:nowrap !important;
    overflow:visible !important;
    text-overflow:clip !important;
    line-height:1.1 !important;
}
div[data-testid="stMetricLabel"],
div[data-testid="stMetricLabel"] > div,
div[data-testid="stMetricLabel"] p {
    overflow:visible !important;
    text-overflow:clip !important;
    white-space:nowrap !important;
    max-width:none !important;
}

div[data-testid="stMetricLabel"] p {
    margin:0 !important;
    padding:0 !important;
    font-size:8px !important;
    line-height:1.1 !important;
}
    div[data-testid="stMetricValue"] {
        color:#123f77 !important;
        font-size:24px !important;
        font-weight:900 !important;
    }

    .module-card {
        background:#ffffff;
        border:1px solid #d3e0ea;
        border-radius:7px;
        padding:7px;
        margin-bottom:8px;
        box-shadow:0 1px 4px rgba(20,65,95,.06);
    }

    .module-title {
        color:#073f78;
        font-size:12px;
        font-weight:900;
        margin-bottom:6px;
    }

    .section-bar {
        background:#07518b;
        color:#ffffff;
        border-radius:4px;
        padding:6px 8px;
        font-size:10px;
        font-weight:900;
        margin:4px 0 6px 0;
    }

    .live-bar {
        background:#ffffff;
        border:1px solid #cbddea;
        border-radius:4px;
        padding:5px 8px;
        color:#4f6678;
        font-size:10px;
        margin-bottom:6px;
    }

    .footer {
        text-align:center;
        color:#627689;
        background:#edf4f8;
        border-top:1px solid #cbdce7;
        padding:7px;
        font-size:10px;
        font-weight:800;
        margin-top:8px;
    }

    .small-note {
        font-size:9px;
        color:#6c7f8f;
    }

    .stDataFrame {
        border:1px solid #d5e0e8;
    }


/* ========================================================
   REFRESH BUTTON — KEEP BELOW HEADER
   ======================================================== */

div[data-testid="stButton"] {
    margin-top: 1px !important;
    margin-bottom: 1px !important;
}



    /* ========================================================
       MATCH CLICKABLE MODULE HEADINGS WITH NORMAL HEADINGS
       ======================================================== */
    [data-testid="stPageLink"] {
        margin-bottom:6px !important;
    }

    [data-testid="stPageLink"],
    [data-testid="stPageLink"] a,
    [data-testid="stPageLink"] a *,
    [data-testid="stPageLink"] p,
    [data-testid="stPageLink"] span,
    [data-testid="stPageLink"] div {
        color:#073f78 !important;
        font-size:12px !important;
        font-weight:900 !important;
        text-decoration:none !important;
    }

    [data-testid="stPageLink"] a:hover,
    [data-testid="stPageLink"] a:hover * {
        color:#073f78 !important;
        text-decoration:none !important;
    }

    </style>
    """,
    unsafe_allow_html=True,
)

#=============================================================
#HEADER CODE
#=============================================================
import streamlit as st
import streamlit.components.v1 as components
import base64
from pathlib import Path
from datetime import datetime


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="PSM Digital Dashboard",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="collapsed"
)


# ============================================================
# REMOVE STREAMLIT TOP SPACE
# ============================================================

st.markdown(
    """
    <style>

    html,
    body {
        margin: 0 !important;
        padding: 0 !important;
    }

    [data-testid="stAppViewContainer"] {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }

    [data-testid="stAppViewContainer"] > .main {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }



    [data-testid="stDecoration"] {
        display: none !important;
        height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    .block-container {
    padding-top: 0 !important;
    margin-top: -30px !important;
    padding-bottom: 0 !important;
    padding-left: 0 !important;
    padding-right: 0 !important;
    max-width: 100% !important;
}

    .stApp {
        margin-top: 0 !important;
        padding-top: 0 !important;
    }

    iframe {
        display: block !important;
        margin-top: -20px !important;
        padding-top: 0 !important;
        border: 0 !important;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# BASE DIRECTORY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent


# ============================================================
# IMAGE TO BASE64
# ============================================================

def image_to_base64(file_path):

    file_path = Path(file_path)

    if not file_path.exists():
        return ""

    try:

        with open(file_path, "rb") as file:

            return base64.b64encode(
                file.read()
            ).decode("utf-8")

    except Exception:

        return ""


# ============================================================
# COMPANY LOGO
# ============================================================

logo_path = BASE_DIR / "jsw_jfe_logo.jpg"

logo_base64 = image_to_base64(logo_path)


# ============================================================
# FILE CHECK
# ============================================================

if not logo_base64:

    st.error(
        "jsw_jfe_logo.jpg not found. "
        "Keep jsw_jfe_logo.jpg in the same folder as this Python file."
    )


# ============================================================
# DATE AND TIME
# ============================================================

now = datetime.now()

current_date = now.strftime(
    "%d %b %Y"
).upper()

current_time = now.strftime(
    "%I:%M %p"
)


# ============================================================
# HEADER HTML
# ============================================================

header_html = """

<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<style>


/* ============================================================
   MAIN HEADER
   ============================================================ */

.psm-header {

    position: relative;

    width: 100%;

    height: 90px;

    overflow: hidden;

    background:
        linear-gradient(
            90deg,
            #031d34 0%,
            #052b49 42%,
            #07385c 74%,
            #052b49 100%
        );

    border-radius: 7px;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.18);

}


/* ============================================================
   LEFT LOGO AREA
   ============================================================ */

.psm-left {

    position: absolute;

    left: 0;

    top: 0;

    width: 100%;

    height: 95px;

    display: flex;

    align-items: center;

    padding-left: 4px;

    box-sizing: border-box;

    z-index: 20;

    pointer-events: none;

}


/* ============================================================
   LOGO PANEL
   ONLY VERTICAL POSITION CHANGED
   ============================================================ */

.logo-panel {

    width: 195px;

    height: 68px;

    background: #ffffff;

    border-radius: 5px;

    padding: 10px;

    display: flex;

    align-items: center;

    justify-content: center;

    flex-shrink: 0;

    box-sizing: border-box;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.20);

    position: relative;

    top: -5px;

    left: 5px;

}


/* ============================================================
   COMPANY LOGO
   ============================================================ */

.company-logo {

    width: 100%;

    height: 100%;

    object-fit: contain;

    object-position: center;

    display: block;

}


/* ============================================================
   LEFT VERTICAL DIVIDER
   ============================================================ */

.vertical-line {

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

    margin-left: 18px;

    margin-right: 20px;

    flex-shrink: 0;

}


/* ============================================================
   CENTER TITLE AREA
   ============================================================ */

.title-area {

    position: absolute;

    left: 50%;

    top: 0;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: center;

    text-align: center;

    min-width: max-content;

    box-sizing: border-box;

    transform: translateX(-50%);

}


/* ============================================================
   MAIN TITLE
   ============================================================ */

.main-title {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 27px;

    font-weight: 900;

    line-height: 1;

    letter-spacing: 0.3px;

    white-space: nowrap;

    margin: 0;

    padding: 0;

}


/* ============================================================
   ORANGE TITLE PART
   ============================================================ */

.main-title-orange {

    color: #f28c00;

}


/* ============================================================
   SUBTITLE
   ============================================================ */

.subtitle {

    color: #ffffff;

    font-family:
        Arial,
        sans-serif;

    font-size: 12px;

    font-weight: 400;

    letter-spacing: 3.6px;

    margin-top: 8px;

    line-height: 1;

    white-space: nowrap;

}


/* ============================================================
   SUB-SUBTITLE / TAGLINE
   ============================================================ */

.tagline {

    color:
        rgba(255,255,255,0.82);

    font-family:
        Arial,
        sans-serif;

    font-size: 7px;

    font-weight: 500;

    letter-spacing: 2.2px;

    margin-top: 6px;

    line-height: 1;

    white-space: nowrap;

}


/* ============================================================
   RIGHT DATE / TIME AREA
   ============================================================ */

.psm-right {

    position: absolute;

    right: 16px;

    top: 0;

    width: 15%;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: flex-end;

    text-align: right;

    color: #ffffff;

    z-index: 30;

    padding-left: 18px;

    box-sizing: border-box;

}


/* ============================================================
   RIGHT VERTICAL DIVIDER
   ============================================================ */

.psm-right::before {

    content: "";

    position: absolute;

    left: 0;

    top: 6px;

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

}


/* ============================================================
   DATE
   ============================================================ */

.date {

    color:
        rgba(255,255,255,0.95);

    font-family:
        Arial,
        sans-serif;

    font-size: 11px;

    font-weight: 400;

    letter-spacing: 0.7px;

    line-height: 1;

    margin: 0;

    padding: 0;

}


/* ============================================================
   TIME
   ============================================================ */

.time {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 22px;

    font-weight: 800;

    margin-top: 4px;

    line-height: 1;

    padding: 0;

}


/* ============================================================
   RIGHT HORIZONTAL LINE
   ============================================================ */

.right-line {

    width: 80px;

    height: 2px;

    background:
        rgba(255,255,255,0.75);

    margin-top: 7px;

    flex-shrink: 0;

}


/* ============================================================
   ORANGE BOTTOM BAR
   ============================================================ */

.orange-bar {

    position: absolute;

    left: 0;

    bottom: 0;

    width: 100%;

    height: 9px;

    background: #f28c00;

    z-index: 50;

}


</style>

</head>


<body>


<!-- ============================================================
     MAIN HEADER
     ============================================================ -->

<div class="psm-header">


    <!-- ========================================================
         LEFT LOGO AREA
         ======================================================== -->

    <div class="psm-left">


        <!-- ====================================================
             LOGO
             ==================================================== -->

        <div class="logo-panel">

            <img
                class="company-logo"
                src="data:image/jpeg;base64,LOGO_IMAGE_BASE64"
                alt="JSW JFE Steel Limited"
            >

        </div>


        <!-- ====================================================
             LEFT VERTICAL LINE
             ==================================================== -->

        <div class="vertical-line"></div>


        <!-- ====================================================
             CENTER TITLE GROUP
             ==================================================== -->

        <div class="title-area">


            <!-- MAIN TITLE -->

            <div class="main-title">

                PROCESS SAFETY INCIDENT(PSI)

                <span class="main-title-orange"></span>

            </div>


            <!-- SUBTITLE -->

            <div class="subtitle">

                PSM DIGITAL DASHBOARD

            </div>


            <!-- SUB-SUBTITLE -->

            <div class="tagline">

                PEOPLE
                &nbsp; | &nbsp;
                PROCESS
                &nbsp; | &nbsp;
                RISK
                &nbsp; | &nbsp;
                COMPLIANCE

            </div>


        </div>


    </div>


    <!-- ========================================================
         RIGHT DATE / TIME
         ======================================================== -->

    <div class="psm-right">


        <div class="date">

            CURRENT_DATE_VALUE

        </div>


        <div class="time">

            CURRENT_TIME_VALUE

        </div>


        <div class="right-line"></div>


    </div>


    <!-- ========================================================
         ORANGE BOTTOM BAR
         ======================================================== -->

    <div class="orange-bar"></div>


</div>


</body>

</html>

"""


# ============================================================
# INSERT LOGO
# ============================================================

header_html = header_html.replace(
    "LOGO_IMAGE_BASE64",
    logo_base64
)


# ============================================================
# INSERT DATE
# ============================================================

header_html = header_html.replace(
    "CURRENT_DATE_VALUE",
    current_date
)


# ============================================================
# INSERT TIME
# ============================================================

header_html = header_html.replace(
    "CURRENT_TIME_VALUE",
    current_time
)


# ============================================================
# DISPLAY HEADER
# ============================================================

components.html(
    header_html,
    height=114,
    scrolling=False
)

st.markdown(
    """
    <style>
    div[data-testid="stVerticalBlock"] > div:has(> iframe) {
        margin-bottom: -65px !important;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# ============================================================
# COLOURS
# ============================================================

JSW_BLUE = "#0057B8"
JSW_RED = "#E31E2F"
GREEN = "#159447"
AMBER = "#E5A400"

DARK_BLUE = "#17365D"
TEXT = "#26374A"
MUTED = "#718096"
BORDER = "#D6E1EC"
LIGHT_BLUE = "#EEF5FB"

# ============================================================
# CSS
# ============================================================

st.markdown(
    """ 
    <style> 

    .stApp { 
        background: #FFFFFF; 
    } 

    header { 
        visibility: hidden; 
    } 

    #MainMenu { 
        visibility: hidden; 
    } 

    footer { 
        visibility: hidden; 
    } 

    .block-container { 
        max-width: 1700px; 
        padding-top: 1rem; 
        padding-left: 1rem; 
        padding-right: 1rem; 
        padding-bottom: 1rem; 
    } 

    div[data-testid="stSelectbox"] label { 
        color: #17365D; 
        font-size: 13px; 
        font-weight: 600; 
    } 

    .kpi-card { 
        position: relative; 
        height: 112px; 
        width: 100%; 
        background: #FFFFFF; 
        border: 1px solid #D6E1EC; 
        border-radius: 10px; 
        padding: 15px 18px 12px 24px; 
        box-sizing: border-box; 
        overflow: hidden; 
        box-shadow: 0 2px 7px rgba(20, 60, 100, 0.04); 
    } 

    .kpi-card::before { 
        content: ""; 
        position: absolute; 
        left: 0; 
        top: 0; 
        bottom: 0; 
        width: 6px; 
        background: var(--accent); 
        border-radius: 10px 0 0 10px; 
    } 

    .kpi-title { 
        font-size: 15px; 
        font-weight: 700; 
        color: var(--accent); 
        text-transform: uppercase; 
        letter-spacing: 0.3px; 
        margin-bottom: 7px; 
        white-space: nowrap; 
        overflow: hidden; 
        text-overflow: ellipsis; 
    } 

    .kpi-number { 
        font-size: 34px; 
        line-height: 1; 
        font-weight: 700; 
        color: var(--accent); 
    } 

    .kpi-subtitle { 
        font-size: 14px; 
        color: var(--accent); 
        margin-top: 10 px; 
        white-space: nowrap; 
        overflow: hidden; 
        text-overflow: ellipsis; 
    } 

    .level-card-space { 
        padding-top: 0; 
    } 

    .level-simple-card { 
        height: 140px; 
        width: 100%; 
        border: 1px solid currentColor; 
        border-radius: 10px; 
        padding: 18px 18px 14px 20px; 
        box-sizing: border-box; 
    } 

    .level-simple-title { 
        font-size: 15px; 
        font-weight: 700; 
        text-transform: uppercase; 
        letter-spacing: 0.25px; 
        margin-bottom: 10px; 
    } 

    .level-simple-number { 
        font-size: 38px; 
        line-height: 1; 
        font-weight: 700; 
    } 

    .level-simple-subtitle { 
        font-size: 13px; 
        margin-top: 10px; 
        opacity: 0.78; 
    } 

    .psi-html-table-wrap { 
        width: 100%; 
        height: 430px; 
        overflow-y: auto; 
        overflow-x: hidden; 
        border: 1px solid #D6E1EC; 
        border-radius: 8px; 
        box-sizing: border-box; 
        scrollbar-width: thin; 
        scrollbar-color: #AEBCCC #F1F4F8; 
    } 

    .psi-html-table-wrap::-webkit-scrollbar { 
        width: 9px; 
    } 

    .psi-html-table-wrap::-webkit-scrollbar-track { 
        background: #F1F4F8; 
        border-radius: 8px; 
    } 

    .psi-html-table-wrap::-webkit-scrollbar-thumb { 
        background: #AEBCCC; 
        border-radius: 8px; 
    } 

    .psi-html-table-wrap::-webkit-scrollbar-thumb:hover { 
        background: #7F91A8; 
    } 

    .psi-html-table { 
        width: 100%; 
        max-width: 100%; 
        table-layout: fixed; 
        border-collapse: collapse; 
        font-size: 13px; 
        color: #26374A; 
        background: #FFFFFF; 
    } 

    .psi-html-table th { 
        position: sticky; 
        top: 0; 
        z-index: 2; 
        background: #17365D; 
        color: #FFFFFF; 
        font-weight: 700; 
        text-align: left; 
        padding: 11px 10px; 
        border-right: 1px solid rgba(255,255,255,0.22); 
        white-space: normal; 
        overflow-wrap: anywhere; 
    } 

    .psi-html-table td { 
        padding: 10px 10px; 
        border-top: 1px solid #E1E7EF; 
        vertical-align: top; 
        overflow-wrap: anywhere; 
        word-break: normal; 
    } 

    .psi-html-table tbody tr:nth-child(even) td { 
        background: #FBFCFE; 
    } 

    .psi-doc-link { 
        display: inline-block; 
        background: #0057B8; 
        color: #FFFFFF !important; 
        text-decoration: none !important; 
        font-weight: 700; 
        font-size: 12px; 
        line-height: 1; 
        padding: 7px 14px; 
        border-radius: 7px; 
        white-space: nowrap; 
        text-align: center; 
        min-width: 34px; 
        box-sizing: border-box; 
    } 

    .psi-doc-link:hover { 
        background: #004A9F; 
        color: #FFFFFF !important; 
        text-decoration: none !important; 
    } 

    .psi-no-doc { 
        color: #718096; 
        white-space: nowrap; 
    } 

    .psi-document-cell { 
        text-align: center; 
        vertical-align: middle !important; 
    } 

    .section-title { 
        font-size: 14px; 
        font-weight: 700; 
        color: #17365D; 
        padding: 3px 2px 8px 2px; 
        letter-spacing: 0.25px; 
    } 

    .vertical-gap { 
        height: 20px; 
        min-height: 20px; 
        margin: 0 !important; 
        padding: 0 !important; 
    } 

    .section-card { 
        margin-top: 0 !important; 
        margin-bottom: 0 !important; 
    } 

    .section-card + .section-card { 
        margin-top: 10px !important; 
    } 

    div[data-testid="stDataFrame"] { 
        border: 1px solid #D6E1EC; 
        border-radius: 8px; 
        overflow: hidden; 
    } 

    </style> 
    """,
    unsafe_allow_html=True
)

# ============================================================
# GOOGLE SHEET CONFIGURATION
# ============================================================

SHEET_ID = "1--X0TT5Ts92EKAxrhV-fQgqeTHBX3rDVc1Egg74MewM"
SHEET_NAME = "PSI"

CSV_URL = (
        "https://docs.google.com/spreadsheets/d/"
        + SHEET_ID
        + "/gviz/tq?tqx=out:csv&sheet="
        + SHEET_NAME
)

XLSX_URL = (
        "https://docs.google.com/spreadsheets/d/"
        + SHEET_ID
        + "/export?format=xlsx"
)


# ============================================================
# DATA LOADING
# ============================================================

@st.cache_data(ttl=60)
def load_csv_data():
    data = pd.read_csv(CSV_URL)

    data.columns = (
        data.columns
        .astype(str)
        .str.strip()
        .str.replace("\n", " ", regex=False)
        .str.replace(r"\s+", " ", regex=True)
    )

    data = data.dropna(how="all")
    data = data.dropna(axis=1, how="all")

    return data


@st.cache_data(ttl=300)
def load_document_links():
    """
    Reads the actual hyperlink target from the Google Sheet XLSX.

    This is intentionally separate from the CSV because CSV export
    can return only the displayed text of a hyperlink instead of
    its real URL.
    """

    links = []

    try:
        import requests
        from openpyxl import load_workbook
        from io import BytesIO

        response = requests.get(
            XLSX_URL,
            timeout=30
        )

        response.raise_for_status()

        workbook = load_workbook(
            filename=BytesIO(response.content),
            data_only=False,
            read_only=False
        )

        if SHEET_NAME not in workbook.sheetnames:
            return []

        sheet = workbook[SHEET_NAME]

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[
                    str(cell.value).strip()
                ] = cell.column

        target_column = None

        preferred_names = [
            "Attach Incident Report",
            "Attach Incident Reports",
            "Incident Report",
            "Attach Document",
            "Attached Document",
            "Document Link"
        ]

        for preferred in preferred_names:
            for header, column_number in headers.items():
                if header.lower() == preferred.lower():
                    target_column = column_number
                    break

            if target_column is not None:
                break

        if target_column is None:
            for header, column_number in headers.items():
                lower_header = header.lower()

                if (
                        "attach incident report" in lower_header
                        or "incident report" in lower_header
                ):
                    target_column = column_number
                    break

        if target_column is None:
            return []

        for row_number in range(
                2,
                sheet.max_row + 1
        ):

            cell = sheet.cell(
                row=row_number,
                column=target_column
            )

            url = None

            if cell.hyperlink is not None:
                url = cell.hyperlink.target

            elif isinstance(cell.value, str):

                text = cell.value.strip()

                if text.startswith("http://"):
                    url = text

                elif text.startswith("https://"):
                    url = text

                elif "HYPERLINK(" in text.upper():

                    first_quote = text.find('"')

                    if first_quote != -1:

                        second_quote = text.find(
                            '"',
                            first_quote + 1
                        )

                        if second_quote != -1:
                            url = text[
                                first_quote + 1:
                                second_quote
                            ]

            if url is not None:
                url = str(url).strip()

                if not (
                        url.startswith("http://")
                        or url.startswith("https://")
                ):
                    url = None

            links.append(url)

        return links

    except Exception:
        return []


try:
    df = load_csv_data()

except Exception as error:

    st.error(
        "Unable to load the PSI tab from Google Sheets."
    )

    st.write(
        "Please check that the Google Sheet is accessible "
        "and the tab is named PSI."
    )

    st.code(str(error))

    st.stop()


# ============================================================
# FIND COLUMN
# ============================================================

def find_column(columns, possible_names):
    for name in possible_names:

        for column in columns:

            if (
                    str(column).strip().lower()
                    == name.lower()
            ):
                return column

    for name in possible_names:

        for column in columns:

            if name.lower() in str(column).lower():
                return column

    return None


# ============================================================
# IDENTIFY COLUMNS
# ============================================================

date_col = find_column(
    df.columns,
    [
        "Incident Date",
        "Incident date",
        "Date of Incident",
        "Event Date"
    ]
)

department_col = find_column(
    df.columns,
    [
        "Department",
        "Dept"
    ]
)

description_col = find_column(
    df.columns,
    [
        "Incident Description",
        "Description",
        "Incident Details"
    ]
)

classification_col = find_column(
    df.columns,
    [
        "Incident Classification",
        "Classification",
        "Incident Type"
    ]
)

level_col = find_column(
    df.columns,
    [
        "Incident Level",
        "PSI Level",
        "Level"
    ]
)

investigation_col = find_column(
    df.columns,
    [
        "Investigation Status",
        "Investigation status",
        "Investigation"
    ])

# ============================================================
# CHECK DATE COLUMN
# ============================================================

if date_col is None:
    st.error(
        "Incident Date column was not found."
    )

    st.write(
        "Columns found in PSI sheet:"
    )

    st.write(list(df.columns))

    st.stop()

# ============================================================
# PRESERVE ALL INCIDENT ROWS
# ============================================================

# IMPORTANT:
# Do NOT drop rows just because Incident Date could not be parsed.
# The dashboard must still count all incidents in the sheet.

df["_Original_Row"] = range(
    2,
    len(df) + 2
)

# ============================================================
# ROBUST DATE PARSING
# ============================================================

original_dates = df[date_col].copy()

parsed_dates = pd.to_datetime(
    original_dates,
    errors="coerce",
    dayfirst=True
)

# Try mixed format for any values that were not parsed.
try:

    missing_date_mask = parsed_dates.isna()

    if missing_date_mask.any():
        parsed_dates.loc[missing_date_mask] = (
            pd.to_datetime(
                original_dates.loc[missing_date_mask],
                errors="coerce",
                format="mixed",
                dayfirst=True
            )
        )

except (TypeError, ValueError):

    pass

df["_Parsed_Date"] = parsed_dates


# ============================================================
# FINANCIAL YEAR
# ============================================================

def get_financial_year(date):
    if pd.isna(date):
        return ""

    if date.month >= 4:
        return (
                str(date.year)
                + "-"
                + str(date.year + 1)[-2:]
        )

    return (
            str(date.year - 1)
            + "-"
            + str(date.year)[-2:]
    )


df["Financial Year"] = (
    df["_Parsed_Date"].apply(
        get_financial_year
    )
)

# ============================================================
# MONTH
# ============================================================

df["Month"] = ""

valid_date_mask = df["_Parsed_Date"].notna()

df.loc[
    valid_date_mask,
    "Month"
] = df.loc[
    valid_date_mask,
    "_Parsed_Date"
].dt.strftime("%b-%y")

df["Month Date"] = pd.NaT

df.loc[
    valid_date_mask,
    "Month Date"
] = (
    df.loc[
        valid_date_mask,
        "_Parsed_Date"
    ]
    .dt.to_period("M")
    .dt.to_timestamp()
)

# ============================================================
# DOCUMENT LINKS
# ============================================================

document_links = load_document_links()

if document_links:

    link_map = {}

    for index, url in enumerate(
            document_links,
            start=2
    ):
        link_map[index] = url

    df["_Document_URL"] = df[
        "_Original_Row"
    ].map(link_map)

else:

    df["_Document_URL"] = None

# ============================================================
# FILTERS
# ============================================================

filter_1, filter_2, filter_3 = st.columns(3)

with filter_1:
    financial_years = sorted(
        [
            year
            for year in df["Financial Year"].unique()
            if year != ""
        ],
        reverse=True
    )

    selected_fy = st.selectbox(
        "Financial Year",
        ["All Financial Years"]
        + financial_years
    )

with filter_2:
    month_source = df.copy()

    if selected_fy != "All Financial Years":
        month_source = month_source[
            month_source["Financial Year"]
            == selected_fy
            ]

    months = (
        month_source[
            month_source["Month"] != ""
            ][
            ["Month", "Month Date"]
        ]
        .drop_duplicates()
        .sort_values(
            "Month Date",
            ascending=False
        )["Month"]
        .tolist()
    )

    selected_month = st.selectbox(
        "Month",
        ["All Months"] + months
    )

with filter_3:
    if department_col is not None:

        departments = sorted(
            df[department_col]
            .dropna()
            .astype(str)
            .str.strip()
            .unique()
            .tolist()
        )

        selected_department = st.selectbox(
            "Department",
            ["All Departments"]
            + departments
        )

    else:

        selected_department = (
            "All Departments"
        )

    # ============================================================
# APPLY FILTERS
# ============================================================

data = df.copy()

if selected_fy != "All Financial Years":
    data = data[
        data["Financial Year"]
        == selected_fy
        ]

if selected_month != "All Months":
    data = data[
        data["Month"]
        == selected_month
        ]

if (
        department_col is not None
        and selected_department
        != "All Departments"
):
    data = data[
        data[department_col]
        .astype(str)
        .str.strip()
        == selected_department
        ]

# ============================================================
# KPI CALCULATIONS
# ============================================================

# This now counts every incident row.
# Therefore the unfiltered total will remain 14
# even if one or more date cells cannot be parsed.

total_incidents = len(data)

completed = 0

if investigation_col is not None:
    status = (
        data[investigation_col]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.lower()
    )

    completed = int(
        status.str.contains(
            "completed|complete|closed|done",
            regex=True,
            na=False
        ).sum()
    )

pending = max(
    total_incidents - completed,
    0
)

if total_incidents > 0:

    compliance = (
            completed
            / total_incidents
            * 100
    )

else:

    compliance = 0.0


# ============================================================
# KPI FUNCTION
# ============================================================

def show_kpi(
        title,
        value,
        subtitle,
        accent
):
    # Keep HTML at the start of the markdown line.
    # Leading indentation makes Streamlit treat the block as code.
    html = f"""<div class="kpi-card" style="--accent:{accent};"> 
<div class="kpi-title">{title}</div> 
<div class="kpi-number">{value}</div> 
<div class="kpi-subtitle">{subtitle}</div> 
</div>"""

    st.markdown(
        html,
        unsafe_allow_html=True
    )


# ============================================================
# MAIN KPI ROW
# ============================================================

kpi_1, kpi_2, kpi_3, kpi_4 = (
    st.columns(4)
)

with kpi_1:
    show_kpi(
        "Total Incidents",
        f"{total_incidents:,}",
        "Incidents recorded",
        JSW_BLUE
    )

with kpi_2:
    show_kpi(
        "Investigation Completed",
        f"{completed:,}",
        f"{compliance:.1f}% of total incidents",
        GREEN
    )

with kpi_3:
    show_kpi(
        "Investigation Pending",
        f"{pending:,}",
        "Yet to be completed",
        JSW_RED
    )

with kpi_4:
    show_kpi(
        "Investigation Compliance",
        f"{compliance:.1f}%",
        "Completed / Total incidents",
        AMBER
    )

# ============================================================
# GAP
# ============================================================

st.markdown(
    '<div class="vertical-gap"></div>',
    unsafe_allow_html=True
)


# ============================================================
# CHART STYLE
# ============================================================

def style_chart(fig):
    fig.update_layout(
        paper_bgcolor="white",
        plot_bgcolor="white",
        font=dict(
            family="Arial",
            color=TEXT,
            size=12
        ),
        margin=dict(
            l=45,
            r=25,
            t=10,
            b=45
        ),
        showlegend=False
    )

    fig.update_xaxes(
        showgrid=False,
        zeroline=False
    )

    fig.update_yaxes(
        showgrid=True,
        gridcolor="#E6EDF4",
        zeroline=False
    )

    return fig


# ============================================================
# DEPARTMENT CHART
# ============================================================

department_fig = None

if department_col is not None:
    department_data = (
        data[department_col]
        .fillna("Not Available")
        .astype(str)
        .str.strip()
        .value_counts()
        .reset_index()
    )

    department_data.columns = [
        "Department",
        "Incidents"
    ]

    department_data = department_data.sort_values(
        "Incidents",
        ascending=True
    )

    department_fig = px.bar(
        department_data,
        x="Incidents",
        y="Department",
        orientation="h",
        text="Incidents"
    )

    department_fig.update_traces(
        marker_color=JSW_BLUE,
        textposition="outside",
        cliponaxis=False
    )

    max_department = (
        int(
            department_data["Incidents"].max()
        )
        if not department_data.empty
        else 1
    )

    department_fig.update_xaxes(
        title=None,
        dtick=1,
        range=[
            0,
            max_department + 0.7
        ]
    )

    department_fig.update_yaxes(
        title=None
    )

    department_fig.update_layout(
        height=310
    )

    department_fig = style_chart(
        department_fig
    )

# ============================================================
# MONTH-WISE TREND
# ============================================================

month_data = (
    data[
        data["Month"] != ""
        ]
    .groupby(
        ["Month Date", "Month"]
    )
    .size()
    .reset_index(
        name="Incidents"
    )
    .sort_values(
        "Month Date"
    )
)

month_fig = px.line(
    month_data,
    x="Month",
    y="Incidents",
    markers=True,
    text="Incidents"
)

# Marker is deliberately the SAME colour as line.
month_fig.update_traces(
    line=dict(
        color=JSW_BLUE,
        width=3
    ),
    marker=dict(
        color=JSW_BLUE,
        size=8
    ),
    textposition="top center"
)

max_month = (
    int(
        month_data["Incidents"].max()
    )
    if not month_data.empty
    else 1
)

month_fig.update_yaxes(
    title=None,
    dtick=1,
    range=[
        0,
        max_month + 0.5
    ]
)

month_fig.update_xaxes(
    title=None
)

month_fig.update_layout(
    height=310
)

month_fig = style_chart(
    month_fig
)

# ============================================================
# CLASSIFICATION CHART
# ============================================================

classification_fig = None

if classification_col is not None:
    classification_data = (
        data[classification_col]
        .fillna("Not Available")
        .astype(str)
        .str.strip()
        .value_counts()
        .reset_index()
    )

    classification_data.columns = [
        "Classification",
        "Incidents"
    ]

    classification_data = (
        classification_data
        .sort_values(
            "Incidents",
            ascending=True
        )
    )

    classification_fig = px.bar(
        classification_data,
        x="Incidents",
        y="Classification",
        orientation="h",
        text="Incidents"
    )

    classification_fig.update_traces(
        marker_color=JSW_RED,
        textposition="outside",
        cliponaxis=False
    )

    max_classification = (
        int(
            classification_data[
                "Incidents"
            ].max()
        )
        if not classification_data.empty
        else 1
    )

    classification_fig.update_xaxes(
        title=None,
        dtick=1,
        range=[
            0,
            max_classification + 0.7
        ]
    )

    classification_fig.update_yaxes(
        title=None
    )

    classification_fig.update_layout(
        height=310
    )

    classification_fig = style_chart(
        classification_fig
    )

# ============================================================
# CHART SECTIONS
# ============================================================

chart_1, chart_2, chart_3 = (
    st.columns(3)
)

with chart_1:
    with st.container(border=True):

        st.markdown(
            '<div class="section-title">'
            'INCIDENTS BY DEPARTMENT'
            '</div>',
            unsafe_allow_html=True
        )

        if department_fig is not None:

            st.plotly_chart(
                department_fig,
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )

        else:

            st.info(
                "Department data unavailable."
            )

with chart_2:
    with st.container(border=True):
        st.markdown(
            '<div class="section-title">'
            'MONTH-WISE INCIDENT TREND'
            '</div>',
            unsafe_allow_html=True
        )

        st.plotly_chart(
            month_fig,
            use_container_width=True,
            config={
                "displayModeBar": False
            }
        )

with chart_3:
    with st.container(border=True):

        st.markdown(
            '<div class="section-title">'
            'INCIDENT CLASSIFICATION'
            '</div>',
            unsafe_allow_html=True
        )

        if classification_fig is not None:

            st.plotly_chart(
                classification_fig,
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )

        else:

            st.info(
                "Classification data unavailable."
            )

        # ============================================================
# GAP
# ============================================================

st.markdown(
    '<div class="vertical-gap"></div>',
    unsafe_allow_html=True
)

# ============================================================
# INCIDENT LEVEL DATA
# ============================================================

if level_col is not None:

    level_data = (
        data[level_col]
        .fillna("Not Available")
        .astype(str)
        .str.strip()
        .value_counts()
        .reset_index()
    )

    level_data.columns = [
        "Level",
        "Incidents"
    ]

else:

    level_data = pd.DataFrame(
        columns=[
            "Level",
            "Incidents"
        ]
    )

# ============================================================
# INCIDENT LEVEL SECTION
# ============================================================

with st.container(border=True):
    st.markdown(
        '<div class="section-title">'
        'INCIDENT LEVEL DISTRIBUTION'
        '</div>',
        unsafe_allow_html=True
    )

    if not level_data.empty:

        level_order = [
            "Level 1",
            "Level 2",
            "Level 3",
            "Level 4"
        ]

        level_colors = {
            "Level 1": GREEN,
            "Level 2": JSW_BLUE,
            "Level 3": AMBER,
            "Level 4": JSW_RED
        }

        level_data["Sort"] = (
            level_data["Level"].apply(
                lambda value:
                level_order.index(value)
                if value in level_order
                else 99
            )
        )

        level_data = (
            level_data
            .sort_values("Sort")
            .drop(columns="Sort")
        )

        # ----------------------------------------------------
        # DONUT + LEVEL CARDS
        # ----------------------------------------------------

        donut_col, cards_col = (
            st.columns(
                [1, 3],
                vertical_alignment="center"
            )
        )

        # ----------------------------------------------------
        # DONUT
        # ----------------------------------------------------

        with donut_col:

            donut_colors = [
                level_colors.get(
                    level,
                    JSW_BLUE
                )
                for level
                in level_data["Level"]
            ]

            donut = go.Figure(
                data=[
                    go.Pie(
                        labels=level_data[
                            "Level"
                        ],
                        values=level_data[
                            "Incidents"
                        ],
                        hole=0.68,
                        textinfo="none",
                        marker=dict(
                            colors=donut_colors,
                            line=dict(
                                color="white",
                                width=3
                            )
                        )
                    )
                ]
            )

            donut.update_layout(
                height=270,
                paper_bgcolor="white",
                plot_bgcolor="white",
                margin=dict(
                    l=5,
                    r=5,
                    t=5,
                    b=5
                ),
                showlegend=True,
                legend=dict(
                    orientation="v",
                    x=1.0,
                    y=0.5
                ),
                annotations=[
                    dict(
                        text=(
                                "<b>"
                                + str(total_incidents)
                                + "</b>"
                                  "<br>"
                                  "<span style="
                                  "'font-size:11px'>"
                                  "TOTAL INCIDENTS"
                                  "</span>"
                        ),
                        x=0.5,
                        y=0.5,
                        showarrow=False,
                        font=dict(
                            size=17,
                            color=DARK_BLUE
                        )
                    )
                ]
            )

            st.plotly_chart(
                donut,
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )

            # ----------------------------------------------------
        # LEVEL KPI CARDS
        # ----------------------------------------------------

        with cards_col:

            card_1, card_2, card_3, card_4 = (
                st.columns(4)
            )

            card_columns = [
                card_1,
                card_2,
                card_3,
                card_4
            ]

            for index, level in enumerate(
                    level_order
            ):

                matching = level_data[
                    level_data["Level"]
                    == level
                    ]

                if matching.empty:

                    count = 0

                else:

                    count = int(
                        matching[
                            "Incidents"
                        ].iloc[0]
                    )

                percentage = (
                    count
                    / total_incidents
                    * 100
                    if total_incidents > 0
                    else 0
                )

                with card_columns[index]:

                    # Soft tinted background based on the donut segment color.
                    level_styles = {
                        "Level 1": (
                            "#159447",
                            "#EAF7EF"
                        ),
                        "Level 2": (
                            "#0057B8",
                            "#EAF2FB"
                        ),
                        "Level 3": (
                            "#E5A400",
                            "#FFF7E0"
                        ),
                        "Level 4": (
                            "#E31E2F",
                            "#FDECEF"
                        )
                    }

                    level_color, level_bg = level_styles.get(
                        level,
                        (JSW_BLUE, "#EAF2FB")
                    )

                    st.markdown(
                        f'''<div class="level-card-space"> 
<div class="level-simple-card" style="color:{level_color}; background:{level_bg}; border-color:{level_color}33;"> 
<div class="level-simple-title">{level}</div> 
<div class="level-simple-number">{count}</div> 
<div class="level-simple-subtitle">{percentage:.1f}% of incidents</div> 
</div> 
</div>''',
                        unsafe_allow_html=True
                    )

    else:

        st.info(
            "Incident level data unavailable."
        )

    # ============================================================
# GAP
# ============================================================

st.markdown(
    '<div class="vertical-gap"></div>',
    unsafe_allow_html=True
)

# ============================================================
# PROCESS SAFETY INCIDENT DETAILS
# ============================================================

with st.container(border=True):
    st.markdown(
        '<div class="section-title">'
        'PROCESS SAFETY INCIDENT DETAILS'
        '</div>',
        unsafe_allow_html=True
    )

    # --------------------------------------------------------
    # TABLE
    # --------------------------------------------------------

    table = pd.DataFrame(
        index=data.index
    )

    if department_col is not None:
        table["Department"] = (
            data[department_col]
            .fillna("")
            .astype(str)
        )

    if description_col is not None:
        table["Incident Description"] = (
            data[description_col]
            .fillna("")
            .astype(str)
        )

        # Preserve original date if parsing failed.
    display_dates = (
        data[date_col]
        .fillna("")
        .astype(str)
    )

    parsed_date_text = data[
        "_Parsed_Date"
    ].dt.strftime("%d-%b-%y")

    for row_index in table.index:

        parsed_value = parsed_date_text.loc[
            row_index
        ]

        if pd.notna(parsed_value):
            display_dates.loc[
                row_index
            ] = parsed_value

    table["Incident Date"] = display_dates

    if classification_col is not None:
        table["Incident Classification"] = (
            data[classification_col]
            .fillna("")
            .astype(str)
        )

    if level_col is not None:
        table["Incident Level"] = (
            data[level_col]
            .fillna("")
            .astype(str)
        )

    if investigation_col is not None:

        table["Investigation Status"] = (
            data[investigation_col]
            .fillna("")
            .astype(str)
        )

    else:

        table["Investigation Status"] = ""

        # --------------------------------------------------------
    # DOCUMENT COLUMN
    # --------------------------------------------------------

    # IMPORTANT:
    # Only use the actual URL extracted from the
    # "Attach Incident Report" column.
    #
    # We do NOT use the displayed text from the sheet.
    # This prevents "View Document" from linking back
    # to the dashboard itself.

    document_values = []

    for row_index in data.index:

        url = data.loc[
            row_index,
            "_Document_URL"
        ]

        if (
                isinstance(url, str)
                and (
                url.startswith("https://")
                or url.startswith("http://")
        )
        ):

            document_values.append(url)

        else:

            document_values.append(
                "No Report Available"
            )

    table["View Document"] = (
        document_values
    )

    # --------------------------------------------------------
    # COLUMN ORDER
    # --------------------------------------------------------

    requested_columns = [
        "Department",
        "Incident Description",
        "Incident Date",
        "Incident Classification",
        "Incident Level",
        "Investigation Status",
        "View Document"
    ]

    table = table[
        [
            column
            for column in requested_columns
            if column in table.columns
        ]
    ]

    # --------------------------------------------------------
    # DISPLAY
    # --------------------------------------------------------

    # Use a real HTML table so only rows with a genuine URL are
    # clickable. Blank report cells show "No Document Available"
    # and have no href, so they cannot redirect to this dashboard.

    display_columns = [
        "Department",
        "Incident Description",
        "Incident Date",
        "Incident Classification",
        "Incident Level",
        "Investigation Status",
        "View Document"
    ]

    table_html = [
        '<div class="psi-html-table-wrap">',
        '<table class="psi-html-table">',
        '<colgroup>',
        '<col style="width:9%;">',  # Department
        '<col style="width:33%;">',  # Incident Description
        '<col style="width:9%;">',  # Incident Date
        '<col style="width:15%;">',  # Incident Classification
        '<col style="width:9%;">',  # Incident Level
        '<col style="width:13%;">',  # Investigation Status
        '<col style="width:12%;">',  # View Document
        '</colgroup>',
        '<thead><tr>'
    ]

    for column in display_columns:
        if column in table.columns:
            table_html.append(
                f"<th>{html.escape(str(column))}</th>"
            )

    table_html.append("</tr></thead><tbody>")

    for row_index, row in table.iterrows():

        table_html.append("<tr>")

        for column in display_columns:

            if column not in table.columns:
                continue

            value = row[column]

            if pd.isna(value):
                value = ""

            value = str(value)

            if column == "View Document":

                url = value.strip()

                if (
                        url.startswith("https://")
                        or url.startswith("http://")
                ):
                    cell_html = (
                        '<a class="psi-doc-link" '
                        f'href="{html.escape(url, quote=True)}" '
                        'target="_blank" rel="noopener noreferrer">'
                        'View'
                        '</a>'
                    )
                else:
                    cell_html = (
                        '<span class="psi-no-doc">'
                        'No Document Available'
                        '</span>'
                    )

            else:
                cell_html = html.escape(value)

            td_class = (
                ' class="psi-document-cell"'
                if column == "View Document"
                else ""
            )
            table_html.append(
                f"<td{td_class}>{cell_html}</td>"
            )

        table_html.append("</tr>")

    table_html.extend([
        "</tbody></table>",
        "</div>"
    ])

    st.markdown(
        "".join(table_html),
        unsafe_allow_html=True
    )

# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """ 
    <div style=" 
        text-align:center; 
        color:#8493A3; 
        font-size:11px; 
        padding:14px 0 5px 0; 
    "> 
        Process Safety Management Dashboard 
    </div> 
    """,
    unsafe_allow_html=True
)

