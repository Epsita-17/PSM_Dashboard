import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import os
import base64
import re
import json
from io import BytesIO
from urllib.request import Request, urlopen
from openpyxl import load_workbook
from streamlit_autorefresh import st_autorefresh
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo

# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="PSM Digital Dashboard",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ============================================================
# REMOVE STREAMLIT TOP SPACE
# ============================================================

st.markdown(
    """
    <style>

    html,
    body {
        margin: 0 !important;
        padding: 0 !important;
    }

    [data-testid="stAppViewContainer"] {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }

    [data-testid="stAppViewContainer"] > .main {
        padding-top: 0 !important;
        margin-top: 0 !important;
    }

    [data-testid="stHeader"] {
    background: transparent !important;
}
        height: 0 !important;
        min-height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    [data-testid="stToolbar"] {
        display: none !important;
    }

    [data-testid="stDecoration"] {
        display: none !important;
        height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    .block-container {
        padding-top: 0 !important;
        margin-top: 0 !important;
        padding-bottom: 0 !important;
        padding-left: 0 !important;
        padding-right: 0 !important;
        max-width: 100% !important;
    }

    .stApp {
        margin-top: 0 !important;
        padding-top: 0 !important;
    }

    iframe {
        display: block !important;
        margin-top: -0px !important;
        padding-top: 0 !important;
        border: 0 !important;
    }

    </style>
    """,
    unsafe_allow_html=True
)

# ============================================================
# BASE DIRECTORY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent


# ============================================================
# IMAGE TO BASE64
# ============================================================

def image_to_base64(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        return ""

    try:

        with open(file_path, "rb") as file:

            return base64.b64encode(
                file.read()
            ).decode("utf-8")

    except Exception:

        return ""


# ============================================================
# COMPANY LOGO
# ============================================================

logo_path = BASE_DIR / "jsw_jfe_logo.jpg"

logo_base64 = image_to_base64(logo_path)

# ============================================================
# FILE CHECK
# ============================================================

if not logo_base64:
    st.error(
        "jsw_jfe_logo.jpg not found. "
        "Keep jsw_jfe_logo.jpg in the same folder as this Python file."
    )

# ============================================================
# DATE AND TIME
# ============================================================

now = datetime.now(ZoneInfo("Asia/Kolkata"))

current_date = now.strftime(
    "%d %b %Y"
).upper()

current_time = now.strftime(
    "%I:%M %p"
)

# ============================================================
# HEADER HTML
# ============================================================

header_html = """

<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<style>


/* ============================================================
   MAIN HEADER
   ============================================================ */

.psm-header {

    position: relative;

    width: 100%;

    height: 90px;

    overflow: hidden;

    background:
        linear-gradient(
            90deg,
            #031d34 0%,
            #052b49 42%,
            #07385c 74%,
            #052b49 100%
        );

    border-radius: 7px;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.18);

}


/* ============================================================
   LEFT LOGO AREA
   ============================================================ */

.psm-left {

    position: absolute;

    left: 0;

    top: 0;

    width: 100%;

    height: 95px;

    display: flex;

    align-items: center;

    padding-left: 4px;

    box-sizing: border-box;

    z-index: 20;

    pointer-events: none;

}


/* ============================================================
   LOGO PANEL
   ONLY VERTICAL POSITION CHANGED
   ============================================================ */

.logo-panel {

    width: 195px;

    height: 68px;

    background: #ffffff;

    border-radius: 5px;

    padding: 4px;

    display: flex;

    align-items: center;

    justify-content: center;

    flex-shrink: 0;

    box-sizing: border-box;

    box-shadow:
        0 3px 9px
        rgba(0,0,0,0.20);

    position: relative;

    top: -6px;

    left: 5px;

}


/* ============================================================
   COMPANY LOGO
   ============================================================ */

.company-logo {

    width: 100%;

    height: 100%;

    object-fit: contain;

    object-position: center;

    display: block;

}


/* ============================================================
   LEFT VERTICAL DIVIDER
   ============================================================ */

.vertical-line {

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

    margin-left: 18px;

    margin-right: 20px;

    flex-shrink: 0;

}


/* ============================================================
   CENTER TITLE AREA
   ============================================================ */

.title-area {

    position: absolute;

    left: 50%;

    top: 0;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: center;

    text-align: center;

    min-width: max-content;

    box-sizing: border-box;

    transform: translateX(-50%);

}


/* ============================================================
   MAIN TITLE
   ============================================================ */

.main-title {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 27px;

    font-weight: 900;

    line-height: 1;

    letter-spacing: 0.3px;

    white-space: nowrap;

    margin: 0;

    padding: 0;

}


/* ============================================================
   ORANGE TITLE PART
   ============================================================ */

.main-title-orange {

    color: #f28c00;

}


/* ============================================================
   SUBTITLE
   ============================================================ */

.subtitle {

    color: #ffffff;

    font-family:
        Arial,
        sans-serif;

    font-size: 12px;

    font-weight: 400;

    letter-spacing: 3.6px;

    margin-top: 8px;

    line-height: 1;

    white-space: nowrap;

}

/* ============================================================
   SUB-SUBTITLE / TAGLINE
   ============================================================ */

.tagline {

    color:
        rgba(255,255,255,0.82);

    font-family:
        Arial,
        sans-serif;

    font-size: 7px;

    font-weight: 500;

    letter-spacing: 2.2px;

    margin-top: 6px;

    line-height: 1;

    white-space: nowrap;

}


/* ============================================================
   RIGHT DATE / TIME AREA
   ============================================================ */

.psm-right {

    position: absolute;

    right: 16px;

    top: 0;

    width: 15%;

    height: 95px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    align-items: flex-end;

    text-align: right;

    color: #ffffff;

    z-index: 30;

    padding-left: 18px;

    box-sizing: border-box;

}


/* ============================================================
   RIGHT VERTICAL DIVIDER
   ============================================================ */

.psm-right::before {

    content: "";

    position: absolute;

    left: 15px;

    top: 6px;

    width: 2px;

    height: 83px;

    background:
        rgba(255,255,255,0.65);

}


/* ============================================================
   DATE
   ============================================================ */

.date {

    color:
        rgba(255,255,255,0.95);

    font-family:
        Arial,
        sans-serif;

    font-size: 11px;

    font-weight: 400;

    letter-spacing: 0.7px;

    line-height: 1;

    margin: 0;

    padding: 0;

}


/* ============================================================
   TIME
   ============================================================ */

.time {

    color: #ffffff;

    font-family:
        "Arial Narrow",
        "Roboto Condensed",
        Arial,
        sans-serif;

    font-size: 22px;

    font-weight: 800;

    margin-top: 4px;

    line-height: 1;

    padding: 0;

}


/* ============================================================
   RIGHT HORIZONTAL LINE
   ============================================================ */

.right-line {

    width: 80px;

    height: 2px;

    background:
        rgba(255,255,255,0.75);

    margin-top: 7px;

    flex-shrink: 0;

}


/* ============================================================
   ORANGE BOTTOM BAR
   ============================================================ */

.orange-bar {

    position: absolute;

    left: 0;

    bottom: 0;

    width: 100%;

    height: 7px;

    background: #f28c00;

    z-index: 50;

}


</style>

</head>


<body>


<!-- ============================================================
     MAIN HEADER
     ============================================================ -->

<div class="psm-header">


    <!-- ========================================================
         LEFT LOGO AREA
         ======================================================== -->

    <div class="psm-left">


        <!-- ====================================================
             LOGO
             ==================================================== -->

        <div class="logo-panel">

            <img
                class="company-logo"
                src="data:image/jpeg;base64,LOGO_IMAGE_BASE64"
                alt="JSW JFE Steel Limited"
            >

        </div>


        <!-- ====================================================
             LEFT VERTICAL LINE
             ==================================================== -->

        <div class="vertical-line"></div>


        <!-- ====================================================
             CENTER TITLE GROUP
             ==================================================== -->

        <div class="title-area">


            <!-- MAIN TITLE -->

            <div class="main-title">

                INTERNAL AUDIT & COMPLIANCE

                <span class="main-title-orange"></span>

            </div>


            <!-- SUBTITLE -->

            <div class="subtitle">

                PSM DIGITAL DASHBOARD

            </div>


            <!-- SUB-SUBTITLE -->

            <div class="tagline">

                PEOPLE
                &nbsp; | &nbsp;
                PROCESS
                &nbsp; | &nbsp;
                RISK
                &nbsp; | &nbsp;
                COMPLIANCE

            </div>


        </div>


    </div>


    <!-- ========================================================
         RIGHT DATE / TIME
         ======================================================== -->

    <div class="psm-right">


        <div class="date">

            CURRENT_DATE_VALUE

        </div>


        <div class="time">

            CURRENT_TIME_VALUE

        </div>


        <div class="right-line"></div>


    </div>


    <!-- ========================================================
         ORANGE BOTTOM BAR
         ======================================================== -->

    <div class="orange-bar"></div>


</div>


</body>

</html>

"""

# ============================================================
# INSERT LOGO
# ============================================================

header_html = header_html.replace(
    "LOGO_IMAGE_BASE64",
    logo_base64
)

# ============================================================
# INSERT DATE
# ============================================================

header_html = header_html.replace(
    "CURRENT_DATE_VALUE",
    current_date
)

# ============================================================
# INSERT TIME
# ============================================================

header_html = header_html.replace(
    "CURRENT_TIME_VALUE",
    current_time
)

# ============================================================
# DISPLAY HEADER
# ============================================================

components.html(
    header_html,
    height=114,
    scrolling=False
)

import streamlit as st
import pandas as pd
import requests
import re
from io import BytesIO
from openpyxl import load_workbook


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="Audit Dashboard",
    layout="wide",
    initial_sidebar_state="collapsed"
)


# =========================================================
# GOOGLE SHEET CONFIGURATION
# =========================================================

SHEET_ID = "1--X0TT5Ts92EKAxrhV-fQgqeTHBX3rDVc1Egg74MewM"
SHEET_NAME = "Audit"

GOOGLE_XLSX_URL = (
    f"https://docs.google.com/spreadsheets/d/"
    f"{SHEET_ID}/export?format=xlsx"
)


# =========================================================
# SETTINGS
# =========================================================

ROWS_PER_PAGE = 5

AUTO_REFRESH_SECONDS = 60


# =========================================================
# SESSION STATE
# =========================================================

if "current_page" not in st.session_state:
    st.session_state.current_page = 1

if "applied_fy" not in st.session_state:
    st.session_state.applied_fy = "All"

if "applied_department" not in st.session_state:
    st.session_state.applied_department = "All"


# =========================================================
# AUTOMATIC REFRESH
# =========================================================

try:

    from streamlit_autorefresh import st_autorefresh

    st_autorefresh(
        interval=AUTO_REFRESH_SECONDS * 1000,
        key="audit_auto_refresh"
    )

except Exception:
    pass


# =========================================================
# CSS
# =========================================================

st.markdown(
    """
<style>

/* =====================================================
   GENERAL
   ===================================================== */

html,
body,
[data-testid="stAppViewContainer"],
[data-testid="stAppViewContainer"] > .main {
    overflow-x: hidden !important;
}

.block-container {
    padding-top: 0.75rem !important;
    padding-left: 0.8rem !important;
    padding-right: 0.8rem !important;
    padding-bottom: 0.8rem !important;
    max-width: 100% !important;
}

header {
    visibility: hidden !important;
    height: 0 !important;
}

[data-testid="stHeader"] {
    display: none !important;
}


/* =====================================================
   HORIZONTAL BLOCK GAP
   ===================================================== */

div[data-testid="stHorizontalBlock"] {
    gap: 0.45rem;
}


/* =====================================================
   TOP FILTER LABEL
   ===================================================== */

.filter-label {
    font-size: 13px;

    font-weight: 600;

    color: #17356D;

    height: 21px;

    line-height: 21px;

    margin-bottom: 8px;
}


/* =====================================================
   SELECT BOX
   ===================================================== */

div[data-baseweb="select"] {
    min-height: 42px !important;
    height: 42px !important;
}

div[data-baseweb="select"] > div {
    min-height: 42px !important;
    height: 42px !important;

    border-radius: 8px !important;

    font-size: 13px !important;
}


/* =====================================================
   SEARCH BUTTON
   ===================================================== */

.search-space {
    padding-top: 29px;
}

.search-space button {
    width: 95px !important;

    height: 42px !important;

    min-height: 42px !important;

    padding: 0 !important;

    margin: 0 !important;

    border-radius: 8px !important;

    font-size: 13px !important;

    font-weight: 600 !important;

    color: #17356D !important;
}


/* =====================================================
   REFRESH BUTTON
   ===================================================== */

.refresh-space {
    padding-top: 29px;
}

.refresh-space button {
    width: 42px !important;

    height: 42px !important;

    min-height: 42px !important;

    padding: 0 !important;

    margin: 0 !important;

    border-radius: 8px !important;

    background: #FFFFFF !important;

    border: 1px solid #D3D8DE !important;

    color: #17356D !important;

    font-size: 19px !important;

    line-height: 42px !important;
}


/* =====================================================
   KPI WRAPPER
   ===================================================== */

.kpi-wrapper {
    display: flex;

    width: 100%;

    gap: 14px;

    margin-top: 20px;

    margin-bottom: 20px;
}


/* =====================================================
   KPI CARD
   ===================================================== */

.kpi-card {
    position: relative;

    flex: 1;

    height: 92px;

    border-radius: 11px;

    padding: 12px 18px 12px 27px;

    box-sizing: border-box;

    display: flex;

    flex-direction: column;

    justify-content: center;

    overflow: hidden;
}


/* =====================================================
   KPI LEFT SIDE BAR
   ===================================================== */

.kpi-card::before {
    content: "";

    position: absolute;

    left: 0;

    top: 0;

    width: 8px;

    height: 100%;
}


/* =====================================================
   KPI 1 - TOTAL DEPARTMENTS
   ===================================================== */

.departments-kpi {
    background: #F1F3F7;

    border: 1px solid #DCE1E8;
}

.departments-kpi::before {
    background: #64748B;
}

.departments-kpi .kpi-title {
    color: #334155;
}

.departments-kpi .kpi-value {
    color: #334155;
}


/* =====================================================
   KPI 2 - AUDITS CONDUCTED
   ===================================================== */

.audit-kpi {
    background: #E7F1FC;

    border: 1px solid #D2E3F5;
}

.audit-kpi::before {
    background: #1479E8;
}

.audit-kpi .kpi-title {
    color: #17356D;
}

.audit-kpi .kpi-value {
    color: #17356D;
}


/* =====================================================
   KPI 3 - DEPARTMENT AUDIT COVERAGE
   ===================================================== */

.coverage-kpi {
    background: #FFF4E5;

    border: 1px solid #F5DEC0;
}

.coverage-kpi::before {
    background: #E88916;
}

.coverage-kpi .kpi-title {
    color: #A85C00;
}

.coverage-kpi .kpi-value {
    color: #A85C00;
}


/* =====================================================
   KPI 4 - AUDIT COMPLIANCE
   ===================================================== */

.compliance-kpi {
    background: #E6F5EA;

    border: 1px solid #C9E8D3;
}

.compliance-kpi::before {
    background: #20A35A;
}

.compliance-kpi .kpi-title {
    color: #168046;
}

.compliance-kpi .kpi-value {
    color: #168046;
}


/* =====================================================
   KPI TEXT
   ===================================================== */

.kpi-title {
    font-size: 14px;

    font-weight: 700;

    line-height: 18px;

    margin-bottom: 4px;
}

.kpi-value {
    font-size: 25px;

    font-weight: 700;

    line-height: 28px;
}


/* =====================================================
   TABLE WRAPPER
   ===================================================== */

.audit-table-wrapper {
    width: 100%;

    overflow: hidden;

    border: 1px solid #D6E1EC;

    border-radius: 10px;
}


/* =====================================================
   TABLE
   ===================================================== */

.audit-table {
    width: 100%;

    border-collapse: separate;

    border-spacing: 0;

    table-layout: fixed;

    font-family: Arial, sans-serif;
}


/* =====================================================
   TABLE HEADER
   ===================================================== */

.audit-table th {
    background: #EAF3FC;

    color: #17356D;

    font-size: 13px;

    font-weight: 700;

    text-align: center;

    padding: 10px 5px;

    height: 43px;

    border-right: 1px solid #D6E1EC;

    border-bottom: 1px solid #D6E1EC;
}


/* =====================================================
   TABLE BODY
   ===================================================== */

.audit-table td {
    color: #17356D;

    font-size: 12px;

    text-align: center;

    padding: 5px;

    height: 43px;

    border-right: 1px solid #D6E1EC;

    border-bottom: 1px solid #E1E8F0;

    overflow: hidden;

    text-overflow: ellipsis;

    white-space: nowrap;

    vertical-align: middle;
}


/* Alternate rows */

.audit-table tr:nth-child(even) td {
    background: #F5F9FD;
}

.audit-table tr:nth-child(odd) td {
    background: #FFFFFF;
}


/* Remove right border */

.audit-table th:last-child,
.audit-table td:last-child {
    border-right: none;
}


/* Remove bottom border */

.audit-table tr:last-child td {
    border-bottom: none;
}


/* =====================================================
   AUDIT REPORT LINK
   ===================================================== */

.audit-link {
    display: inline-block;

    width: 100px;

    padding: 6px 10px;

    border-radius: 9px;

    background: #DCEBFA;

    color: #2864A5 !important;

    text-decoration: none !important;

    font-size: 12px;

    font-weight: 500;

    text-align: center;
}


/* =====================================================
   COMPLIANCE REPORT LINK
   ===================================================== */

.compliance-link {
    display: inline-block;

    width: 100px;

    padding: 6px 10px;

    border-radius: 9px;

    background: #DDF3E4;

    color: #18864B !important;

    text-decoration: none !important;

    font-size: 12px;

    font-weight: 500;

    text-align: center;
}


/* =====================================================
   PAGINATION
   ===================================================== */

.pagination-area {
    width: 100%;

    margin-top: 12px;
}


/* =====================================================
   PAGINATION INFO
   ===================================================== */

.pagination-info {
    text-align: right;

    padding-top: 8px;

    font-size: 13px;

    color: #17356D;

    white-space: nowrap;
}

.pagination-records {
    color: #60708A;
}


/* =====================================================
   PAGINATION BUTTON
   ===================================================== */

.pagination-button {
    display: flex;

    justify-content: flex-end;
}

.pagination-button button {
    width: 36px !important;

    height: 36px !important;

    min-height: 36px !important;

    padding: 0 !important;

    margin: 0 !important;

    border-radius: 9px !important;

    background: #FFFFFF !important;

    border: 1px solid #D3D8DE !important;

    color: #AEB6BF !important;

    font-size: 18px !important;

    font-weight: 400 !important;

    line-height: 36px !important;
}

.pagination-button button:hover {
    background: #F7F9FB !important;

    border-color: #C5CCD4 !important;
}

.pagination-button button:disabled {
    background: #FFFFFF !important;

    color: #BFC5CB !important;

    border-color: #D7DBE0 !important;
}


/* =====================================================
   VERTICAL SPACING
   ===================================================== */

div[data-testid="stVerticalBlock"] {
    gap: 0.3rem;
}

</style>
""",
    unsafe_allow_html=True
)


# =========================================================
# DOWNLOAD GOOGLE SHEET
# =========================================================

def download_google_sheet():

    response = requests.get(
        GOOGLE_XLSX_URL,
        timeout=30
    )

    response.raise_for_status()

    return BytesIO(
        response.content
    )


# =========================================================
# CLEAN COLUMN NAME
# =========================================================

def clean_column_name(column):

    return (
        str(column)
        .strip()
        .replace("\n", " ")
        .replace("\r", " ")
    )


# =========================================================
# FIND COLUMN
# =========================================================

def find_column(
    df,
    possible_names
):

    normalized = {}

    for col in df.columns:

        key = (
            str(col)
            .strip()
            .lower()
            .replace("_", " ")
            .replace("-", " ")
        )

        key = re.sub(
            r"\s+",
            " ",
            key
        )

        normalized[key] = col


    # Exact match

    for name in possible_names:

        key = (
            str(name)
            .strip()
            .lower()
            .replace("_", " ")
            .replace("-", " ")
        )

        key = re.sub(
            r"\s+",
            " ",
            key
        )

        if key in normalized:

            return normalized[key]


    # Partial match

    for key, original_col in normalized.items():

        for name in possible_names:

            name_key = (
                str(name)
                .strip()
                .lower()
                .replace("_", " ")
                .replace("-", " ")
            )

            name_key = re.sub(
                r"\s+",
                " ",
                name_key
            )

            if (
                name_key in key
                or key in name_key
            ):

                return original_col


    return None


# =========================================================
# LOAD AUDIT DATA
# =========================================================

@st.cache_data(ttl=50)
def load_audit_data():

    excel_file = download_google_sheet()

    df = pd.read_excel(
        excel_file,
        sheet_name=SHEET_NAME
    )

    df.columns = [
        clean_column_name(c)
        for c in df.columns
    ]

    df = df.dropna(
        how="all"
    ).reset_index(drop=True)

    return df


# =========================================================
# EXTRACT URL FROM EXCEL CELL
# =========================================================

def extract_url(cell):

    if cell is None:

        return ""


    # -----------------------------------------------------
    # Normal hyperlink
    # -----------------------------------------------------

    try:

        if cell.hyperlink is not None:

            target = cell.hyperlink.target

            if target:

                return str(
                    target
                ).strip()

    except Exception:

        pass


    value = cell.value

    if value is None:

        return ""


    value = str(
        value
    ).strip()


    # -----------------------------------------------------
    # HYPERLINK formula
    # -----------------------------------------------------

    match = re.search(
        r'HYPERLINK\s*\(\s*"([^"]+)"',
        value,
        re.IGNORECASE
    )

    if match:

        return match.group(1).strip()


    # -----------------------------------------------------
    # Direct URL
    # -----------------------------------------------------

    if (
        value.startswith("http://")
        or value.startswith("https://")
    ):

        return value


    return ""


# =========================================================
# LOAD REPORT LINKS
# =========================================================

@st.cache_data(ttl=50)
def load_audit_links():

    excel_file = download_google_sheet()

    workbook = load_workbook(
        filename=excel_file,
        data_only=False
    )


    if SHEET_NAME not in workbook.sheetnames:

        return pd.DataFrame(
            columns=[
                "_AuditReportURL",
                "_ComplianceReportURL"
            ]
        )


    ws = workbook[SHEET_NAME]


    # -----------------------------------------------------
    # Find header row
    # -----------------------------------------------------

    header_row = None

    for row in ws.iter_rows(
        min_row=1,
        max_row=min(
            ws.max_row,
            10
        )
    ):

        values = [
            str(cell.value).strip().lower()
            if cell.value is not None
            else ""
            for cell in row
        ]

        joined = " ".join(
            values
        )


        if (
            "department" in joined
            or "audit date" in joined
            or "audit score" in joined
        ):

            header_row = row[0].row

            break


    if header_row is None:

        header_row = 1


    # -----------------------------------------------------
    # Find report columns
    # -----------------------------------------------------

    audit_report_col = None

    compliance_report_col = None


    for cell in ws[header_row]:

        if cell.value is None:

            continue


        header = clean_column_name(
            cell.value
        ).lower()


        header = (
            header
            .replace("_", " ")
            .replace("-", " ")
        )


        header = re.sub(
            r"\s+",
            " ",
            header
        ).strip()


        if (
            "audit report" in header
            and "compliance"
            not in header
        ):

            audit_report_col = cell.column


        if (
            "compliance report"
            in header
        ):

            compliance_report_col = cell.column


    # -----------------------------------------------------
    # Extract links
    # -----------------------------------------------------

    audit_urls = []

    compliance_urls = []


    for row_num in range(
        header_row + 1,
        ws.max_row + 1
    ):

        audit_url = ""

        compliance_url = ""


        if audit_report_col:

            audit_url = extract_url(
                ws.cell(
                    row=row_num,
                    column=audit_report_col
                )
            )


        if compliance_report_col:

            compliance_url = extract_url(
                ws.cell(
                    row=row_num,
                    column=compliance_report_col
                )
            )


        audit_urls.append(
            audit_url
        )

        compliance_urls.append(
            compliance_url
        )


    return pd.DataFrame(
        {
            "_AuditReportURL":
                audit_urls,

            "_ComplianceReportURL":
                compliance_urls
        }
    )


# =========================================================
# LOAD DATA
# =========================================================

try:

    df = load_audit_data()

    link_df = load_audit_links()

except Exception as e:

    st.error(
        f"Unable to load Audit data: {e}"
    )

    st.stop()


# =========================================================
# MATCH LINKS WITH DATA
# =========================================================

if len(link_df) < len(df):

    link_df = link_df.reindex(
        range(len(df))
    )

elif len(link_df) > len(df):

    link_df = link_df.iloc[
        :len(df)
    ].copy()


df["_AuditReportURL"] = (
    link_df["_AuditReportURL"]
    .fillna("")
    .astype(str)
    .values
)


df["_ComplianceReportURL"] = (
    link_df["_ComplianceReportURL"]
    .fillna("")
    .astype(str)
    .values
)


# =========================================================
# FIND COLUMNS
# =========================================================

department_col = find_column(
    df,
    [
        "Department",
        "Dept",
        "Department Name"
    ]
)


audit_date_col = find_column(
    df,
    [
        "Audit Date",
        "Date of Audit",
        "Audit Dt",
        "Date"
    ]
)


audit_score_col = find_column(
    df,
    [
        "Audit Score",
        "Score",
        "Audit Rating"
    ]
)


audit_report_col = find_column(
    df,
    [
        "Audit Report",
        "Audit Reports"
    ]
)


compliance_report_col = find_column(
    df,
    [
        "Compliance Report",
        "Compliance Reports"
    ]
)


# =========================================================
# VALIDATE DEPARTMENT
# =========================================================

if department_col is None:

    st.error(
        "Department column was not found."
    )

    st.write(
        "Available columns:",
        list(df.columns)
    )

    st.stop()


# =========================================================
# CLEAN DEPARTMENT
# =========================================================

df[department_col] = (
    df[department_col]
    .fillna("")
    .astype(str)
    .str.strip()
)


# =========================================================
# PARSE AUDIT DATE
# =========================================================

if audit_date_col is not None:

    df["_AuditDateParsed"] = pd.to_datetime(
        df[audit_date_col],
        errors="coerce",
        dayfirst=True
    )

else:

    df["_AuditDateParsed"] = pd.NaT


# =========================================================
# FINANCIAL YEAR
# =========================================================

def get_financial_year(
    date_value
):

    if pd.isna(date_value):

        return None


    year = date_value.year

    month = date_value.month


    if month >= 4:

        start_year = year

    else:

        start_year = year - 1


    end_year = start_year + 1


    return (
        f"FY {start_year}-{str(end_year)[-2:]}"
    )


df["_FinancialYear"] = (
    df["_AuditDateParsed"]
    .apply(
        get_financial_year
    )
)


# =========================================================
# FINANCIAL YEAR OPTIONS
# =========================================================

available_fy = sorted(
    [
        fy
        for fy in df["_FinancialYear"]
        .dropna()
        .unique()
    ],
    reverse=True
)


financial_year_options = [
    "All"
] + available_fy


# =========================================================
# DEPARTMENT OPTIONS
# =========================================================

available_departments = sorted(
    [
        d
        for d in df[department_col].unique()
        if str(d).strip() != ""
    ]
)


department_options = [
    "All"
] + available_departments


# =========================================================
# TOP FILTER ROW
# =========================================================
#
# Search and refresh are immediately adjacent.
#
# =========================================================

year_col, dept_col, search_col, refresh_col = st.columns(
    [
        0.31,
        0.31,
        0.10,
        0.045
    ],
    gap="small"
)


# =========================================================
# YEAR FILTER
# =========================================================

with year_col:

    st.markdown(
        '<div class="filter-label">'
        'Select Year'
        '</div>',
        unsafe_allow_html=True
    )


    selected_fy = st.selectbox(
        "Select Year",

        options=financial_year_options,

        index=(
            financial_year_options.index(
                st.session_state.applied_fy
            )
            if (
                st.session_state.applied_fy
                in financial_year_options
            )
            else 0
        ),

        label_visibility="collapsed",

        key="financial_year_selector"
    )


# =========================================================
# DEPARTMENT FILTER
# =========================================================

with dept_col:

    st.markdown(
        '<div class="filter-label">'
        'Department'
        '</div>',
        unsafe_allow_html=True
    )


    selected_department = st.selectbox(
        "Department",

        options=department_options,

        index=(
            department_options.index(
                st.session_state.applied_department
            )
            if (
                st.session_state.applied_department
                in department_options
            )
            else 0
        ),

        label_visibility="collapsed",

        key="department_selector"
    )


# =========================================================
# SEARCH BUTTON
# =========================================================

with search_col:

    st.markdown(
        '<div class="search-space">',
        unsafe_allow_html=True
    )


    search_clicked = st.button(
        "Search",
        key="search_button"
    )


    st.markdown(
        '</div>',
        unsafe_allow_html=True
    )


# =========================================================
# REFRESH BUTTON
# =========================================================

with refresh_col:

    st.markdown(
        '<div class="refresh-space">',
        unsafe_allow_html=True
    )


    refresh_clicked = st.button(
        "↻",
        key="refresh_button",
        help="Refresh data"
    )


    st.markdown(
        '</div>',
        unsafe_allow_html=True
    )


# =========================================================
# MANUAL REFRESH
# =========================================================

if refresh_clicked:

    load_audit_data.clear()

    load_audit_links.clear()

    st.session_state.current_page = 1

    st.rerun()


# =========================================================
# SEARCH ACTION
# =========================================================

if search_clicked:

    st.session_state.applied_fy = (
        selected_fy
    )

    st.session_state.applied_department = (
        selected_department
    )

    st.session_state.current_page = 1

    st.rerun()


# =========================================================
# FILTER DATA
# =========================================================

filtered_df = df.copy()


# =========================================================
# FINANCIAL YEAR FILTER
# =========================================================

if (
    st.session_state.applied_fy
    != "All"
):

    filtered_df = filtered_df[
        filtered_df["_FinancialYear"]
        == st.session_state.applied_fy
    ].copy()


# =========================================================
# DEPARTMENT FILTER
# =========================================================

if (
    st.session_state.applied_department
    != "All"
):

    filtered_df = filtered_df[
        filtered_df[department_col]
        == st.session_state.applied_department
    ].copy()


# =========================================================
# AUDIT REPORT PRESENCE
# =========================================================
#
# IMPORTANT:
#
# We determine whether an audit was conducted from the
# CONTENT of the Audit Report column.
#
# This is deliberately NOT dependent on whether
# openpyxl successfully extracted the hyperlink URL.
#
# Therefore a populated Google Sheets hyperlink cell
# counts as an audit even if URL metadata isn't detected.
#
# =========================================================

if audit_report_col is not None:

    audit_report_present_mask = (
        filtered_df[audit_report_col]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
    )

else:

    audit_report_present_mask = (
        filtered_df["_AuditReportURL"]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
    )


# =========================================================
# AUDITED DEPARTMENTS
# =========================================================

audited_department_series = (
    filtered_df.loc[
        audit_report_present_mask,
        department_col
    ]
    .dropna()
    .astype(str)
    .str.strip()
)


audited_department_series = (
    audited_department_series[
        audited_department_series != ""
    ]
)


total_audited_departments = (
    audited_department_series.nunique()
)


# =========================================================
# TOTAL DEPARTMENTS
# =========================================================

if (
    st.session_state.applied_fy
    == "All"
):

    total_department_series = (
        df[department_col]
        .dropna()
        .astype(str)
        .str.strip()
    )

else:

    total_department_series = (
        filtered_df[department_col]
        .dropna()
        .astype(str)
        .str.strip()
    )


total_department_series = (
    total_department_series[
        total_department_series != ""
    ]
)


total_departments = (
    total_department_series.nunique()
)


# =========================================================
# DEPARTMENT AUDIT COVERAGE
# =========================================================
#
# Audited Departments
# ------------------- × 100
# Total Departments
#
# =========================================================

if total_departments > 0:

    audit_coverage_percentage = (
        total_audited_departments
        / total_departments
    ) * 100

else:

    audit_coverage_percentage = 0


# =========================================================
# COMPLIANCE REPORT PRESENCE
# =========================================================
#
# Again, use the actual CONTENT of the sheet cell,
# rather than depending only on hyperlink extraction.
#
# =========================================================

if compliance_report_col is not None:

    compliance_report_present_mask = (
        filtered_df[compliance_report_col]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
    )

else:

    compliance_report_present_mask = (
        filtered_df["_ComplianceReportURL"]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
    )


# =========================================================
# COMPLIANCE DEPARTMENTS
# =========================================================

compliance_department_series = (
    filtered_df.loc[
        compliance_report_present_mask,
        department_col
    ]
    .dropna()
    .astype(str)
    .str.strip()
)


compliance_department_series = (
    compliance_department_series[
        compliance_department_series != ""
    ]
)


total_compliance_departments = (
    compliance_department_series.nunique()
)


# =========================================================
# AUDIT COMPLIANCE
# =========================================================
#
# Compliance Reports
# ------------------ × 100
# Audited Departments
#
# Example:
#
# 1 / 8 × 100 = 12.5%
#
# =========================================================

if total_audited_departments > 0:

    compliance_percentage = (
        total_compliance_departments
        / total_audited_departments
    ) * 100

else:

    compliance_percentage = 0


# =========================================================
# LIMIT PERCENTAGES
# =========================================================

audit_coverage_percentage = max(
    0,
    min(
        100,
        audit_coverage_percentage
    )
)


compliance_percentage = max(
    0,
    min(
        100,
        compliance_percentage
    )
)


# =========================================================
# KPI CARDS
# =========================================================

kpi_html = f"""
<div class="kpi-wrapper">


<!-- ===================================================
     1. TOTAL DEPARTMENTS
     =================================================== -->

<div class="kpi-card departments-kpi">

    <div class="kpi-title">
        Total Departments
    </div>

    <div class="kpi-value">
        {total_departments}
    </div>

</div>


<!-- ===================================================
     2. AUDITS CONDUCTED
     =================================================== -->

<div class="kpi-card audit-kpi">

    <div class="kpi-title">
        Audits Conducted
    </div>

    <div class="kpi-value">
        {total_audited_departments}
    </div>

</div>


<!-- ===================================================
     3. DEPARTMENT AUDIT COVERAGE
     =================================================== -->

<div class="kpi-card coverage-kpi">

    <div class="kpi-title">
        Department Audit Coverage
    </div>

    <div class="kpi-value">
        {audit_coverage_percentage:.1f}%
    </div>

</div>


<!-- ===================================================
     4. AUDIT COMPLIANCE
     =================================================== -->

<div class="kpi-card compliance-kpi">

    <div class="kpi-title">
        Audit Compliance
    </div>

    <div class="kpi-value">
        {compliance_percentage:.1f}%
    </div>

</div>


</div>
"""


# =========================================================
# DISPLAY KPI
# =========================================================

try:

    st.html(
        kpi_html
    )

except AttributeError:

    st.markdown(
        kpi_html,
        unsafe_allow_html=True
    )


# =========================================================
# TABLE RECORD COUNT
# =========================================================

total_records = len(
    filtered_df
)


# =========================================================
# PAGINATION CALCULATION
# =========================================================

total_pages = max(
    1,
    (
        total_records
        + ROWS_PER_PAGE
        - 1
    )
    // ROWS_PER_PAGE
)


if (
    st.session_state.current_page
    > total_pages
):

    st.session_state.current_page = (
        total_pages
    )


if (
    st.session_state.current_page
    < 1
):

    st.session_state.current_page = 1


current_page = (
    st.session_state.current_page
)


# =========================================================
# CURRENT PAGE DATA
# =========================================================

start_idx = (
    current_page
    - 1
) * ROWS_PER_PAGE


end_idx = (
    start_idx
    + ROWS_PER_PAGE
)


page_df = filtered_df.iloc[
    start_idx:end_idx
].copy()


# =========================================================
# BUILD TABLE ROWS
# =========================================================

table_rows = ""


for _, row in page_df.iterrows():

    # -----------------------------------------------------
    # DEPARTMENT
    # -----------------------------------------------------

    department = str(
        row.get(
            department_col,
            ""
        )
    ).strip()


    if (
        not department
        or department.lower() == "nan"
    ):

        department = "-"


    # -----------------------------------------------------
    # AUDIT DATE
    # -----------------------------------------------------

    if audit_date_col is not None:

        audit_date = row.get(
            audit_date_col,
            ""
        )

    else:

        audit_date = ""


    if pd.isna(audit_date):

        audit_date_display = "-"

    else:

        parsed_date = pd.to_datetime(
            audit_date,
            errors="coerce",
            dayfirst=True
        )


        if pd.isna(parsed_date):

            audit_date_text = str(
                audit_date
            ).strip()


            if (
                not audit_date_text
                or audit_date_text.lower()
                == "nan"
            ):

                audit_date_display = "-"

            else:

                audit_date_display = (
                    audit_date_text
                )

        else:

            audit_date_display = (
                parsed_date.strftime(
                    "%d-%m-%Y"
                )
            )


    # -----------------------------------------------------
    # AUDIT SCORE
    # -----------------------------------------------------

    if audit_score_col is not None:

        audit_score = row.get(
            audit_score_col,
            ""
        )


        if pd.isna(audit_score):

            audit_score_display = "-"

        else:

            audit_score_display = str(
                audit_score
            ).strip()


            if (
                not audit_score_display
                or audit_score_display.lower()
                == "nan"
            ):

                audit_score_display = "-"

    else:

        audit_score_display = "-"


    # -----------------------------------------------------
    # AUDIT REPORT LINK
    # -----------------------------------------------------

    audit_url = str(
        row.get(
            "_AuditReportURL",
            ""
        )
    ).strip()


    if audit_url:

        audit_report_html = (
            f'<a class="audit-link" '
            f'href="{audit_url}" '
            f'target="_blank">'
            f'View'
            f'</a>'
        )

    else:

        audit_report_html = "-"


    # -----------------------------------------------------
    # COMPLIANCE REPORT LINK
    # -----------------------------------------------------

    compliance_url = str(
        row.get(
            "_ComplianceReportURL",
            ""
        )
    ).strip()


    if compliance_url:

        compliance_report_html = (
            f'<a class="compliance-link" '
            f'href="{compliance_url}" '
            f'target="_blank">'
            f'View'
            f'</a>'
        )

    else:

        compliance_report_html = "-"


    # -----------------------------------------------------
    # TABLE ROW
    # -----------------------------------------------------

    table_rows += f"""
<tr>

<td>
{department}
</td>

<td>
{audit_date_display}
</td>

<td>
{audit_score_display}
</td>

<td>
{audit_report_html}
</td>

<td>
{compliance_report_html}
</td>

</tr>
"""


# =========================================================
# TABLE HTML
# =========================================================

table_html = f"""
<div class="audit-table-wrapper">

<table class="audit-table">

<colgroup>

<col style="width:22%;">

<col style="width:15%;">

<col style="width:14%;">

<col style="width:24%;">

<col style="width:25%;">

</colgroup>


<thead>

<tr>

<th>
Department
</th>

<th>
Audit Date
</th>

<th>
Audit Score
</th>

<th>
Audit Report
</th>

<th>
Compliance Report
</th>

</tr>

</thead>


<tbody>

{table_rows}

</tbody>

</table>

</div>
"""


# =========================================================
# DISPLAY TABLE
# =========================================================

try:

    st.html(
        table_html
    )

except AttributeError:

    st.markdown(
        table_html,
        unsafe_allow_html=True
    )


# =========================================================
# PAGINATION
# =========================================================

st.markdown(
    '<div class="pagination-area"></div>',
    unsafe_allow_html=True
)


# =========================================================
# PAGINATION RIGHT SIDE
# =========================================================

blank_col, info_col, prev_col, next_col = st.columns(
    [
        7.3,
        1.8,
        0.40,
        0.40
    ],
    gap="small"
)


# =========================================================
# PAGE INFORMATION
# =========================================================

with info_col:

    page_info_html = f"""
<div class="pagination-info">

<b>
Page {current_page} of {total_pages}
</b>

<span class="pagination-records">

&nbsp;&nbsp;·&nbsp;&nbsp;

{total_records} records

</span>

</div>
"""


    try:

        st.html(
            page_info_html
        )

    except AttributeError:

        st.markdown(
            page_info_html,
            unsafe_allow_html=True
        )


# =========================================================
# PREVIOUS
# =========================================================

with prev_col:

    st.markdown(
        '<div class="pagination-button">',
        unsafe_allow_html=True
    )


    previous_clicked = st.button(
        "‹",

        key="previous_page",

        disabled=(
            current_page <= 1
        )
    )


    st.markdown(
        '</div>',
        unsafe_allow_html=True
    )


# =========================================================
# NEXT
# =========================================================

with next_col:

    st.markdown(
        '<div class="pagination-button">',
        unsafe_allow_html=True
    )


    next_clicked = st.button(
        "›",

        key="next_page",

        disabled=(
            current_page >= total_pages
        )
    )


    st.markdown(
        '</div>',
        unsafe_allow_html=True
    )


# =========================================================
# PREVIOUS ACTION
# =========================================================

if previous_clicked:

    st.session_state.current_page = (
        current_page - 1
    )

    st.rerun()


# =========================================================
# NEXT ACTION
# =========================================================

if next_clicked:

    st.session_state.current_page = (
        current_page + 1
    )

    st.rerun()

