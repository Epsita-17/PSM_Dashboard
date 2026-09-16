
import streamlit as st
import pandas as pd
import numpy as np
import re
import requests
from html.parser import HTMLParser
from html import unescape
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from urllib.parse import quote
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
import base64
import streamlit.components.v1 as components

# ============================================================
# PSM DOCUMENT REPOSITORY DASHBOARD
# Google Sheets:
#   - MOM
#   - TRAINING PPTS
#   - PSM Procedure
#
# NOTE:
# The Google Sheet must be accessible to the Streamlit app
# (for example: "Anyone with the link -> Viewer").
# ============================================================

st.set_page_config(
    page_title="PSM Document Repository",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ============================================================
# BASE DIRECTORY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent


# ============================================================
# IMAGE TO BASE64
# ============================================================

def image_to_base64(file_path):

    file_path = Path(file_path)

    if not file_path.exists():
        return ""

    try:

        with open(file_path, "rb") as file:

            return base64.b64encode(
                file.read()
            ).decode("utf-8")

    except Exception:

        return ""


    # ============================================================
# COMPANY LOGO
# ============================================================

logo_path = BASE_DIR / "jsw_jfe_logo.jpg"

logo_base64 = image_to_base64(logo_path)


# ============================================================
# FILE CHECK
# ============================================================

if not logo_base64:

    st.error(
        "jsw_jfe_logo.jpg not found. "
        "Keep jsw_jfe_logo.jpg in the same folder as this Python file."
    )


# ============================================================
# DATE AND TIME
# ============================================================

now = datetime.now(ZoneInfo("Asia/Kolkata"))

current_date = now.strftime(
    "%d %b %Y"
).upper()

current_time = now.strftime(
    "%I:%M %p"
)


# ============================================================
# HEADER HTML
# ============================================================

header_html = """ 

<!DOCTYPE html> 

<html> 

<head> 

<meta charset="UTF-8"> 

<style> 


/* ============================================================ 
   MAIN HEADER 
   ============================================================ */ 

.psm-header { 

    position: relative; 

    width: 100%; 

    height: 90px; 

    overflow: hidden; 

    background: 
        linear-gradient( 
            90deg, 
            #031d34 0%, 
            #052b49 42%, 
            #07385c 74%, 
            #052b49 100% 
        ); 

    border-radius: 7px; 

    box-shadow: 
        0 3px 9px 
        rgba(0,0,0,0.18); 

} 


/* ============================================================ 
   LEFT LOGO AREA 
   ============================================================ */ 

.psm-left { 

    position: absolute; 

    left: 0; 

    top: 0; 

    width: 100%; 

    height: 95px; 

    display: flex; 

    align-items: center; 

    padding-left: 4px; 

    box-sizing: border-box; 

    z-index: 20; 

    pointer-events: none; 

} 


/* ============================================================ 
   LOGO PANEL 
   ONLY VERTICAL POSITION CHANGED 
   ============================================================ */ 

.logo-panel { 

    width: 195px; 

    height: 68px; 

    background: #ffffff; 

    border-radius: 5px; 

    padding: 4px; 

    display: flex; 

    align-items: center; 

    justify-content: center; 

    flex-shrink: 0; 

    box-sizing: border-box; 

    box-shadow: 
        0 3px 9px 
        rgba(0,0,0,0.20); 

    position: relative; 

    top: -6px; 

    left: 6px; 

} 


/* ============================================================ 
   COMPANY LOGO 
   ============================================================ */ 

.company-logo { 

    width: 100%; 

    height: 100%; 

    object-fit: contain; 

    object-position: center; 

    display: block; 

} 


/* ============================================================ 
   LEFT VERTICAL DIVIDER 
   ============================================================ */ 

.vertical-line { 

    width: 2px; 

    height: 83px; 

    background: 
        rgba(255,255,255,0.65); 

    margin-left: 18px; 

    margin-right: 20px; 

    flex-shrink: 0; 

} 


/* ============================================================ 
   CENTER TITLE AREA 
   ============================================================ */ 

.title-area { 

    position: absolute; 

    left: 50%; 

    top: 0; 

    height: 95px; 

    display: flex; 

    flex-direction: column; 

    justify-content: center; 

    align-items: center; 

    text-align: center; 

    min-width: max-content; 

    box-sizing: border-box; 

    transform: translateX(-50%); 

} 


/* ============================================================ 
   MAIN TITLE 
   ============================================================ */ 

.main-title { 

    color: #ffffff; 

    font-family: 
        "Arial Narrow", 
        "Roboto Condensed", 
        Arial, 
        sans-serif; 

    font-size: 27px; 

    font-weight: 900; 

    line-height: 1; 

    letter-spacing: 0.3px; 

    white-space: nowrap; 

    margin: 0; 

    padding: 0; 

} 


/* ============================================================ 
   ORANGE TITLE PART 
   ============================================================ */ 

.main-title-orange { 

    color: #f28c00; 

} 


/* ============================================================ 
   SUBTITLE 
   ============================================================ */ 

.subtitle { 

    color: #ffffff; 

    font-family: 
        Arial, 
        sans-serif; 

    font-size: 12px; 

    font-weight: 400; 

    letter-spacing: 3.6px; 

    margin-top: 8px; 

    line-height: 1; 

    white-space: nowrap; 

} 

/* ============================================================ 
   SUB-SUBTITLE / TAGLINE 
   ============================================================ */ 

.tagline { 

    color: 
        rgba(255,255,255,0.82); 

    font-family: 
        Arial, 
        sans-serif; 

    font-size: 7px; 

    font-weight: 500; 

    letter-spacing: 2.2px; 

    margin-top: 6px; 

    line-height: 1; 

    white-space: nowrap; 

} 


/* ============================================================ 
   RIGHT DATE / TIME AREA 
   ============================================================ */ 

.psm-right { 

    position: absolute; 

    right: 16px; 

    top: 0; 

    width: 15%; 

    height: 95px; 

    display: flex; 

    flex-direction: column; 

    justify-content: center; 

    align-items: flex-end; 

    text-align: right; 

    color: #ffffff; 

    z-index: 30; 

    padding-left: 18px; 

    box-sizing: border-box; 

} 


/* ============================================================ 
   RIGHT VERTICAL DIVIDER 
   ============================================================ */ 

.psm-right::before { 

    content: ""; 

    position: absolute; 

    left: 0; 

    top: 6px; 

    width: 2px; 

    height: 83px; 

    background: 
        rgba(255,255,255,0.65); 

} 


/* ============================================================ 
   DATE 
   ============================================================ */ 

.date { 

    color: 
        rgba(255,255,255,0.95); 

    font-family: 
        Arial, 
        sans-serif; 

    font-size: 11px; 

    font-weight: 400; 

    letter-spacing: 0.7px; 

    line-height: 1; 

    margin: 0; 

    padding: 0; 

} 


/* ============================================================ 
   TIME 
   ============================================================ */ 

.time { 

    color: #ffffff; 

    font-family: 
        "Arial Narrow", 
        "Roboto Condensed", 
        Arial, 
        sans-serif; 

    font-size: 22px; 

    font-weight: 800; 

    margin-top: 4px; 

    line-height: 1; 

    padding: 0; 

} 


/* ============================================================ 
   RIGHT HORIZONTAL LINE 
   ============================================================ */ 

.right-line { 

    width: 80px; 

    height: 2px; 

    background: 
        rgba(255,255,255,0.75); 

    margin-top: 7px; 

    flex-shrink: 0; 

} 


/* ============================================================ 
   ORANGE BOTTOM BAR 
   ============================================================ */ 

.orange-bar { 

    position: absolute; 

    left: 0; 

    bottom: 0; 

    width: 100%; 

    height: 9px; 

    background: #f28c00; 

    z-index: 50; 

} 


</style> 

</head> 


<body> 


<!-- ============================================================ 
     MAIN HEADER 
     ============================================================ --> 

<div class="psm-header"> 


    <!-- ======================================================== 
         LEFT LOGO AREA 
         ======================================================== --> 

    <div class="psm-left"> 


        <!-- ==================================================== 
             LOGO 
             ==================================================== --> 

        <div class="logo-panel"> 

            <img 
                class="company-logo" 
                src="data:image/jpeg;base64,LOGO_IMAGE_BASE64" 
                alt="JSW JFE Steel Limited" 
            > 

        </div> 


        <!-- ==================================================== 
             LEFT VERTICAL LINE 
             ==================================================== --> 

        <div class="vertical-line"></div> 


        <!-- ==================================================== 
             CENTER TITLE GROUP 
             ==================================================== --> 

        <div class="title-area"> 


            <!-- MAIN TITLE --> 

            <div class="main-title"> 

                PROCESS HAZARD ANALYSIS (PHA) 

                <span class="main-title-orange"></span> 

            </div> 


            <!-- SUBTITLE --> 

            <div class="subtitle"> 

                PSM DIGITAL DASHBOARD 

            </div> 


            <!-- SUB-SUBTITLE --> 

            <div class="tagline"> 

                PEOPLE 
                &nbsp; | &nbsp; 
                PROCESS 
                &nbsp; | &nbsp; 
                RISK 
                &nbsp; | &nbsp; 
                COMPLIANCE 

            </div> 


        </div> 


    </div> 


    <!-- ======================================================== 
         RIGHT DATE / TIME 
         ======================================================== --> 

    <div class="psm-right"> 


        <div class="date"> 

            CURRENT_DATE_VALUE 

        </div> 


        <div class="time"> 

            CURRENT_TIME_VALUE 

        </div> 


        <div class="right-line"></div> 


    </div> 


    <!-- ======================================================== 
         ORANGE BOTTOM BAR 
         ======================================================== --> 

    <div class="orange-bar"></div> 


</div> 


</body> 

</html> 

"""


# ============================================================
# INSERT LOGO
# ============================================================

header_html = header_html.replace(
    "LOGO_IMAGE_BASE64",
    logo_base64
)


# ============================================================
# INSERT DATE
# ============================================================

header_html = header_html.replace(
    "CURRENT_DATE_VALUE",
    current_date
)


# ============================================================
# INSERT TIME
# ============================================================

header_html = header_html.replace(
    "CURRENT_TIME_VALUE",
    current_time
)


# ============================================================
# DISPLAY HEADER
# ============================================================

components.html(
    header_html,
    height=114,
    scrolling=False
)



# -----------------------------
# Google Sheet configuration
# -----------------------------
SHEET_ID = "1--X0TT5Ts92EKAxrhV-fQgqeTHBX3rDVc1Egg74MewM"

SHEET_NAMES = {
    "mom": "MOM",
    "training": "TRAINING PPTS",
    "procedure": "PSM Procedure",
}

# -----------------------------
# Global CSS
# -----------------------------
st.markdown(
    """
    <style>
    /* Remove Streamlit's default top/header area */
    [data-testid="stHeader"] {
        display: none;
    }

    .stApp {
        background: #ffffff;
    }

    /* Compact dashboard: designed to fit the complete repository on one
       normal desktop viewport with minimal/no vertical scrolling. */
    .main .block-container {
        padding-top: 0.10rem;
        padding-bottom: 0.15rem;
        padding-left: 0.55rem;
        padding-right: 0.55rem;
        max-width: none;
        width: 100%;
    }

    /* Use the screen width while keeping the dashboard compact vertically. */
    .main {
        zoom: 0.82;
    }

    /* Hide sidebar completely */
    [data-testid="stSidebar"] {
        display: none;
    }

    /* ---------- Horizontal year selector ---------- */
    .year-inline-label {
        font-size: 15px;
        font-weight: 700;
        color: #0c2b72;
        white-space: nowrap;
        height: 36px;
        display: flex;
        align-items: center;
        justify-content: flex-start;
        margin: 0;
    }

    /* Keep year dropdown and refresh icon on exactly the same baseline. */
    div[data-testid="stSelectbox"] {
        margin: 0 !important;
    }

    div[data-testid="stSelectbox"] > div {
        margin: 0 !important;
    }

    /* ---------- KPI cards ---------- */
    .kpi-card {
        min-height: 78px;
        border-radius: 8px;
        padding: 8px 18px 7px 24px;
        position: relative;
        overflow: hidden;
        box-shadow: 0 1px 5px rgba(0,0,0,0.05);
        border: 1px solid rgba(30,80,140,0.08);
    }

    .kpi-card:before {
        content: "";
        position: absolute;
        left: 0;
        top: 0;
        bottom: 0;
        width: 7px;
        border-radius: 8px 0 0 8px;
    }

    .kpi-red {
        background: linear-gradient(90deg, #fff1f3 0%, #fde8eb 100%);
    }

    .kpi-red:before {
        background: #e3182d;
    }

    .kpi-training {
        background: linear-gradient(90deg, #fffaf0 0%, #fff7e6 100%);
    }

    .kpi-training:before {
        background: #f28c18;
    }

    .kpi-training .kpi-title {
        color: #c66a00;
    }

    .kpi-training .kpi-number {
        color: #c66a00;
    }

    .kpi-training .kpi-subtitle {
        color: #9a5a00;
    }

    .kpi-number.training-number {
        color: #c66a00 !important;
    }

    /* Compact refresh control: small icon button aligned with year selector */
    div[data-testid="stButton"] > button {
        width: 38px !important;
        min-width: 38px !important;
        max-width: 38px !important;
        height: 36px !important;
        padding: 0 !important;
        margin-top: 0 !important;
        border: 1px solid #8da4c4 !important;
        border-radius: 6px !important;
        background: #ffffff !important;
        color: #0c438e !important;
        font-size: 18px !important;
        line-height: 1 !important;
    }

    div[data-testid="stButton"] > button:hover {
        background: #eef5fc !important;
        border-color: #0c6fd8 !important;
        color: #0759ad !important;
    }

    .kpi-blue {
        background: linear-gradient(90deg, #edf6ff 0%, #e3f0fc 100%);
    }

    .kpi-blue:before {
        background: #0877df;
    }

    .kpi-title {
        font-size: 16px;
        font-weight: 700;
        color: #09266d;
        line-height: 1.15;
    }

    .kpi-number {
        font-size: 32px;
        font-weight: 800;
        line-height: 1.0;
        margin-top: 3px;
    }

    .kpi-number.red {
        color: #d7192d;
    }

    .kpi-number.blue {
        color: #06489f;
    }

    .kpi-subtitle {
        color: #455a85;
        font-size: 13px;
        margin-top: 1px;
    }

    /* ---------- Section ---------- */
    .section-title {
        font-size: 21px;
        font-weight: 800;
        line-height: 1.1;
        margin: 0;
        padding: 0 0 5px 1px;
        position: relative;
    }

    .section-title:after {
        content: "";
        position: absolute;
        left: 1px;
        bottom: 0;
        width: 55px;
        height: 3px;
        border-radius: 4px;
    }

    .procedure-title {
        color: #e21c2b;
    }

    .procedure-title:after {
        background: #e21c2b;
    }

    .meeting-title {
        color: #0036b3;
    }

    .meeting-title:after {
        background: #0036b3;
    }

    .training-title {
        color: #e21c2b;
    }

    .training-title:after {
        background: #e21c2b;
    }

    /* ---------- Search boxes ---------- */
    div[data-testid="stTextInput"] {
        margin-top: -4px;
    }

    div[data-testid="stTextInput"] input {
        height: 32px;
        border: 1px solid #8da0bd;
        border-radius: 7px;
        color: #18356e;
        font-size: 13px;
        padding-left: 12px;
    }

    div[data-testid="stTextInput"] input:focus {
        border-color: #147dd1;
        box-shadow: 0 0 0 1px #147dd1;
    }

    /* ---------- Table ---------- */
    .table-wrap {
        border: 1px solid #dce5ee;
        border-radius: 3px;
        overflow: hidden;
        box-shadow: 0 1px 3px rgba(0,0,0,0.04);
        background: #ffffff;
        padding: 0 !important;
        margin: 0 !important;
        line-height: 0;
    }

    .repo-table {
        width: 100%;
        height: auto !important;
        min-height: 0 !important;
        margin: 0 !important;
        padding: 0 !important;
        border-collapse: collapse;
        border-spacing: 0 !important;
        table-layout: fixed;
        font-size: 14px;
        line-height: normal;
    }

    .repo-table th {
        background: #e5f2fc;
        color: #08276d;
        font-weight: 800;
        text-align: left;
        height: 34px !important;
        padding: 5px 10px !important;
        box-sizing: border-box;
        border-right: 1px solid #d4e1ed;
        border-bottom: 1px solid #d4e1ed;
        white-space: nowrap;
        line-height: 20px !important;
    }

    .repo-table td {
        color: #153d80;
        height: 28px !important;
        max-height: 28px !important;
        padding: 0 10px !important;
        margin: 0 !important;
        line-height: 20px !important;
        box-sizing: border-box;
        border-right: 1px solid #e1e8ef;
        border-bottom: 1px solid #e1e8ef;
        vertical-align: middle;
        background: #ffffff;
        overflow: hidden;
    }

    .repo-table tr {
        height: 28px !important;
        max-height: 28px !important;
    }

    .repo-table tbody {
        height: auto !important;
        min-height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }

    .repo-table tbody tr:last-child,
    .repo-table tbody tr:last-child td {
        height: 28px !important;
        max-height: 28px !important;
        padding-top: 0 !important;
        padding-bottom: 0 !important;
        margin-bottom: 0 !important;
        border-bottom: none;
    }

    .repo-table td > * {
        margin-top: 0 !important;
        margin-bottom: 0 !important;
    }

    .repo-table th:last-child,
    .repo-table td:last-child {
        border-right: none;
        text-align: center;
    }

    .repo-table tr:hover td {
        background: #f8fbfe;
    }

    .view-btn {
        display: inline-flex;
        align-items: center;
        justify-content: center;
        min-width: 78px;
        height: 23px;
        padding: 0 16px;
        background: #0d7fdc;
        color: #ffffff !important;
        text-decoration: none !important;
        border-radius: 4px;
        font-size: 13px;
        font-weight: 600;
        box-shadow: inset 0 -1px 0 rgba(0,0,0,0.08);
    }

    .view-btn:hover {
        background: #086bbd;
        color: #ffffff !important;
    }

    .no-document {
        color: #8794a8;
        font-size: 12px;
    }

    /* Attendance bar */
    .attendance-cell {
        display: flex;
        align-items: center;
        gap: 12px;
    }

    .attendance-track {
        width: 225px;
        max-width: 72%;
        height: 11px;
        background: #dce8f3;
        border-radius: 10px;
        overflow: hidden;
    }

    .attendance-fill {
        height: 100%;
        border-radius: 10px;
    }

    .attendance-value {
        min-width: 40px;
        color: #1e376e;
        font-size: 13px;
    }

    /* ---------- Table pagination ---------- */
    .pagination-row {
        display: flex;
        justify-content: flex-end;
        align-items: center;
        gap: 8px;
        width: 100%;
        margin-top: 5px;
    }

    .page-indicator {
        height: 36px;
        display: flex;
        align-items: center;
        justify-content: flex-end;
        color: #48628d;
        font-size: 13px;
        font-weight: 600;
        white-space: nowrap !important;
        overflow: visible !important;
    }

    /* Pagination buttons: icon only, fixed width, never wrap. */
    div[data-testid="stButton"] > button.pagination-button,
    div[data-testid="stButton"] > button {
        min-height: 36px !important;
        height: 36px !important;
        min-width: 42px !important;
        width: 42px !important;
        white-space: nowrap !important;
        word-break: keep-all !important;
        overflow-wrap: normal !important;
        padding: 0 !important;
        display: flex !important;
        align-items: center !important;
        justify-content: center !important;
    }

    div[data-testid="stButton"] > button p {
        white-space: nowrap !important;
        word-break: keep-all !important;
        overflow-wrap: normal !important;
        margin: 0 !important;
    }

    /* Empty message */
    .empty-msg {
        padding: 18px;
        color: #66758e;
        text-align: center;
        border: 1px solid #dce5ee;
        border-radius: 4px;
        background: #fbfdff;
    }

    /* Footer */
    .footer {
        margin-top: 6px;
        padding-top: 8px;
        border-top: 1px solid #d9e2eb;
        color: #274b83;
        font-size: 12px;
    }

    /* Streamlit selectbox */
    div[data-baseweb="select"] > div {
        min-height: 36px;
        border-color: #91a4bd;
        border-radius: 5px;
    }

    /* Refresh button */
    div[data-testid="stButton"] > button {
        height: 31px;
        border-radius: 6px;
        border: 1px solid #9aabc2;
        color: #123776;
        background: #ffffff;
        font-size: 12px;
        padding: 0 13px;
    }

    div[data-testid="stButton"] > button:hover {
        border-color: #167bd1;
        color: #0a65b5;
    }

    .data-status {
        color: #71819a;
        font-size: 11px;
        text-align: right;
        margin-top: 3px;
    }

    /* Remove extra gaps around columns */
    div[data-testid="column"] {
        padding-top: 0 !important;
        padding-bottom: 0 !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# Utility functions
# ============================================================

def clean_col_name(value):
    """Normalize a column name for flexible matching."""
    value = str(value).strip().lower()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalize_dataframe(df):
    """Clean column names and empty rows."""
    if df is None:
        return pd.DataFrame()

    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    # Remove completely empty rows
    df = df.dropna(how="all").reset_index(drop=True)

    # Convert all values to strings only where useful later.
    return df


def find_column(df, candidates):
    """
    Find a column using a list of possible names.
    Matching is case-insensitive and tolerant of spaces/symbols.
    """
    if df.empty:
        return None

    normalized = {
        clean_col_name(col): col
        for col in df.columns
    }

    # Exact normalized match
    for candidate in candidates:
        key = clean_col_name(candidate)
        if key in normalized:
            return normalized[key]

    # Partial match
    for candidate in candidates:
        key = clean_col_name(candidate)
        for norm_col, original_col in normalized.items():
            if key in norm_col or norm_col in key:
                return original_col

    return None



class _GoogleSheetHTMLParser(HTMLParser):
    """Extract hyperlink targets from Google Sheets HTML output by cell."""
    def __init__(self):
        super().__init__()
        self.row = -1
        self.col = -1
        self.in_cell = False
        self.current_href = None
        self.links = {}

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag.lower() == "tr":
            self.row += 1
            self.col = -1
        elif tag.lower() in ("td", "th"):
            self.col += 1
            self.in_cell = True
            self.current_href = None
        elif tag.lower() == "a" and self.in_cell:
            href = attrs.get("href")
            if href:
                self.current_href = href

    def handle_endtag(self, tag):
        if tag.lower() in ("td", "th"):
            if self.current_href:
                self.links[(self.row, self.col)] = self.current_href
            self.in_cell = False
            self.current_href = None

@st.cache_data(ttl=60, show_spinner=False)
def fetch_gviz_html_links(sheet_name):
    """
    Google Sheets rich-text links created with Insert -> Link can sometimes
    be omitted from the XLSX cell value/hyperlink metadata. The gviz HTML
    representation still exposes the actual href, so use it as a fallback.
    """
    try:
        url = (
            f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/gviz/tq"
            f"?tqx=out:html&sheet={quote(sheet_name)}"
        )
        response = requests.get(
            url,
            timeout=20,
            headers={"Cache-Control": "no-cache", "User-Agent": "Mozilla/5.0"},
        )
        response.raise_for_status()

        parser = _GoogleSheetHTMLParser()
        parser.feed(response.text)
        return parser.links
    except Exception:
        return {}

def find_document_column(df):
    """
    Find the most likely document/link column. If the expected header is
    missing or named differently, scan the sheet for cells containing URLs.
    """
    if df.empty:
        return None

    candidates = [
        "Document Link", "Document", "Document URL", "File Link",
        "File URL", "File", "PPT Link", "PPT", "MOM Link",
        "MOM", "Attachment", "URL", "Link", "Drive Link",
        "Google Drive Link", "Document Path"
    ]

    col = find_column(df, candidates)
    if col:
        # A Google Sheets hyperlink may appear in the exported dataframe only
        # as its visible text (for example: "PHA Procedure JJSL.pdf").
        # Therefore the presence of a URL in the cell is NOT required to
        # identify the document column.
        return col

    best_col = None
    best_score = 0

    for column in df.columns:
        score = sum(1 for value in df[column].tolist() if extract_url(value))
        if score > best_score:
            best_score = score
            best_col = column

    return best_col


def _column_letters(index_zero_based):
    """Convert a zero-based dataframe column index to Excel letters."""
    n = index_zero_based + 1
    result = ""
    while n:
        n, rem = divmod(n - 1, 26)
        result = chr(65 + rem) + result
    return result


@st.cache_data(ttl=60, show_spinner=False)
def fetch_gviz_column_links(sheet_name, column_index):
    """
    Fetch hyperlink targets from one Google Sheet column using the Google
    Visualization HTML endpoint.
    """
    try:
        col_letter = _column_letters(column_index)
        query = quote(f"select {col_letter}", safe="")
        url = (
            f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/gviz/tq"
            f"?tqx=out:html&sheet={quote(sheet_name)}&tq={query}"
        )
        response = requests.get(
            url,
            timeout=20,
            headers={
                "Cache-Control": "no-cache",
                "Pragma": "no-cache",
                "User-Agent": "Mozilla/5.0",
            },
        )
        response.raise_for_status()

        parser = _GoogleSheetHTMLParser()
        parser.feed(response.text)
        links = {}

        for (html_row, _html_col), href in parser.links.items():
            data_row = html_row - 1
            if data_row >= 0:
                href = unescape(str(href)).strip()
                if href.startswith("//"):
                    href = "https:" + href
                elif href.startswith("/"):
                    href = "https://docs.google.com" + href
                if href.startswith(("http://", "https://")):
                    links[data_row] = href

        # Extra safety if Google changes the table markup.
        if not links:
            hrefs = re.findall(
                r"href\s*=\s*[\"']([^\"']+)[\"']",
                response.text,
                flags=re.IGNORECASE,
            )
            data_row = 0
            for href in hrefs:
                href = unescape(href).strip()
                if href.startswith("//"):
                    href = "https:" + href
                elif href.startswith("/"):
                    href = "https://docs.google.com" + href
                if href.startswith(("http://", "https://")):
                    links[data_row] = href
                    data_row += 1

        return links
    except Exception:
        return {}


def extract_url(value):
    """
    Extract a document URL from:
      - a normal URL
      - HYPERLINK("url","text") / =HYPERLINK(...)
      - HTML anchor
      - Google Drive/Docs/Sheets/Slides links
    """
    if pd.isna(value):
        return None

    text = str(value).strip()
    if not text:
        return None

    # HYPERLINK formula first (important when the sheet exports formulas)
    patterns = [
        r'HYPERLINK\s*\(\s*["\'](https?://[^"\']+)["\']',
        r'HYPERLINK\s*\(\s*"(https?://[^"]+)"',
        r'HYPERLINK\s*\(\s*\'(https?://[^\']+)\'',
    ]
    for pattern in patterns:
        m = re.search(pattern, text, flags=re.IGNORECASE)
        if m:
            return m.group(1)

    # HTML anchor
    m = re.search(
        r'href\s*=\s*["\'](https?://[^"\']+)["\']',
        text,
        flags=re.IGNORECASE
    )
    if m:
        return m.group(1)

    # Any ordinary web URL
    m = re.search(r'https?://[^\s"\'<>]+', text)
    if m:
        return m.group(0).rstrip(').,;')

    # Common Google Drive file ID pasted as text.
    # Keep the existing behavior for IDs, but only after URL extraction.
    if text.startswith("drive.google.com/"):
        return "https://" + text

    # If the cell contains a Google Drive file ID only, make it a
    # Drive URL. This is also useful when users paste just the ID.
    if re.fullmatch(r'[A-Za-z0-9_-]{20,}', text):
        return f'https://drive.google.com/open?id={text}'

    return None


def display_value(value):
    """Convert a sheet value to safe display text."""
    if pd.isna(value):
        return ""

    # Excel/Google Sheets dates can arrive as pandas/Python datetime objects.
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.strftime("%d %b %Y")

    text = str(value).strip()

    # Hide HYPERLINK formula itself if formula somehow comes through.
    if text.upper().startswith("=HYPERLINK"):
        m = re.search(
            r'HYPERLINK\s*\(\s*["\'][^"\']+["\']\s*,\s*["\']([^"\']+)["\']',
            text,
            flags=re.IGNORECASE
        )
        if m:
            return m.group(1)

    return text


def display_date(value):
    """Display only the meeting date; never show 00:00:00/time."""
    if pd.isna(value):
        return ""

    if isinstance(value, (pd.Timestamp, datetime)):
        return value.strftime("%d %b %Y")

    text = str(value).strip()
    if not text:
        return ""

    parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if not pd.isna(parsed):
        return parsed.strftime("%d %b %Y")

    text = re.sub(r"\s+00:00:00(?:\.0+)?$", "", text)
    return text



def _excel_col_to_number(col_letters):
    """Convert Excel column letters (A, B, ..., AA) to a zero-based index."""
    number = 0
    for ch in col_letters.upper():
        number = number * 26 + (ord(ch) - ord("A") + 1)
    return number - 1


def extract_xlsx_hyperlinks(xlsx_bytes, workbook, sheet_name):
    """
    Low-level XLSX hyperlink fallback.

    Google Sheets can export a visually linked cell as an XLSX hyperlink
    relationship even when openpyxl does not expose it through
    cell.hyperlink. This function reads the worksheet XML and its .rels file
    directly and returns {(zero_based_row, zero_based_col): target_url}.
    """
    result = {}

    try:
        with ZipFile(BytesIO(xlsx_bytes)) as z:
            workbook_xml = ET.fromstring(
                z.read("xl/workbook.xml")
            )

            ns = {
                "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
            }

            # Find the worksheet relationship ID for the requested sheet.
            target_rid = None
            for sheet in workbook_xml.findall("main:sheets/main:sheet", ns):
                if sheet.attrib.get("name") == sheet_name:
                    target_rid = sheet.attrib.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )
                    break

            if not target_rid:
                return result

            workbook_rels = ET.fromstring(
                z.read("xl/_rels/workbook.xml.rels")
            )

            worksheet_target = None
            for rel in workbook_rels.findall("rel:Relationship", ns):
                if rel.attrib.get("Id") == target_rid:
                    worksheet_target = rel.attrib.get("Target")
                    break

            if not worksheet_target:
                return result

            if worksheet_target.startswith("/"):
                worksheet_path = worksheet_target.lstrip("/")
            else:
                worksheet_path = "xl/" + worksheet_target.lstrip("./")

            worksheet_path = worksheet_path.replace(
                "xl/xl/", "xl/"
            )

            if worksheet_path not in z.namelist():
                return result

            worksheet_xml = ET.fromstring(z.read(worksheet_path))

            # Relationship file for this worksheet.
            folder = worksheet_path.rsplit("/", 1)[0]
            filename = worksheet_path.rsplit("/", 1)[1]
            rels_path = f"{folder}/_rels/{filename}.rels"

            relationships = {}
            if rels_path in z.namelist():
                rels_xml = ET.fromstring(z.read(rels_path))
                for rel in rels_xml.findall("rel:Relationship", ns):
                    relationships[rel.attrib.get("Id")] = rel.attrib.get("Target")

            # Extract <hyperlink ref="D2" r:id="rIdX"/>.
            for hyperlink in worksheet_xml.findall(
                    ".//main:hyperlinks/main:hyperlink", ns
            ):
                ref = hyperlink.attrib.get("ref", "")
                rid = hyperlink.attrib.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )

                if not ref or not rid:
                    continue

                target = relationships.get(rid)
                if not target:
                    continue

                # Normalize relative targets.
                if target.startswith("/"):
                    url = target
                elif target.startswith("http://") or target.startswith("https://"):
                    url = target
                else:
                    url = target

                m = re.fullmatch(r"([A-Za-z]+)(\d+)", ref)
                if not m:
                    continue

                col_letters, row_number = m.groups()
                col_index = _excel_col_to_number(col_letters)
                row_index = int(row_number) - 1

                result[(row_index, col_index)] = url

    except Exception:
        return {}

    return result


@st.cache_data(ttl=60, show_spinner=False)
def download_google_workbook():
    """Download the Google workbook once and reuse it for all three tabs."""
    xlsx_url = (
        f"https://docs.google.com/spreadsheets/d/{SHEET_ID}"
        f"/export?format=xlsx"
    )

    response = requests.get(
        xlsx_url,
        timeout=12,
        headers={
            "Cache-Control": "no-cache",
            "Pragma": "no-cache",
            "User-Agent": "Mozilla/5.0",
        },
    )
    response.raise_for_status()
    return response.content


@st.cache_data(ttl=60, show_spinner=False)
def load_google_sheet(sheet_name):
    """
    Load the live Google Sheet.

    XLSX export is attempted first because it preserves cell hyperlinks.
    That means a Google Sheets cell using Insert -> Link or HYPERLINK()
    can still produce a working View button in the dashboard.

    CSV/gviz is used as a fallback for environments where XLSX export
    is unavailable.
    """
    try:
        import openpyxl

        xlsx_bytes = download_google_workbook()

        workbook = openpyxl.load_workbook(
            BytesIO(xlsx_bytes),
            data_only=False,
            read_only=False,
        )

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' was not found. "
                f"Available sheets: {', '.join(workbook.sheetnames)}"
            )

        ws = workbook[sheet_name]

        rows = list(ws.iter_rows())
        if not rows:
            return pd.DataFrame(), None

        # First non-empty row is treated as the header.
        header_row_index = 0
        for i, row in enumerate(rows):
            if any(
                    cell.value is not None and str(cell.value).strip() != ""
                    for cell in row
            ):
                header_row_index = i
                break

        header_cells = rows[header_row_index]
        headers = []
        seen = {}

        for i, cell in enumerate(header_cells):
            header = cell.value
            header = "" if header is None else str(header).strip()

            if not header:
                header = f"Column {i + 1}"

            # Make duplicate headers unique for pandas.
            base = header
            count = seen.get(base, 0)
            if count:
                header = f"{base}.{count}"
            seen[base] = count + 1
            headers.append(header)

        # Low-level hyperlink map. This is the important fallback for
        # Google Sheets rich-text hyperlinks that openpyxl may not expose.
        raw_xlsx_links = extract_xlsx_hyperlinks(
            xlsx_bytes,
            workbook,
            sheet_name,
        )

        data = []

        for row in rows[header_row_index + 1:]:
            values = []

            for cell in row[:len(headers)]:
                value = cell.value

                # First preference: hyperlink exposed by openpyxl.
                if cell.hyperlink and cell.hyperlink.target:
                    value = cell.hyperlink.target

                values.append(value)

            # Ignore completely empty rows.
            if any(
                    value is not None and str(value).strip() != ""
                    for value in values
            ):
                data.append(values)

        df = pd.DataFrame(data, columns=headers)

        # Apply raw XLSX hyperlink relationships. Excel row 1 is the header,
        # therefore worksheet row N maps to dataframe row N-2.
        for (xlsx_row, xlsx_col), href in raw_xlsx_links.items():
            data_row = xlsx_row - (header_row_index + 1)

            if (
                    0 <= data_row < len(df)
                    and 0 <= xlsx_col < len(df.columns)
            ):
                current = df.iat[data_row, xlsx_col]
                if not extract_url(current):
                    df.iat[data_row, xlsx_col] = href

        # Final fallback for rich-text links. The Google Visualization HTML
        # endpoint may expose hrefs even when XLSX does not.
        if not raw_xlsx_links and not df.empty:
            html_links = fetch_gviz_html_links(sheet_name)

            # HTML row 0 is the header, so data rows start at 1.
            for (html_row, html_col), href in html_links.items():
                data_row = html_row - 1
                if (
                        0 <= data_row < len(df)
                        and 0 <= html_col < len(df.columns)
                ):
                    current = df.iat[data_row, html_col]
                    if not extract_url(current):
                        df.iat[data_row, html_col] = href

        return normalize_dataframe(df), None

    except Exception as xlsx_error:
        # Fallback: Google Visualization CSV endpoint.
        csv_url = (
            f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/gviz/tq"
            f"?tqx=out:csv&sheet={quote(sheet_name)}"
            f"&cachebust={int(datetime.now().timestamp())}"
        )

        try:
            df = pd.read_csv(
                csv_url,
                storage_options={"User-Agent": "Mozilla/5.0"},
            )
            return normalize_dataframe(df), None
        except Exception as csv_error:
            return (
                pd.DataFrame(),
                f"XLSX error: {xlsx_error}; CSV error: {csv_error}"
            )


@st.cache_data(ttl=60, show_spinner=False)
def enrich_document_links(df, sheet_name):
    """Merge real Google Sheets document hrefs into the dataframe."""
    if df.empty:
        return df

    df = df.copy()
    doc_col = find_document_column(df)
    if not doc_col:
        return df

    col_index = list(df.columns).index(doc_col)
    link_map = fetch_gviz_column_links(sheet_name, col_index)

    for data_row, href in link_map.items():
        if 0 <= data_row < len(df):
            current = df.iat[data_row, col_index]
            if not extract_url(current):
                df.iat[data_row, col_index] = href

    return df


def get_year_series(df):
    """
    Find a year column first. If there is no year column, try date columns.
    Returns a Series of nullable integer years.
    """
    if df.empty:
        return pd.Series(dtype="Int64")

    year_col = find_column(
        df,
        ["Year", "Financial Year", "FY", "Calendar Year"]
    )

    if year_col:
        raw = df[year_col].astype(str).str.extract(r"(20\d{2}|19\d{2})")[0]
        return pd.to_numeric(raw, errors="coerce").astype("Int64")

    date_col = find_column(
        df,
        [
            "Meeting Date",
            "Date",
            "Document Date",
            "Training Date",
            "Created Date",
            "Date of Meeting",
        ]
    )

    if date_col:
        dates = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True)
        return dates.dt.year.astype("Int64")

    # Last fallback: inspect every column for a four-digit year.
    for col in df.columns:
        extracted = df[col].astype(str).str.extract(r"\b(20\d{2}|19\d{2})\b")[0]
        if extracted.notna().sum() > 0:
            years = pd.to_numeric(extracted, errors="coerce")
            if years.notna().sum() >= max(1, int(len(df) * 0.30)):
                return years.astype("Int64")

    return pd.Series(pd.NA, index=df.index, dtype="Int64")


def filter_by_year(df, selected_year):
    if df.empty:
        return df.copy()

    years = get_year_series(df)

    # If no year can be detected, don't remove data.
    if years.notna().sum() == 0:
        return df.copy()

    return df[years == int(selected_year)].copy()


def available_years(dataframes):
    years = set()

    for df in dataframes:
        if df.empty:
            continue
        y = get_year_series(df)
        for value in y.dropna().tolist():
            try:
                years.add(int(value))
            except Exception:
                pass

    # Always provide current year as a useful option.
    years.add(datetime.now().year)

    return sorted(years, reverse=True)


def parse_percentage(value):
    """Convert 92%, 92, 0.92 etc. into a percentage number."""
    if pd.isna(value):
        return np.nan

    text = str(value).strip().replace(",", "")

    m = re.search(r"-?\d+(?:\.\d+)?", text)
    if not m:
        return np.nan

    number = float(m.group())

    if "%" in text:
        return max(0, min(100, number))

    # If a decimal fraction is supplied, e.g. 0.92
    if 0 <= number <= 1:
        return number * 100

    return max(0, min(100, number))


def attendance_html(value):
    pct = parse_percentage(value)

    if np.isnan(pct):
        return '<span class="no-document">—</span>'

    # Attendance colour rules:
    # 0–40%   = red
    # 41–90%  = yellow
    # >90%    = green
    if pct <= 40:
        style = "background:#e3182d;"
    elif pct <= 90:
        style = "background:#f4b400;"
    else:
        style = "background:#10b981;"

    return f"""
    <div class="attendance-cell">
        <div class="attendance-track">
            <div class="attendance-fill" style="width:{pct:.0f}%;{style}"></div>
        </div>
        <span class="attendance-value">{pct:.0f}%</span>
    </div>
    """


def document_link_html(value):
    """
    Create the View button from the document URL stored in Google Sheets.
    Supports direct URLs, Google Drive, Google Docs, Google Slides, and
    HYPERLINK formulas.
    """
    url = extract_url(value)

    if not url:
        return '<span class="no-document">Not available</span>'

    safe_url = (
        url.replace("&", "&amp;")
        .replace('"', "&quot;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    return (
        f'<a class="view-btn" href="{safe_url}" '
        f'target="_blank" rel="noopener noreferrer">View</a>'
    )


def search_dataframe(df, query):
    if df.empty or not query:
        return df

    q = str(query).strip().lower()
    if not q:
        return df

    mask = pd.Series(False, index=df.index)

    for col in df.columns:
        mask = mask | df[col].astype(str).str.lower().str.contains(
            q, regex=False, na=False
        )

    return df[mask].copy()


def make_table(headers, rows, widths=None):
    """
    Build an HTML table.
    rows should already contain HTML-safe strings.
    """
    width_html = ""

    if widths:
        width_html = "<colgroup>"
        for width in widths:
            width_html += f'<col style="width:{width};">'
        width_html += "</colgroup>"

    html = '<div class="table-wrap"><table class="repo-table">'
    html += width_html
    html += "<thead><tr>"

    for h in headers:
        html += f"<th>{h}</th>"

    html += "</tr></thead><tbody>"

    if not rows:
        html += (
            f'<tr><td colspan="{len(headers)}" '
            f'style="text-align:center;color:#718096;padding:15px;">'
            f"No records found</td></tr>"
        )
    else:
        for row in rows:
            html += "<tr>"
            for cell in row:
                html += f"<td>{cell}</td>"
            html += "</tr>"

    html += "</tbody></table></div>"

    return html


ROWS_PER_PAGE = 5


def paginated_table(headers, rows, widths, state_key):
    """
    Display a maximum of five rows at a time with Previous/Next controls.
    The controls appear only when more than five records exist.
    """
    total_rows = len(rows)
    total_pages = max(1, (total_rows + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE)

    if state_key not in st.session_state:
        st.session_state[state_key] = 0

    page = int(st.session_state[state_key])
    page = max(0, min(page, total_pages - 1))
    st.session_state[state_key] = page

    start_row = page * ROWS_PER_PAGE
    end_row = min(start_row + ROWS_PER_PAGE, total_rows)
    visible_rows = rows[start_row:end_row]

    st.markdown(
        make_table(headers, visible_rows, widths),
        unsafe_allow_html=True,
    )

    if total_pages > 1:
        # All pagination elements stay on one horizontal line at the
        # lower-right: "Page 1 of 2 · 10 records"  ‹  ›
        spacer, indicator_col, previous_col, next_col = st.columns(
            [7.8, 2.0, 0.42, 0.42],
            gap="small",
        )

        with indicator_col:
            st.markdown(
                f'<div class="page-indicator">Page {page + 1} of {total_pages}'
                f' &nbsp;·&nbsp; {total_rows} records</div>',
                unsafe_allow_html=True,
            )

        with previous_col:
            previous = st.button(
                "‹",
                key=f"{state_key}_previous",
                disabled=(page == 0),
                help="Previous page",
                use_container_width=True,
            )

        with next_col:
            next_page = st.button(
                "›",
                key=f"{state_key}_next",
                disabled=(page >= total_pages - 1),
                help="Next page",
                use_container_width=True,
            )

        if previous:
            st.session_state[state_key] = max(0, page - 1)
            st.rerun()

        if next_page:
            st.session_state[state_key] = min(total_pages - 1, page + 1)
            st.rerun()


# ============================================================
# Load data
# ============================================================

mom_df, mom_error = load_google_sheet(SHEET_NAMES["mom"])
training_df, training_error = load_google_sheet(SHEET_NAMES["training"])
procedure_df, procedure_error = load_google_sheet(SHEET_NAMES["procedure"])

# Resolve rich-text hyperlinks separately when the XLSX export contains only
# the visible filename.
mom_df = enrich_document_links(mom_df, SHEET_NAMES["mom"])
training_df = enrich_document_links(training_df, SHEET_NAMES["training"])
procedure_df = enrich_document_links(procedure_df, SHEET_NAMES["procedure"])

all_errors = []
for name, error in [
    ("MOM", mom_error),
    ("TRAINING PPTS", training_error),
    ("PSM Procedure", procedure_error),
]:
    if error:
        all_errors.append(f"{name}: {error}")

if all_errors:
    st.error(
        "Unable to read one or more Google Sheet tabs. "
        "Make sure the spreadsheet is shared as 'Anyone with the link – Viewer'.\n\n"
        + "\n".join(all_errors)
    )
    st.stop()


# ============================================================
# Year selector
# ============================================================

years = available_years([mom_df, training_df, procedure_df])

default_year = datetime.now().year
if default_year not in years:
    default_year = years[0] if years else datetime.now().year

year_label_col, year_filter_col, refresh_col, spacer = st.columns(
    [0.85, 2.0, 0.42, 6.73],
    gap="small"
)

with year_label_col:
    st.markdown(
        '<div class="year-inline-label">Select Year:</div>',
        unsafe_allow_html=True,
    )

with year_filter_col:
    selected_year = st.selectbox(
        "Select Year",
        years,
        index=years.index(default_year) if default_year in years else 0,
        label_visibility="collapsed",
    )

with refresh_col:
    refresh_clicked = st.button(
        "↻",
        help="Refresh latest Google Sheet data",
        key="refresh_data",
    )
    if refresh_clicked:
        st.cache_data.clear()
        st.rerun()

# Filter all three repositories
mom = filter_by_year(mom_df, selected_year)
training = filter_by_year(training_df, selected_year)
procedures = filter_by_year(procedure_df, selected_year)

# Meetings are always displayed newest first.
# Use the meeting date column when available; invalid/blank dates go last.
mom_date_sort_col = find_column(
    mom,
    ["Meeting Date", "Date", "Date of Meeting", "MOM Date"]
)
if mom_date_sort_col and not mom.empty:
    _mom_sort_dates = pd.to_datetime(
        mom[mom_date_sort_col],
        errors="coerce",
        dayfirst=True,
    )
    mom = (
        mom.assign(_sort_meeting_date=_mom_sort_dates)
        .sort_values(
            "_sort_meeting_date",
            ascending=False,
            na_position="last",
            kind="stable",
        )
        .drop(columns=["_sort_meeting_date"])
        .reset_index(drop=True)
    )


# ============================================================
# KPI cards
# ============================================================

k1, k2, k3 = st.columns(3, gap="medium")

with k1:
    st.markdown(
        f"""
        <div class="kpi-card kpi-red">
            <div class="kpi-title">PSM Procedures</div>
            <div class="kpi-number red">{len(procedures):,}</div>
            <div class="kpi-subtitle">Documents</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with k2:
    st.markdown(
        f"""
        <div class="kpi-card kpi-blue">
            <div class="kpi-title">Meetings</div>
            <div class="kpi-number blue">{len(mom):,}</div>
            <div class="kpi-subtitle">Total Meetings</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with k3:
    st.markdown(
        f"""
        <div class="kpi-card kpi-training">
            <div class="kpi-title">Training PPTs</div>
            <div class="kpi-number training-number">{len(training):,}</div>
            <div class="kpi-subtitle">Documents</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.markdown("<div style='height:6px'></div>", unsafe_allow_html=True)


# ============================================================
# PSM PROCEDURES
# ============================================================

title_col, search_col = st.columns([5.2, 2.3], gap="medium")

with title_col:
    st.markdown(
        '<div class="section-title procedure-title">PSM Procedures</div>',
        unsafe_allow_html=True,
    )

with search_col:
    procedure_search = st.text_input(
        "Search Procedures",
        placeholder="⌕  Search Procedures...",
        label_visibility="collapsed",
        key="procedure_search",
    )

if st.session_state.get("_last_procedure_search") != procedure_search:
    st.session_state["procedure_page"] = 0
    st.session_state["_last_procedure_search"] = procedure_search

procedure_view = search_dataframe(procedures, procedure_search)

proc_title = find_column(
    procedure_view,
    ["Procedure Title", "Title", "Procedure Name", "Document Title", "Name"]
)
proc_no = find_column(
    procedure_view,
    ["Document No.", "Document No", "Document Number", "Doc No", "No."]
)
proc_revision = find_column(
    procedure_view,
    ["Revision", "Rev", "Revision No.", "Rev No"]
)
proc_document = find_document_column(procedure_view)

proc_rows = []

for _, row in procedure_view.iterrows():
    title = display_value(row[proc_title]) if proc_title else ""
    doc_no = display_value(row[proc_no]) if proc_no else ""
    revision = display_value(row[proc_revision]) if proc_revision else ""
    document = document_link_html(row[proc_document]) if proc_document else (
        '<span class="no-document">Not available</span>'
    )

    proc_rows.append([
        title,
        doc_no,
        revision,
        document,
    ])

paginated_table(
    ["Procedure Title", "Document No.", "Revision", "Document"],
    proc_rows,
    ["36%", "20%", "20%", "24%"],
    "procedure_page",
)

st.markdown("<div style='height:6px'></div>", unsafe_allow_html=True)


# ============================================================
# MEETINGS / MOM
# ============================================================

title_col, search_col = st.columns([5.2, 2.3], gap="medium")

with title_col:
    st.markdown(
        '<div class="section-title meeting-title">Meetings</div>',
        unsafe_allow_html=True,
    )

with search_col:
    meeting_search = st.text_input(
        "Search Meetings",
        placeholder="⌕  Search Meetings (Name / Date)...",
        label_visibility="collapsed",
        key="meeting_search",
    )

if st.session_state.get("_last_meeting_search") != meeting_search:
    st.session_state["meeting_page"] = 0
    st.session_state["_last_meeting_search"] = meeting_search

mom_view = search_dataframe(mom, meeting_search)

meeting_name = find_column(
    mom_view,
    ["Meeting Name", "Meeting", "Name", "Title", "Meeting Title"]
)
meeting_date = find_column(
    mom_view,
    ["Meeting Date", "Date", "Date of Meeting", "MOM Date"]
)
attendance = find_column(
    mom_view,
    ["Attendance", "Attendance %", "Attendance Percentage", "Present %"]
)
mom_document = find_document_column(mom_view)

mom_rows = []

for _, row in mom_view.iterrows():
    name = display_value(row[meeting_name]) if meeting_name else ""
    date = display_date(row[meeting_date]) if meeting_date else ""
    attend = attendance_html(row[attendance]) if attendance else (
        '<span class="no-document">—</span>'
    )
    document = document_link_html(row[mom_document]) if mom_document else (
        '<span class="no-document">Not available</span>'
    )

    mom_rows.append([
        name,
        date,
        attend,
        document,
    ])

paginated_table(
    ["Meeting Name", "Meeting Date", "Attendance", "Document"],
    mom_rows,
    ["28%", "20%", "36%", "16%"],
    "meeting_page",
)

st.markdown("<div style='height:6px'></div>", unsafe_allow_html=True)


# ============================================================
# TRAINING PPTS
# ============================================================

title_col, search_col = st.columns([5.2, 2.3], gap="medium")

with title_col:
    st.markdown(
        '<div class="section-title training-title">Training PPTs</div>',
        unsafe_allow_html=True,
    )

with search_col:
    training_search = st.text_input(
        "Search Training",
        placeholder="⌕  Search Training PPTs...",
        label_visibility="collapsed",
        key="training_search",
    )

if st.session_state.get("_last_training_search") != training_search:
    st.session_state["training_page"] = 0
    st.session_state["_last_training_search"] = training_search

training_view = search_dataframe(training, training_search)

training_title = find_column(
    training_view,
    ["Title", "Training Title", "PPT Title", "Name", "Document Title"]
)
training_category = find_column(
    training_view,
    ["Category", "Type", "Training Category"]
)
training_document = find_document_column(training_view)

training_rows = []

for _, row in training_view.iterrows():
    title = display_value(row[training_title]) if training_title else ""
    category = display_value(row[training_category]) if training_category else ""
    document = document_link_html(row[training_document]) if training_document else (
        '<span class="no-document">Not available</span>'
    )

    training_rows.append([
        title,
        category,
        document,
    ])

paginated_table(
    ["Title", "Category", "Document"],
    training_rows,
    ["48%", "32%", "20%"],
    "training_page",
)


# ============================================================
# Footer
# ============================================================

st.markdown(
    '<div class="data-status">Data source: Google Sheets • Use “↻ Refresh Data” after making sheet changes.</div>',
    unsafe_allow_html=True,
)
